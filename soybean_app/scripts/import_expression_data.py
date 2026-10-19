from __future__ import annotations

import argparse
import hashlib
import math
import sys
from dataclasses import dataclass
from pathlib import Path
from statistics import fmean
from typing import Any, Iterable

from openpyxl import load_workbook

from soybean_app.core.foundation import (
    BASE_DIR,
    DB_CONFIG,
    EXPRESSION_DATASETS,
    EXPRESSION_TIMEPOINTS,
    get_species_option,
    psycopg,
)


SCHEMA_PATH = BASE_DIR / "deploy" / "postgres" / "initdb" / "04-expression-data.sql"
REQUIRED_METADATA_COLUMNS = ("Protein accession", "Protein description", "Gene name")


@dataclass(frozen=True)
class ExpressionProtein:
    protein_accession: str
    gene_name: str
    protein_description: str


@dataclass(frozen=True)
class ExpressionMeasurement:
    protein_accession: str
    timepoint: str
    replicate_no: int
    abundance: float


@dataclass(frozen=True)
class ParsedExpressionDataset:
    species_key: str
    dataset_name: str
    organism: str
    source_file: str
    source_sheet: str
    source_sha256: str
    proteins: tuple[ExpressionProtein, ...]
    measurements: tuple[ExpressionMeasurement, ...]


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _text(value: Any) -> str:
    return str(value or "").strip()


def _measurement_number(value: Any, *, row_number: int, column_name: str) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, str) and value.strip().startswith("#"):
        return None
    if isinstance(value, bool):
        raise ValueError(f"Row {row_number}, column {column_name}: boolean is not a valid abundance.")
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(
            f"Row {row_number}, column {column_name}: {value!r} is not a numeric abundance."
        ) from exc
    if not math.isfinite(number):
        raise ValueError(f"Row {row_number}, column {column_name}: abundance must be finite.")
    return number


def parse_expression_workbook(species_key: str) -> ParsedExpressionDataset:
    config = EXPRESSION_DATASETS.get(species_key)
    if not config:
        raise ValueError(f"Unknown expression dataset: {species_key}")

    path = Path(config["path"])
    sheet_name = str(config["sheet"])
    if not path.is_file():
        raise FileNotFoundError(f"Expression workbook not found: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet {sheet_name!r} was not found in {path.name}.")
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration as exc:
            raise ValueError(f"Worksheet {sheet_name!r} is empty.") from exc

        headers = [_text(value) for value in header_row]
        column_index = {header: index for index, header in enumerate(headers) if header}
        required_columns = set(REQUIRED_METADATA_COLUMNS)
        required_columns.update(EXPRESSION_TIMEPOINTS)
        required_columns.update(
            f"{timepoint}_{replicate_no}"
            for timepoint in EXPRESSION_TIMEPOINTS
            for replicate_no in range(1, 4)
        )
        missing_columns = sorted(required_columns.difference(column_index))
        if missing_columns:
            raise ValueError(f"Missing required columns: {', '.join(missing_columns)}")

        proteins: list[ExpressionProtein] = []
        measurements: list[ExpressionMeasurement] = []
        accession_rows: dict[str, int] = {}

        for row_number, row in enumerate(rows, start=2):
            accession = _text(row[column_index["Protein accession"]])
            if not accession:
                if any(value not in (None, "") for value in row):
                    raise ValueError(f"Row {row_number}: Protein accession is required.")
                continue
            if accession in accession_rows:
                raise ValueError(
                    f"Duplicate Protein accession {accession!r} in rows "
                    f"{accession_rows[accession]} and {row_number}."
                )
            accession_rows[accession] = row_number

            proteins.append(
                ExpressionProtein(
                    protein_accession=accession,
                    gene_name=_text(row[column_index["Gene name"]]),
                    protein_description=_text(row[column_index["Protein description"]]),
                )
            )

            for timepoint in EXPRESSION_TIMEPOINTS:
                reported_mean = _measurement_number(
                    row[column_index[timepoint]],
                    row_number=row_number,
                    column_name=timepoint,
                )
                replicate_values: list[float] = []
                for replicate_no in range(1, 4):
                    column_name = f"{timepoint}_{replicate_no}"
                    abundance = _measurement_number(
                        row[column_index[column_name]],
                        row_number=row_number,
                        column_name=column_name,
                    )
                    if abundance is None:
                        continue
                    replicate_values.append(abundance)
                    measurements.append(
                        ExpressionMeasurement(
                            protein_accession=accession,
                            timepoint=timepoint,
                            replicate_no=replicate_no,
                            abundance=abundance,
                        )
                    )

                if reported_mean is not None and not replicate_values:
                    raise ValueError(
                        f"Row {row_number}, {timepoint}: a reported mean exists without replicate values."
                    )
                if reported_mean is not None and not math.isclose(
                    reported_mean,
                    fmean(replicate_values),
                    rel_tol=1e-9,
                    abs_tol=1e-9,
                ):
                    raise ValueError(
                        f"Row {row_number}, {timepoint}: reported mean {reported_mean!r} does not match "
                        f"the replicate mean {fmean(replicate_values)!r}."
                    )

    finally:
        workbook.close()

    species_option = get_species_option(species_key)
    return ParsedExpressionDataset(
        species_key=species_key,
        dataset_name=str(config["label"]),
        organism=str(species_option.get("display", species_key)),
        source_file=path.name,
        source_sheet=sheet_name,
        source_sha256=_file_sha256(path),
        proteins=tuple(proteins),
        measurements=tuple(measurements),
    )


def _schema_statements(path: Path = SCHEMA_PATH) -> Iterable[str]:
    if not path.is_file():
        raise FileNotFoundError(f"Expression schema file not found: {path}")
    sql_lines = [
        line
        for line in path.read_text(encoding="utf-8").splitlines()
        if not line.lstrip().startswith("\\")
    ]
    for statement in "\n".join(sql_lines).split(";"):
        clean_statement = statement.strip()
        if clean_statement:
            yield clean_statement


def ensure_expression_schema(cursor: Any) -> None:
    for statement in _schema_statements():
        cursor.execute(statement)


def replace_expression_dataset(cursor: Any, dataset: ParsedExpressionDataset) -> tuple[int, int]:
    cursor.execute(
        """
        INSERT INTO ptm.expression_dataset (
          dataset_key,
          dataset_name,
          organism,
          source_file,
          source_sheet,
          source_sha256,
          protein_count,
          measurement_count,
          imported_at
        )
        VALUES (%s, %s, %s, %s, %s, %s, 0, 0, NOW())
        ON CONFLICT (dataset_key) DO UPDATE SET
          dataset_name = EXCLUDED.dataset_name,
          organism = EXCLUDED.organism,
          source_file = EXCLUDED.source_file,
          source_sheet = EXCLUDED.source_sheet,
          source_sha256 = EXCLUDED.source_sha256,
          protein_count = 0,
          measurement_count = 0,
          imported_at = NOW()
        RETURNING dataset_id;
        """,
        (
            dataset.species_key,
            dataset.dataset_name,
            dataset.organism,
            dataset.source_file,
            dataset.source_sheet,
            dataset.source_sha256,
        ),
    )
    dataset_id = int(cursor.fetchone()[0])

    cursor.execute("DELETE FROM ptm.expression_protein WHERE dataset_id = %s;", (dataset_id,))
    with cursor.copy(
        """
        COPY ptm.expression_protein (
          dataset_id,
          protein_accession,
          gene_name,
          protein_description
        ) FROM STDIN
        """
    ) as copy:
        for protein in dataset.proteins:
            copy.write_row(
                (
                    dataset_id,
                    protein.protein_accession,
                    protein.gene_name,
                    protein.protein_description,
                )
            )

    cursor.execute(
        """
        SELECT expression_protein_id, protein_accession
        FROM ptm.expression_protein
        WHERE dataset_id = %s;
        """,
        (dataset_id,),
    )
    protein_ids = {str(accession): int(protein_id) for protein_id, accession in cursor.fetchall()}
    if len(protein_ids) != len(dataset.proteins):
        raise RuntimeError("Imported protein count does not match the parsed workbook.")

    with cursor.copy(
        """
        COPY ptm.expression_measurement (
          expression_protein_id,
          timepoint,
          replicate_no,
          abundance
        ) FROM STDIN
        """
    ) as copy:
        for measurement in dataset.measurements:
            copy.write_row(
                (
                    protein_ids[measurement.protein_accession],
                    measurement.timepoint,
                    measurement.replicate_no,
                    measurement.abundance,
                )
            )

    cursor.execute(
        """
        UPDATE ptm.expression_dataset
        SET protein_count = %s,
            measurement_count = %s
        WHERE dataset_id = %s;
        """,
        (len(dataset.proteins), len(dataset.measurements), dataset_id),
    )
    return len(dataset.proteins), len(dataset.measurements)


def _connect() -> Any:
    if psycopg is None:
        raise RuntimeError("psycopg is required. Install requirements.txt before importing data.")
    return psycopg.connect(
        host=DB_CONFIG["host"],
        port=DB_CONFIG["port"],
        dbname=DB_CONFIG["dbname"],
        user=DB_CONFIG["user"],
        password=DB_CONFIG["password"],
        connect_timeout=DB_CONFIG["connect_timeout"],
        autocommit=False,
    )


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Import quantitative proteome Excel workbooks into PostgreSQL."
    )
    parser.add_argument(
        "--species",
        action="append",
        choices=sorted(EXPRESSION_DATASETS),
        help="Import only this dataset. Repeat the option to select more than one.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Validate and summarize the workbooks without changing PostgreSQL.",
    )
    return parser.parse_args()


def main() -> int:
    args = _parse_args()
    species_keys = args.species or list(EXPRESSION_DATASETS)
    parsed_datasets: list[ParsedExpressionDataset] = []

    for species_key in species_keys:
        print(f"Validating {species_key}...", flush=True)
        dataset = parse_expression_workbook(species_key)
        parsed_datasets.append(dataset)
        print(
            f"  {len(dataset.proteins):,} proteins, "
            f"{len(dataset.measurements):,} replicate measurements, "
            f"SHA-256 {dataset.source_sha256[:12]}..."
        )

    if args.dry_run:
        print("Dry run completed. PostgreSQL was not changed.")
        return 0

    connection = _connect()
    try:
        with connection.cursor() as cursor:
            ensure_expression_schema(cursor)
            for dataset in parsed_datasets:
                protein_count, measurement_count = replace_expression_dataset(cursor, dataset)
                print(
                    f"Imported {dataset.species_key}: {protein_count:,} proteins and "
                    f"{measurement_count:,} measurements."
                )
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()

    print("Expression data import completed successfully.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"Import failed: {exc}", file=sys.stderr)
        raise SystemExit(1) from exc
