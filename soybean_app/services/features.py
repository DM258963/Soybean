from soybean_app.core.foundation import *

def _feature_int(value: Any) -> int | None:
    try:
        number = int(value)
    except (TypeError, ValueError):
        return None
    return number if number > 0 else None


def _uniprot_feature_ranges(accession: str, protein_length: int) -> list[dict[str, Any]]:
    normalized = _normalize_accession(accession)
    if not normalized or protein_length <= 0:
        return []

    entry, _ = fetch_uniprot_entry(normalized)
    if not entry:
        return []

    allowed_types = {
        "DOMAIN",
        "REGION",
        "MOTIF",
        "REPEAT",
        "TOPOLOGICAL DOMAIN",
        "ZINC FINGER",
    }
    feature_rows: list[dict[str, Any]] = []
    for feat in entry.get("features", []) or []:
        raw_type = str(feat.get("type", "")).strip()
        normalized_type = raw_type.upper()
        if normalized_type not in allowed_types:
            continue

        start = _feature_int(_nested_value(feat, ["location", "start", "value"]))
        end = _feature_int(_nested_value(feat, ["location", "end", "value"]))
        if start is None or end is None:
            continue
        start = max(1, min(start, protein_length))
        end = max(1, min(end, protein_length))
        if end < start:
            start, end = end, start

        description = str(feat.get("description", "")).strip()
        label = description or raw_type.title()
        feature_rows.append(
            {
                "type": raw_type.title(),
                "start": start,
                "end": end,
                "label": label,
                "title": f"{raw_type.title()} {start}-{end}: {label}",
            }
        )
        if len(feature_rows) >= 12:
            break

    return feature_rows


def _assign_feature_lanes(features: list[dict[str, Any]], max_lanes: int = 3) -> int:
    lane_ends = [0] * max_lanes
    lane_count = 1
    for feature in sorted(features, key=lambda item: (item["start"], item["end"])):
        selected_lane = 0
        for lane_idx, lane_end in enumerate(lane_ends):
            if feature["start"] > lane_end:
                selected_lane = lane_idx
                break
        else:
            selected_lane = min(range(max_lanes), key=lambda idx: lane_ends[idx])

        feature["lane"] = selected_lane
        lane_ends[selected_lane] = max(lane_ends[selected_lane], int(feature["end"]))
        lane_count = max(lane_count, selected_lane + 1)
    return lane_count


def build_protein_ptm_map(
    sequence: str,
    result_rows: list[dict[str, Any]],
    result_meta: dict[str, Any],
) -> dict[str, Any]:
    positions = [
        int(row.get("Position"))
        for row in result_rows
        if isinstance(row.get("Position"), int) and int(row.get("Position")) > 0
    ]
    protein_length = len(sequence) if sequence else (max(positions) if positions else 0)
    if protein_length <= 0:
        return {"available": False}

    left = 72
    width = 856

    def x_for_position(position: int) -> float:
        if protein_length <= 1:
            return left + width / 2
        return left + ((position - 1) / (protein_length - 1)) * width

    feature_blocks = _uniprot_feature_ranges(str(result_meta.get("protein_accession", "")), protein_length)
    feature_lane_count = _assign_feature_lanes(feature_blocks) if feature_blocks else 1
    axis_y = 148 + feature_lane_count * 28 + 12
    if feature_blocks:
        feature_palette = {
            "Domain": "#7fb8d6",
            "Region": "#d8e2c4",
            "Motif": "#efc36f",
            "Repeat": "#8cc7b0",
            "Topological Domain": "#b9d18a",
            "Zinc Finger": "#c4a0d8",
        }
        for feature in feature_blocks:
            start = int(feature["start"])
            end = int(feature["end"])
            feature["x"] = round(x_for_position(start), 2)
            feature["width"] = round(max(x_for_position(end) - x_for_position(start), 6), 2)
            feature["y"] = 148 + int(feature.get("lane", 0)) * 28
            feature["fill"] = feature_palette.get(str(feature.get("type", "")), "#d8e2c4")
            feature["show_label"] = feature["width"] >= 58
    else:
        feature_blocks = [
            {
                "type": "Protein",
                "start": 1,
                "end": protein_length,
                "label": "Protein sequence",
                "title": f"Protein sequence 1-{protein_length}",
                "lane": 0,
                "x": left,
                "width": width,
                "y": 148,
                "fill": "#d8e2c4",
                "show_label": True,
            }
        ]

    site_buckets: dict[int, dict[str, Any]] = {}
    for row in result_rows:
        position = row.get("Position")
        if not isinstance(position, int) or position <= 0 or position > protein_length:
            continue
        ptm_code = str(row.get("PTM types", "")).strip()
        if not is_visual_ptm_code(ptm_code):
            continue
        bucket = site_buckets.setdefault(
            position,
            {
                "position": position,
                "amino_acids": set(),
                "ptm_codes": set(),
                "site_ids": [],
            },
        )
        amino_acid = str(row.get("Amino acid display") or row.get("Amino acid") or "").strip()
        if amino_acid and amino_acid != "-":
            bucket["amino_acids"].add(amino_acid)
        bucket["ptm_codes"].add(ptm_code)
        site_id = str(row.get("Site id", "")).strip()
        if site_id:
            bucket["site_ids"].append(site_id)

    lane_right_edges = [-1000.0] * 4
    sites: list[dict[str, Any]] = []
    for bucket in sorted(site_buckets.values(), key=lambda item: item["position"]):
        position = int(bucket["position"])
        x_pos = x_for_position(position)
        lane = 0
        for idx, right_edge in enumerate(lane_right_edges):
            if x_pos - right_edge >= 28:
                lane = idx
                break
        lane_right_edges[lane] = x_pos

        ordered_codes = order_ptm_codes(set(bucket["ptm_codes"]))
        primary_code = ordered_codes[0] if ordered_codes else ""
        style = get_ptm_style(primary_code)
        is_multiple = len(ordered_codes) > 1
        amino_text = "/".join(sorted(bucket["amino_acids"])) or "-"
        y_pos = 78 - lane * 17
        sites.append(
            {
                "site_id": bucket["site_ids"][0] if bucket["site_ids"] else "",
                "site_ids": ",".join(bucket["site_ids"]),
                "position": position,
                "amino_acid": amino_text,
                "ptm_codes": ordered_codes,
                "ptm_text": ", ".join(ordered_codes),
                "x": round(x_pos, 2),
                "y": y_pos,
                "line_y": 132,
                "bg": "#000000" if is_multiple else style["bg"],
                "fg": "#ffffff" if is_multiple else style["fg"],
                "label": f"{amino_text}{position}",
                "show_label": len(site_buckets) <= 24,
                "title": f"Position {position} | {amino_text} | PTM: {', '.join(ordered_codes)}",
            }
        )

    accession = str(result_meta.get("protein_accession", "")).strip()
    gene = str(result_meta.get("gene_name", "")).strip()
    description = str(result_meta.get("protein_description", "")).strip()
    title_parts = [part for part in [accession or gene, description] if part]
    return {
        "available": bool(sites),
        "protein_title": ": ".join(title_parts) if title_parts else "Protein PTM Map",
        "length": protein_length,
        "site_count": len(sites),
        "features": feature_blocks,
        "feature_lane_count": feature_lane_count,
        "sites": sites,
        "view_height": 238 + max(feature_lane_count - 1, 0) * 28,
        "axis": {
            "left": left,
            "right": left + width,
            "width": width,
            "track_y": 132,
            "axis_y": axis_y,
            "label_y": axis_y + 20,
            "start": 1,
            "middle": max(1, round(protein_length / 2)),
            "end": protein_length,
            "middle_x": round(x_for_position(max(1, round(protein_length / 2))), 2),
        },
    }


def build_sequence_rows(sequence: str, site_rows: list[dict[str, Any]], line_width: int = 75) -> list[dict[str, Any]]:
    if not sequence:
        return []

    def build_plain_cell(amino_acid: str, pos: int) -> dict[str, Any]:
        return {
            "aa": amino_acid,
            "position": pos,
            "highlight": False,
            "ptm_text": "",
            "bg_style": "",
            "fg": "",
            "ptm_codes": [],
            "primary_ptm": "",
        }

    site_map: dict[int, dict[str, set[str]]] = {}
    for row in site_rows:
        position = row.get("Position")
        if not isinstance(position, int):
            continue
        site_map.setdefault(position, {"ptm_types": set()})
        site_map[position]["ptm_types"].add(row.get("PTM types", ""))

    rendered_rows: list[dict[str, Any]] = []
    for start in range(0, len(sequence), line_width):
        segment = sequence[start:start + line_width]
        cells: list[dict[str, Any]] = []
        for offset, amino_acid in enumerate(segment):
            pos = start + offset + 1
            ptm_info = site_map.get(pos)
            if ptm_info:
                ptm_set = {item for item in ptm_info["ptm_types"] if is_visual_ptm_code(item)}
                ordered_codes = order_ptm_codes(ptm_set)
                if not ordered_codes:
                    cells.append(build_plain_cell(amino_acid, pos))
                    continue
                primary_ptm = ordered_codes[0] if ordered_codes else ""
                style = get_ptm_style(primary_ptm)
                cells.append(
                    {
                        "aa": amino_acid,
                        "position": pos,
                        "highlight": True,
                        "ptm_text": ", ".join(ordered_codes),
                        "bg_style": build_multi_ptm_background(ordered_codes),
                        "fg": style["fg"] if len(ordered_codes) == 1 else "#ffffff",
                        "ptm_codes": ordered_codes,
                        "primary_ptm": primary_ptm,
                    }
                )
            else:
                cells.append(build_plain_cell(amino_acid, pos))

        rendered_rows.append(
            {
                "start": start + 1,
                "end": start + len(segment),
                "cells": cells,
                "is_full": len(segment) == line_width,
            }
        )
    return rendered_rows


def build_result_rows_and_meta(matched_rows: list[dict[str, str]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    def to_float(value: Any) -> float | None:
        if value is None:
            return None
        text = str(value).strip()
        if not text:
            return None
        try:
            return float(text)
        except (TypeError, ValueError):
            return None

    def mean(values: list[float]) -> float | None:
        if not values:
            return None
        return sum(values) / len(values)

    def safe_pow2(value: float | None) -> float | None:
        if value is None:
            return None
        return 2 ** value

    result_rows: list[dict[str, Any]] = []
    for idx, row in enumerate(matched_rows, start=1):
        try:
            position = int(row.get("position", ""))
        except (TypeError, ValueError):
            continue

        w4_reps_log2_raw = [to_float(row.get("w4m_1")), to_float(row.get("w4m_2")), to_float(row.get("w4m_3"))]
        w12_reps_log2_raw = [to_float(row.get("w12m_1")), to_float(row.get("w12m_2")), to_float(row.get("w12m_3"))]
        w4_reps_log2 = [v for v in w4_reps_log2_raw if v is not None]
        w12_reps_log2 = [v for v in w12_reps_log2_raw if v is not None]
        w4_reps_linear = [safe_pow2(v) for v in w4_reps_log2]
        w12_reps_linear = [safe_pow2(v) for v in w12_reps_log2]

        w4_mean_linear = mean([v for v in w4_reps_linear if v is not None])
        w12_mean_linear = mean([v for v in w12_reps_linear if v is not None])

        ratio = to_float(row.get("w12m_w4m_ratio"))
        if ratio is None and w4_mean_linear and w12_mean_linear:
            ratio = w12_mean_linear / w4_mean_linear

        result_rows.append(
            {
                "Site id": f"site-{idx}",
                "Position": position,
                "Position display": str(position) if position > 0 else "-",
                "Amino acid": row.get("amino_acid", ""),
                "Amino acid display": (row.get("amino_acid", "") or "").strip() if position > 0 else "-",
                "PTM types": row.get("ptm_types", ""),
                "Subcellular localization": (row.get("subcellular_localization") or "").strip(),
                "W4 reps log2": w4_reps_log2,
                "W12 reps log2": w12_reps_log2,
                "W4 mean linear": w4_mean_linear,
                "W12 mean linear": w12_mean_linear,
                "W12/W4 ratio": ratio,
                "W12/W4 P value": to_float(row.get("w12m_w4m_p_value")),
                "W12/W4 Type": (row.get("w12m_w4m_type") or "").strip(),
            }
        )

    if not result_rows:
        return [], {}

    first = matched_rows[0]
    result_meta = {
        "gene_name": first.get("gene_name", ""),
        "organism": first.get("organism", ""),
        "protein_accession": first.get("protein_accession", ""),
        "transcript_id": first.get("transcript_id", ""),
        "gene_names": first.get("gene_names", ""),
        "protein_description": first.get("protein_description", ""),
        "subcellular_localization": first.get("subcellular_localization", ""),
        "total_sites": len(result_rows),
    }
    return result_rows, result_meta


def build_site_chart_data(result_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    chart_data: list[dict[str, Any]] = []
    for row in result_rows:
        chart_data.append(
            {
                "site_id": row.get("Site id", ""),
                "position": row.get("Position"),
                "amino_acid": row.get("Amino acid", ""),
                "ptm_type": row.get("PTM types", ""),
                "w4_reps_log2": row.get("W4 reps log2", []),
                "w12_reps_log2": row.get("W12 reps log2", []),
                "w4_mean_linear": row.get("W4 mean linear"),
                "w12_mean_linear": row.get("W12 mean linear"),
                "ratio": row.get("W12/W4 ratio"),
                "subcellular_localization": row.get("Subcellular localization", ""),
                "p_value": row.get("W12/W4 P value"),
                "ratio_type": row.get("W12/W4 Type", ""),
            }
        )
    return chart_data


def _safe_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _mean(values: list[float]) -> float | None:
    return sum(values) / len(values) if values else None


@lru_cache(maxsize=4)
def load_expression_dataset(species: str) -> dict[str, dict[str, Any]]:
    species_key = normalize_species_key(species)
    config = EXPRESSION_DATASETS.get(species_key)
    if not config:
        return {}

    path = Path(config["path"])
    if not path.exists():
        return {}

    try:
        from openpyxl import load_workbook
    except Exception:
        return {}

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[str(config["sheet"])]
        rows = worksheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows)]
        index = {header: idx for idx, header in enumerate(headers)}
        if "Protein accession" not in index:
            return {}

        dataset: dict[str, dict[str, Any]] = {}
        for row in rows:
            accession_idx = index["Protein accession"]
            accession = str(row[accession_idx] or "").strip() if accession_idx < len(row) else ""
            if not accession:
                continue

            timepoints: list[dict[str, Any]] = []
            for timepoint in EXPRESSION_TIMEPOINTS:
                mean_idx = index.get(timepoint)
                mean_value = _safe_float(row[mean_idx]) if mean_idx is not None and mean_idx < len(row) else None
                reps: list[float] = []
                for replicate_idx in range(1, 4):
                    col_idx = index.get(f"{timepoint}_{replicate_idx}")
                    value = _safe_float(row[col_idx]) if col_idx is not None and col_idx < len(row) else None
                    if value is not None:
                        reps.append(value)

                if mean_value is None:
                    mean_value = _mean(reps)
                sd_value = statistics.stdev(reps) if len(reps) > 1 else None
                timepoints.append(
                    {
                        "label": timepoint,
                        "mean": mean_value,
                        "sd": sd_value,
                        "replicates": reps,
                        "replicate_count": len(reps),
                    }
                )

            dataset[accession] = {
                "protein_accession": accession,
                "protein_description": str(row[index["Protein description"]] or "").strip()
                if "Protein description" in index and index["Protein description"] < len(row)
                else "",
                "gene_name": str(row[index["Gene name"]] or "").strip()
                if "Gene name" in index and index["Gene name"] < len(row)
                else "",
                "dataset_label": str(config["label"]),
                "dataset_display": str(config["display"]),
                "timepoints": timepoints,
            }
    finally:
        workbook.close()

    return dataset


def build_expression_profile(accession: str, species: str) -> dict[str, Any]:
    clean_accession = str(accession or "").strip()
    species_key = normalize_species_key(species)
    config = EXPRESSION_DATASETS.get(species_key)
    profile: dict[str, Any] = {
        "available": False,
        "protein_accession": clean_accession,
        "dataset_label": str(config["label"]) if config else "",
        "dataset_display": str(config["display"]) if config else "",
        "timepoints": [],
        "axis_max": 0,
    }
    if not clean_accession:
        return profile

    record = load_expression_dataset(species_key).get(clean_accession)
    if not record:
        return profile

    raw_timepoints = list(record.get("timepoints", []))
    max_value = 0.0
    for item in raw_timepoints:
        mean_value = item.get("mean")
        sd_value = item.get("sd")
        if isinstance(mean_value, (int, float)) and math.isfinite(float(mean_value)):
            max_value = max(max_value, float(mean_value) + (float(sd_value) if isinstance(sd_value, (int, float)) else 0.0))

    axis_max = build_chart_axis_max(max_value)
    timepoints: list[dict[str, Any]] = []
    plot_scale = axis_max if axis_max > 0 else 1.0
    for item in raw_timepoints:
        mean_value = item.get("mean")
        sd_value = item.get("sd")
        mean_number = float(mean_value) if isinstance(mean_value, (int, float)) and math.isfinite(float(mean_value)) else None
        sd_number = float(sd_value) if isinstance(sd_value, (int, float)) and math.isfinite(float(sd_value)) else None
        bar_percent = (mean_number / plot_scale * 100) if mean_number is not None else 0
        err_percent = (sd_number / plot_scale * 100) if sd_number is not None else 0
        error_bottom_percent = max(0, bar_percent - err_percent) if mean_number is not None else 0
        timepoints.append(
            {
                **item,
                "mean_display": format_chart_value(mean_number),
                "sd_display": format_chart_value(sd_number),
                "bar_percent": max(2, min(100, bar_percent)) if mean_number is not None else 0,
                "error_percent": max(0, min(100, err_percent)),
                "error_bottom_percent": min(100, error_bottom_percent),
                "error_top_percent": max(0, min(100, bar_percent + err_percent)) if mean_number is not None else 0,
            }
        )

    return {
        **profile,
        **record,
        "available": any(item.get("mean") is not None for item in timepoints),
        "axis_max": axis_max,
        "axis_labels": [format_chart_value(axis_max * ratio) for ratio in (1, 0.75, 0.5, 0.25, 0)],
        "timepoints": timepoints,
    }


def build_chart_axis_max(max_value: float | None) -> float:
    if max_value is None or not math.isfinite(max_value) or max_value <= 0:
        return 1.0
    padded = max_value * 1.1
    exponent = math.floor(math.log10(padded))
    fraction = padded / (10 ** exponent)
    if fraction <= 1:
        nice_fraction = 1
    elif fraction <= 2:
        nice_fraction = 2
    elif fraction <= 5:
        nice_fraction = 5
    else:
        nice_fraction = 10
    return nice_fraction * (10 ** exponent)


def format_chart_value(value: float | None) -> str:
    if value is None or not math.isfinite(value):
        return "-"
    if abs(value) >= 1000 or (abs(value) > 0 and abs(value) < 0.001):
        return f"{value:.2e}"
    if abs(value) >= 10:
        return f"{value:.1f}".rstrip("0").rstrip(".")
    return f"{value:.4f}".rstrip("0").rstrip(".")


def _expression_na_value(value: float | None) -> str:
    return "NA" if value is None else format_chart_value(value)


def _expression_overview_value(value: float | None) -> str:
    if value is None or not math.isfinite(value):
        return "NA"
    return f"{value:.2f}"


def _expression_timepoint_values(timepoints: list[dict[str, Any]]) -> dict[str, str]:
    values = {timepoint: "NA" for timepoint in EXPRESSION_TIMEPOINTS}
    for point in timepoints:
        label = str(point.get("label", ""))
        if label in values:
            values[label] = _expression_overview_value(point.get("mean"))
    return values


def _expression_replicate_rows(timepoints: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for point in timepoints:
        replicates = list(point.get("replicates", []))
        rows.append(
            {
                "label": point.get("label", ""),
                "mean": _expression_na_value(point.get("mean")),
                "sd": _expression_na_value(point.get("sd")),
                "replicate_count": len(replicates),
                "replicates": [_expression_na_value(value) for value in replicates],
            }
        )
    return rows


def _expression_search_text(record: dict[str, Any]) -> str:
    return " ".join(
        str(record.get(key, ""))
        for key in ("protein_accession", "gene_name", "protein_description")
        if str(record.get(key, "")).strip()
    )


def _expression_identifier_candidates(record: dict[str, Any]) -> list[str]:
    description = str(record.get("protein_description", "") or record.get("description", "")).strip()
    description_ids = re.findall(r"\bGN=([^\s;|,/]+)", description)
    description_ids.extend(
        f"Glyma.{matched.group(1).upper()}"
        for matched in re.finditer(r"\bGlyma[._](\d+G\d+)\b", description, flags=re.IGNORECASE)
    )
    return collect_identifier_tokens(
        record.get("protein_accession", ""),
        record.get("gene_name", ""),
        *description_ids,
    )


def _expression_phytozome_id(record: dict[str, Any]) -> str:
    description = str(record.get("protein_description", "") or record.get("description", "")).strip()
    return _format_plaza_gene_id(
        str(record.get("transcript_id", "")),
        str(record.get("gene_names", "")),
        str(record.get("gene_name", "")) or description,
    )


def _expression_record_matches_token(record: dict[str, Any], token: str) -> bool:
    clean_token = token.strip()
    if not clean_token:
        return False
    lowered_token = clean_token.lower()
    normalized_token = normalize_identifier_token(clean_token)
    for candidate in _expression_identifier_candidates(record):
        if lowered_token == candidate.lower():
            return True
        if normalized_token and normalized_token == normalize_identifier_token(candidate):
            return True
    return False


def _expression_record_by_identifier(records: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    index: dict[str, dict[str, Any]] = {}
    for record in records:
        for candidate in _expression_identifier_candidates(record):
            normalized = normalize_identifier_token(candidate)
            if normalized:
                index.setdefault(normalized, record)
    return index


def build_proteome_overview_row(
    record: dict[str, Any],
    species: str,
    *,
    query: str = "",
    has_expression: bool = True,
) -> dict[str, Any]:
    accession = str(record.get("protein_accession", "") or record.get("gene_name", "") or query).strip()
    gene_name = str(record.get("gene_name", "")).strip()
    description = str(record.get("protein_description", "") or record.get("description", "")).strip()
    timepoints = list(record.get("timepoints", [])) if has_expression else []
    return {
        "query": query,
        "display_id": accession or gene_name or query,
        "protein_accession": accession,
        "uniprot_ids": accession,
        "gene_name": gene_name,
        "plaza_gene_id": _expression_phytozome_id(record),
        "description": description,
        "organism": str(get_species_option(species).get("display", species)),
        "dataset_label": str(record.get("dataset_label", "")),
        "dataset_display": str(record.get("dataset_display", "")),
        "timepoint_values": _expression_timepoint_values(timepoints),
        "has_expression": has_expression,
        "status": "Matched" if has_expression else "NA",
        "select_key": accession or gene_name or query,
    }


def search_expression_records(query: str, species: str, limit: int = 200) -> list[dict[str, Any]]:
    clean_query = query.strip()
    if not clean_query:
        return []

    records = list(load_expression_dataset(species).values())
    tokens = split_identifier_query(clean_query) or [clean_query]
    record_index = _expression_record_by_identifier(records)
    matched_rows: list[dict[str, Any]] = []
    fuzzy_matches: list[dict[str, Any]] = []
    seen: set[str] = set()

    def emit_expression_record(record: dict[str, Any], token: str = "") -> None:
        accession = str(record.get("protein_accession", "")).strip()
        key = normalize_identifier_token(accession) or accession.lower()
        if not key or key in seen:
            return
        matched_rows.append(build_proteome_overview_row(record, species, query=token or clean_query, has_expression=True))
        seen.add(key)

    def emit_ptm_overview(row: dict[str, Any], token: str) -> None:
        accession = str(row.get("protein_accession", "")).strip()
        gene_name = str(row.get("gene_name", "")).strip()
        normalized_candidates = [
            normalize_identifier_token(accession),
            normalize_identifier_token(gene_name),
            normalize_identifier_token(row.get("display_id", "")),
            normalize_identifier_token(row.get("query", "")),
        ]
        expression_record = next(
            (
                record_index[candidate]
                for candidate in normalized_candidates
                if candidate and candidate in record_index
            ),
            None,
        )
        if expression_record:
            emit_expression_record(expression_record, token)
            return

        key = normalize_identifier_token(accession or gene_name or row.get("display_id", "") or token)
        if not key or key in seen:
            return
        fallback_record = {
            "protein_accession": accession or str(row.get("display_id", "")).strip() or token,
            "gene_name": gene_name,
            "protein_description": str(row.get("description") or row.get("protein_description") or "").strip(),
            "dataset_label": "",
            "dataset_display": "",
            "timepoints": [],
        }
        overview = build_proteome_overview_row(fallback_record, species, query=token, has_expression=False)
        overview.update(
            {
                "display_id": row.get("display_id") or overview.get("display_id"),
                "description": row.get("description") or row.get("protein_description") or overview.get("description"),
                "plaza_gene_id": row.get("plaza_gene_id") or overview.get("plaza_gene_id"),
                "uniprot_ids": row.get("uniprot_ids") or row.get("protein_accession") or overview.get("uniprot_ids"),
                "select_key": row.get("select_key") or overview.get("select_key"),
                "status": "NA",
            }
        )
        matched_rows.append(overview)
        seen.add(key)

    for token in tokens:
        token_had_direct_match = False
        for record in records:
            if _expression_record_matches_token(record, token):
                emit_expression_record(record, token)
                token_had_direct_match = True

        if token_had_direct_match:
            continue

        try:
            for row in build_protein_summary_rows([token], species):
                if row.get("has_result"):
                    emit_ptm_overview(row, token)
        except Exception:
            pass

        if len(matched_rows) >= limit:
            return matched_rows[:limit]

    if matched_rows:
        return matched_rows[:limit]

    normalized_query = normalize_identifier_token(clean_query)
    if len(normalized_query) < 3:
        return []

    for record in records:
        accession = str(record.get("protein_accession", "")).strip()
        search_text = _expression_search_text(record).lower()
        normalized_text = normalize_identifier_token(search_text)
        if normalized_query in normalized_text and accession not in seen:
            fuzzy_matches.append(record)
            seen.add(accession)
        if len(fuzzy_matches) >= limit:
            break

    fuzzy_matches.sort(
        key=lambda item: (
            str(item.get("gene_name", "")),
            str(item.get("protein_accession", "")),
        )
    )
    return [
        build_proteome_overview_row(record, species, query=clean_query, has_expression=True)
        for record in fuzzy_matches[:limit]
    ]


def build_proteome_blast_rows(blast_rows: list[dict[str, Any]], species: str) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for row in blast_rows:
        accession = str(row.get("protein_accession") or row.get("gene_name") or row.get("display_id") or "").strip()
        profile = build_expression_profile(accession, species)
        overview_record = profile if profile.get("available") else {
            "protein_accession": accession,
            "gene_name": str(row.get("gene_name", "")).strip(),
            "protein_description": str(row.get("description") or row.get("protein_description") or "").strip(),
        }
        overview = build_proteome_overview_row(
            overview_record,
            species,
            query=str(row.get("display_id", "")),
            has_expression=bool(profile.get("available")),
        )
        overview.update(
            {
                "display_id": row.get("display_id") or overview.get("display_id"),
                "description": row.get("description") or overview.get("description"),
                "plaza_gene_id": row.get("plaza_gene_id") or overview.get("plaza_gene_id"),
                "uniprot_ids": row.get("uniprot_ids") or row.get("protein_accession") or overview.get("uniprot_ids"),
                "e_value_display": row.get("e_value_display", "-"),
                "bit_score_display": row.get("bit_score_display", "-"),
                "identity_display": row.get("identity_display", "-"),
                "query_cover_display": row.get("query_cover_display", "-"),
                "gaps": row.get("gaps", "-"),
            }
        )
        output.append(overview)
    return output


def build_proteome_detail(identifier: str, species: str) -> dict[str, Any]:
    clean_identifier = identifier.strip()
    profile = build_expression_profile(clean_identifier, species)
    if not profile.get("available"):
        matches = search_expression_records(clean_identifier, species, limit=1)
        if matches:
            profile = build_expression_profile(str(matches[0].get("protein_accession", "")), species)

    accession = str(profile.get("protein_accession") or clean_identifier).strip()
    gene_name = str(profile.get("gene_name", "")).strip()
    description = str(profile.get("protein_description", "")).strip()
    sequence = fetch_protein_sequence_by_identifiers([accession, gene_name, clean_identifier])
    return {
        "profile": profile,
        "meta": {
            "protein_accession": accession,
            "gene_name": gene_name,
            "description": description,
            "organism": str(get_species_option(species).get("display", species)),
            "sequence_length": len(sequence),
        },
        "replicate_rows": _expression_replicate_rows(list(profile.get("timepoints", []))),
    }


def _format_csv_cell(value: Any) -> str:
    text = "" if value is None else str(value)
    if any(ch in text for ch in ['"', ",", "\r", "\n"]):
        return '"' + text.replace('"', '""') + '"'
    return text


def _csv_response(filename: str, headers: list[str], rows: list[list[Any]]) -> Response:
    output = io.StringIO()
    output.write(",".join(_format_csv_cell(header) for header in headers) + "\r\n")
    for row in rows:
        output.write(",".join(_format_csv_cell(cell) for cell in row) + "\r\n")
    return Response(
        content="\ufeff" + output.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def build_protein_summary_rows(identifier_tokens: list[str], species: str = DEFAULT_SPECIES) -> list[dict[str, Any]]:
    summary_rows: list[dict[str, Any]] = []
    emitted: set[tuple[str, str, str, str]] = set()

    for token in identifier_tokens:
        matched_rows = fetch_sites_for_identifier(token, species)
        if not matched_rows:
            summary_rows.append(build_protein_overview_row(token, []))
            continue

        grouped: dict[tuple[str, str, str, str, str, str], list[dict[str, str]]] = {}
        for row in matched_rows:
            key = (
                str(row.get("gene_name", "")).strip(),
                str(row.get("protein_accession", "")).strip(),
                str(row.get("organism", "")).strip(),
                str(row.get("protein_description", "")).strip(),
                str(row.get("transcript_id", "")).strip(),
                str(row.get("gene_names", "")).strip(),
            )
            grouped.setdefault(key, []).append(row)

        for key, rows in grouped.items():
            emit_key = (token.lower(), key[0].lower(), key[1].lower(), key[2].lower())
            if emit_key in emitted:
                continue
            emitted.add(emit_key)
            summary_rows.append(build_protein_overview_row(token, rows))

    return summary_rows


def build_site_download_rows(identifier: str, species: str = DEFAULT_SPECIES) -> tuple[list[list[Any]], str]:
    matched_rows = fetch_sites_for_identifier(identifier, species)
    matched_rows = keep_primary_target_rows(matched_rows)
    result_rows, result_meta = build_result_rows_and_meta(matched_rows)
    if not result_rows:
        return [], identifier

    export_rows: list[list[Any]] = []
    for row in result_rows:
        export_rows.append(
            [
                result_meta.get("gene_name", ""),
                result_meta.get("protein_accession", ""),
                result_meta.get("organism", ""),
                result_meta.get("protein_description", ""),
                row.get("Position display", ""),
                row.get("Amino acid display", ""),
                row.get("PTM types", ""),
                row.get("Subcellular localization", ""),
                row.get("W4 mean linear", ""),
                row.get("W12 mean linear", ""),
                row.get("W12/W4 ratio", ""),
                row.get("W12/W4 P value", ""),
                row.get("W12/W4 Type", ""),
            ]
        )
    filename_key = str(result_meta.get("gene_name") or result_meta.get("protein_accession") or identifier)
    return export_rows, filename_key


def _safe_download_name(value: str, fallback: str = "protein") -> str:
    cleaned = re.sub(r'[\\/:*?"<>|\s]+', "_", value.strip())
    cleaned = cleaned.strip("._")
    return cleaned or fallback


AUTH_ROLE_ADMIN = "admin"
AUTH_ROLE_TESTER = "tester"
_AUTH_TABLE_READY = False
_AUTH_TABLE_LOCK = threading.Lock()


def _base64_urlsafe(data: bytes) -> str:
    return base64.urlsafe_b64encode(data).decode("ascii").rstrip("=")


def _base64_urlsafe_decode(value: str) -> bytes:
    padding = "=" * (-len(value) % 4)
    return base64.urlsafe_b64decode((value + padding).encode("ascii"))


def _password_hash(password: str) -> str:
    salt = secrets.token_bytes(16)
    iterations = 260000
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, iterations)
    return f"pbkdf2_sha256${iterations}${_base64_urlsafe(salt)}${_base64_urlsafe(digest)}"


def _verify_password(password: str, stored_hash: str) -> bool:
    try:
        algorithm, iterations_raw, salt_raw, digest_raw = stored_hash.split("$", 3)
        if algorithm != "pbkdf2_sha256":
            return False
        iterations = int(iterations_raw)
        salt = _base64_urlsafe_decode(salt_raw)
        expected = _base64_urlsafe_decode(digest_raw)
    except Exception:
        return False

    actual = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, iterations)
    return hmac.compare_digest(actual, expected)


def _db_bool(value: Any) -> bool:
    return str(value or "").strip().lower() in {"1", "t", "true", "yes", "y", "on"}


def _normalize_username(value: str) -> str:
    return str(value or "").strip().lower()


def _valid_username(username: str) -> bool:
    return bool(re.fullmatch(r"[a-z0-9_.@-]{3,64}", username))


def _valid_role(role: str) -> str:
    clean = str(role or "").strip().lower()
    return clean if clean in {AUTH_ROLE_ADMIN, AUTH_ROLE_TESTER} else AUTH_ROLE_TESTER


def ensure_auth_tables() -> None:
    global _AUTH_TABLE_READY
    if not PTM_AUTH_ENABLED or _AUTH_TABLE_READY:
        return

    with _AUTH_TABLE_LOCK:
        if _AUTH_TABLE_READY:
            return
        run_psql_csv_query("CREATE SCHEMA IF NOT EXISTS ptm;")
        run_psql_csv_query(
            """
            CREATE TABLE IF NOT EXISTS ptm.app_user (
              user_id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
              username TEXT NOT NULL UNIQUE,
              password_hash TEXT NOT NULL,
              role TEXT NOT NULL DEFAULT 'tester' CHECK (role IN ('admin', 'tester')),
              can_download BOOLEAN NOT NULL DEFAULT TRUE,
              is_active BOOLEAN NOT NULL DEFAULT TRUE,
              failed_login_count INT NOT NULL DEFAULT 0,
              locked_until TIMESTAMPTZ,
              last_login_at TIMESTAMPTZ,
              created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
              updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            );
            """
        )
        run_psql_csv_query("ALTER TABLE ptm.app_user ADD COLUMN IF NOT EXISTS failed_login_count INT NOT NULL DEFAULT 0;")
        run_psql_csv_query("ALTER TABLE ptm.app_user ADD COLUMN IF NOT EXISTS locked_until TIMESTAMPTZ;")
        _AUTH_TABLE_READY = True


def _auth_user_from_row(row: dict[str, Any], include_password: bool = False) -> dict[str, Any]:
    user: dict[str, Any] = {
        "user_id": int(row.get("user_id", "0") or 0),
        "username": str(row.get("username", "")).strip(),
        "role": str(row.get("role", AUTH_ROLE_TESTER)).strip() or AUTH_ROLE_TESTER,
        "can_download": _db_bool(row.get("can_download")),
        "is_active": _db_bool(row.get("is_active")),
        "failed_login_count": int(row.get("failed_login_count", "0") or 0),
        "locked_until": str(row.get("locked_until", "") or ""),
        "is_locked": _db_bool(row.get("is_locked")),
        "last_login_at": str(row.get("last_login_at", "") or ""),
        "created_at": str(row.get("created_at", "") or ""),
    }
    if include_password:
        user["password_hash"] = str(row.get("password_hash", "") or "")
    return user


def get_app_user_by_username(username: str, include_password: bool = False) -> dict[str, Any] | None:
    if not PTM_AUTH_ENABLED:
        return None
    ensure_auth_tables()
    clean_username = _normalize_username(username)
    if not clean_username:
        return None

    password_select = ", password_hash" if include_password else ""
    rows = run_psql_csv_query(
        f"""
        SELECT
          user_id,
          username,
          role,
          can_download,
          is_active,
          failed_login_count,
          locked_until,
          (locked_until IS NOT NULL AND locked_until > NOW()) AS is_locked,
          last_login_at,
          created_at{password_select}
        FROM ptm.app_user
        WHERE username = {pg_quote(clean_username)}
        LIMIT 1;
        """
    )
    return _auth_user_from_row(rows[0], include_password=include_password) if rows else None


def get_app_user_by_id(user_id: int) -> dict[str, Any] | None:
    if not PTM_AUTH_ENABLED or user_id <= 0:
        return None
    ensure_auth_tables()
    rows = run_psql_csv_query(
        f"""
        SELECT
          user_id,
          username,
          role,
          can_download,
          is_active,
          failed_login_count,
          locked_until,
          (locked_until IS NOT NULL AND locked_until > NOW()) AS is_locked,
          last_login_at,
          created_at
        FROM ptm.app_user
        WHERE user_id = {int(user_id)}
        LIMIT 1;
        """
    )
    return _auth_user_from_row(rows[0]) if rows else None


def list_app_users() -> list[dict[str, Any]]:
    ensure_auth_tables()
    rows = run_psql_csv_query(
        """
        SELECT
          user_id,
          username,
          role,
          can_download,
          is_active,
          failed_login_count,
          locked_until,
          (locked_until IS NOT NULL AND locked_until > NOW()) AS is_locked,
          last_login_at,
          created_at
        FROM ptm.app_user
        ORDER BY CASE role WHEN 'admin' THEN 0 ELSE 1 END, username;
        """
    )
    return [_auth_user_from_row(row) for row in rows]


def create_app_user(username: str, password: str, role: str, can_download: bool) -> tuple[bool, str]:
    ensure_auth_tables()
    clean_username = _normalize_username(username)
    if not _valid_username(clean_username):
        return False, "Username must be 3-64 characters: letters, numbers, dot, underscore, hyphen, or @."
    if len(password) < 8:
        return False, "Password must be at least 8 characters."
    if get_app_user_by_username(clean_username):
        return False, "Username already exists."

    clean_role = _valid_role(role)
    run_psql_csv_query(
        f"""
        INSERT INTO ptm.app_user (username, password_hash, role, can_download, is_active)
        VALUES (
          {pg_quote(clean_username)},
          {pg_quote(_password_hash(password))},
          {pg_quote(clean_role)},
          {'TRUE' if can_download else 'FALSE'},
          TRUE
        );
        """
    )
    return True, f"User {clean_username} created."


def record_login_success(user_id: int) -> None:
    run_psql_csv_query(
        f"""
        UPDATE ptm.app_user
        SET last_login_at = NOW(),
            failed_login_count = 0,
            locked_until = NULL,
            updated_at = NOW()
        WHERE user_id = {int(user_id)};
        """
    )


def record_login_failure(user_id: int) -> int:
    rows = run_psql_csv_query(
        f"""
        WITH next_values AS (
          SELECT
            user_id,
            CASE
              WHEN locked_until IS NOT NULL AND locked_until <= NOW() THEN 1
              ELSE failed_login_count + 1
            END AS next_failed_count
          FROM ptm.app_user
          WHERE user_id = {int(user_id)}
        )
        UPDATE ptm.app_user u
        SET failed_login_count = next_values.next_failed_count,
            locked_until = CASE
              WHEN next_values.next_failed_count >= {int(PTM_AUTH_MAX_FAILED_ATTEMPTS)}
              THEN NOW() + ({pg_quote(str(PTM_AUTH_LOCKOUT_MINUTES) + " minutes")})::INTERVAL
              ELSE NULL
            END,
            updated_at = NOW()
        FROM next_values
        WHERE u.user_id = next_values.user_id
        RETURNING u.failed_login_count;
        """
    )
    if not rows:
        return 0
    try:
        return int(rows[0].get("failed_login_count", "0") or 0)
    except (TypeError, ValueError):
        return 0


def update_app_user(
    *,
    target_user_id: int,
    acting_user_id: int,
    role: str,
    can_download: bool,
    is_active: bool,
    new_password: str,
) -> tuple[bool, str]:
    ensure_auth_tables()
    target = get_app_user_by_id(target_user_id)
    if not target:
        return False, "User not found."

    clean_role = _valid_role(role)
    if target_user_id == acting_user_id and (not is_active or clean_role != AUTH_ROLE_ADMIN):
        return False, "You cannot remove your own active admin access."

    password_sql = ""
    if new_password:
        if len(new_password) < 8:
            return False, "New password must be at least 8 characters."
        password_sql = f", password_hash = {pg_quote(_password_hash(new_password))}"

    run_psql_csv_query(
        f"""
        UPDATE ptm.app_user
        SET role = {pg_quote(clean_role)},
            can_download = {'TRUE' if can_download else 'FALSE'},
            is_active = {'TRUE' if is_active else 'FALSE'},
            failed_login_count = 0,
            locked_until = NULL,
            updated_at = NOW()
            {password_sql}
        WHERE user_id = {int(target_user_id)};
        """
    )
    return True, f"User {target['username']} updated."


def delete_app_user(target_user_id: int, acting_user_id: int) -> tuple[bool, str]:
    ensure_auth_tables()
    target = get_app_user_by_id(target_user_id)
    if not target:
        return False, "User not found."
    if target_user_id == acting_user_id:
        return False, "You cannot delete your own account while logged in."

    run_psql_csv_query(
        f"""
        DELETE FROM ptm.app_user
        WHERE user_id = {int(target_user_id)};
        """
    )
    return True, f"User {target['username']} deleted."


def ensure_bootstrap_admin() -> None:
    if not PTM_AUTH_ENABLED or not PTM_BOOTSTRAP_ADMIN_USERNAME or not PTM_BOOTSTRAP_ADMIN_PASSWORD:
        return
    ensure_auth_tables()
    username = _normalize_username(PTM_BOOTSTRAP_ADMIN_USERNAME)
    if not _valid_username(username) or get_app_user_by_username(username):
        return
    create_app_user(username, PTM_BOOTSTRAP_ADMIN_PASSWORD, AUTH_ROLE_ADMIN, True)


def _session_signature(payload: str) -> str:
    return hmac.new(PTM_AUTH_SECRET_KEY.encode("utf-8"), payload.encode("ascii"), hashlib.sha256).hexdigest()


def create_session_token(user_id: int) -> str:
    expires_at = int(time.time()) + (PTM_AUTH_SESSION_DAYS * 24 * 60 * 60)
    payload = _base64_urlsafe(json.dumps({"user_id": int(user_id), "exp": expires_at}, separators=(",", ":")).encode("utf-8"))
    return f"{payload}.{_session_signature(payload)}"


def parse_session_token(token: str) -> int | None:
    try:
        payload, signature = token.split(".", 1)
        if not hmac.compare_digest(signature, _session_signature(payload)):
            return None
        data = json.loads(_base64_urlsafe_decode(payload).decode("utf-8"))
        if int(data.get("exp", 0)) < int(time.time()):
            return None
        return int(data.get("user_id", 0))
    except Exception:
        return None


def get_current_user(request: Request) -> dict[str, Any] | None:
    if not PTM_AUTH_ENABLED:
        return None
    cached = getattr(request.state, "current_user", None)
    if cached is not None:
        return cached if cached else None

    token = request.cookies.get(PTM_AUTH_COOKIE_NAME, "")
    user_id = parse_session_token(token) if token else None
    try:
        user = get_app_user_by_id(user_id) if user_id else None
    except Exception:
        user = None
    if user and not user.get("is_active"):
        user = None
    request.state.current_user = user or {}
    return user


def current_user_can_download(request: Request) -> bool:
    user = get_current_user(request)
    return bool(user and user.get("can_download") and user.get("is_active"))


def current_downloads_enabled(request: Request) -> bool:
    return bool(PTM_DOWNLOADS_ENABLED or current_user_can_download(request))


def _require_downloads_enabled(request: Request) -> None:
    if not current_downloads_enabled(request):
        raise HTTPException(status_code=403, detail=DOWNLOADS_DISABLED_MESSAGE)


def _format_e_value(value: float) -> str:
    if not math.isfinite(value):
        return "-"
    if value == 0:
        return "0"
    if value < 0.001 or value >= 10000:
        return f"{value:.2e}"
    return f"{value:.4f}".rstrip("0").rstrip(".")


def _wrap_fasta_sequence(sequence: str, line_width: int = 80) -> str:
    clean = normalize_sequence_input(sequence)
    return "\n".join(clean[index:index + line_width] for index in range(0, len(clean), line_width))


def _blast_db_prefix(species: str) -> Path:
    return BLAST_DB_DIR / normalize_species_key(species) / "all_proteins"


def _blast_metadata_path(species: str) -> Path:
    return BLAST_DB_DIR / normalize_species_key(species) / "all_proteins.metadata.json"


def _blast_fasta_path(species: str) -> Path:
    return BLAST_DB_DIR / normalize_species_key(species) / "all_proteins.fasta"


def _blast_db_exists(prefix: Path) -> bool:
    return any(prefix.with_suffix(suffix).exists() for suffix in [".pin", ".psq", ".phr", ".pdb"])


def _blast_metadata_excludes_non_modification_classes(metadata: dict[str, Any]) -> bool:
    excluded = {code.upper() for code in NON_MODIFICATION_CLASSES}
    for item in metadata.values():
        if not isinstance(item, dict):
            return False
        ptm_codes = {code.strip().upper() for code in str(item.get("ptm_types", "")).split(",") if code.strip()}
        if ptm_codes.intersection(excluded):
            return False
    return True


def _check_blast_tools() -> tuple[str, str]:
    blastp_path = shutil.which("blastp")
    makeblastdb_path = shutil.which("makeblastdb")
    if not blastp_path or not makeblastdb_path:
        raise RuntimeError(
            "NCBI BLAST+ is not installed or not on PATH. Please install BLAST+ so blastp and makeblastdb are available."
        )
    return blastp_path, makeblastdb_path


def ensure_blast_database(species: str = DEFAULT_SPECIES) -> tuple[Path, dict[str, dict[str, Any]], int]:
    _, makeblastdb_path = _check_blast_tools()
    selected_species = normalize_species_key(species)
    prefix = _blast_db_prefix(selected_species)
    metadata_path = _blast_metadata_path(selected_species)
    fasta_path = _blast_fasta_path(selected_species)

    if _blast_db_exists(prefix) and metadata_path.exists():
        try:
            with metadata_path.open("r", encoding="utf-8") as fh:
                metadata = json.load(fh)
            if (
                isinstance(metadata, dict)
                and metadata
                and all("transcript_id" in item and "gene_names" in item for item in metadata.values() if isinstance(item, dict))
                and all(item.get("database_scope") == BLAST_METADATA_SCOPE for item in metadata.values() if isinstance(item, dict))
                and _blast_metadata_excludes_non_modification_classes(metadata)
            ):
                return prefix, metadata, len(metadata)
        except Exception:
            pass

    prefix.parent.mkdir(parents=True, exist_ok=True)
    candidates = fetch_blast_candidate_sequences(selected_species)
    if not candidates:
        raise RuntimeError("No protein sequences are available for this species.")

    metadata: dict[str, dict[str, Any]] = {}
    fasta_parts: list[str] = []
    for idx, candidate in enumerate(candidates, start=1):
        sequence = normalize_sequence_input(str(candidate.get("protein", "")))
        if not sequence:
            continue
        blast_id = f"ptm_{idx}"
        gene_name = str(candidate.get("gene_name", "")).strip()
        accession = str(candidate.get("protein_accession", "")).strip()
        metadata[blast_id] = {
            "gene_name": gene_name,
            "protein_accession": accession,
            "organism": str(candidate.get("organism", "")).strip(),
            "protein_description": str(candidate.get("protein_description", "")).strip(),
            "transcript_id": str(candidate.get("transcript_id", "")).strip(),
            "gene_names": str(candidate.get("gene_names", "")).strip(),
            "modification_type_count": str(candidate.get("modification_type_count", "0") or "0"),
            "total_sites": str(candidate.get("total_sites", "0") or "0"),
            "ptm_types": str(candidate.get("ptm_types", "")).strip(),
            "select_key": gene_name or accession,
            "database_scope": BLAST_METADATA_SCOPE,
        }
        fasta_parts.append(f">{blast_id}\n{_wrap_fasta_sequence(sequence)}")

    if not metadata:
        raise RuntimeError("No valid protein sequences are available for BLAST database creation.")

    fasta_path.write_text("\n".join(fasta_parts) + "\n", encoding="utf-8")
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")

    result = subprocess.run(
        [
            makeblastdb_path,
            "-in",
            str(fasta_path),
            "-dbtype",
            "prot",
            "-parse_seqids",
            "-out",
            str(prefix),
        ],
        capture_output=True,
        text=True,
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip() or result.stdout.strip() or "makeblastdb failed.")

    return prefix, metadata, len(metadata)


def build_blast_summary_rows(
    query_sequence: str,
    species: str = DEFAULT_SPECIES,
    max_hits: int = 20,
    identity_min: float = BLAST_DEFAULT_IDENTITY_MIN,
    e_value_max: float = BLAST_DEFAULT_E_VALUE_MAX,
    query_cover_min: float = BLAST_DEFAULT_QUERY_COVER_MIN,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    normalized_query = normalize_sequence_input(query_sequence)
    blastp_path, _ = _check_blast_tools()
    db_prefix, metadata, candidate_count = ensure_blast_database(species)
    hit_rows: list[dict[str, Any]] = []
    task = "blastp-short" if len(normalized_query) < 30 else "blastp"
    outfmt_fields = "6 qseqid sseqid pident length mismatch gaps qstart qend sstart send evalue bitscore qcovs"
    selected_max_hits = _clamp_int(max_hits, 5, 100)
    blast_target_limit = min(max(selected_max_hits * 5, 100), 500)

    with tempfile.TemporaryDirectory() as tmpdir:
        query_path = Path(tmpdir) / "query.fasta"
        output_path = Path(tmpdir) / "blast.tsv"
        query_path.write_text(f">Query_1\n{_wrap_fasta_sequence(normalized_query)}\n", encoding="utf-8")

        result = subprocess.run(
            [
                blastp_path,
                "-task",
                task,
                "-query",
                str(query_path),
                "-db",
                str(db_prefix),
                "-outfmt",
                outfmt_fields,
                "-evalue",
                str(e_value_max),
                "-max_target_seqs",
                str(blast_target_limit),
                "-out",
                str(output_path),
            ],
            capture_output=True,
            text=True,
        )
        if result.returncode != 0:
            raise RuntimeError(result.stderr.strip() or result.stdout.strip() or "blastp failed.")
        blast_lines = output_path.read_text(encoding="utf-8").splitlines() if output_path.exists() else []

    for line in blast_lines:
        parts = line.split("\t")
        if len(parts) < 13:
            continue
        (
            _query_id,
            subject_id,
            pident,
            alignment_length,
            _mismatch,
            gaps,
            _qstart,
            _qend,
            _sstart,
            _send,
            evalue,
            bitscore,
            qcovs,
        ) = parts[:13]
        candidate = metadata.get(subject_id, {})
        if not candidate:
            continue
        try:
            e_value_float = float(evalue)
        except ValueError:
            e_value_float = float("inf")
        try:
            bit_score_float = float(bitscore)
        except ValueError:
            bit_score_float = 0.0
        try:
            identity_float = float(pident)
        except ValueError:
            identity_float = 0.0
        try:
            query_cover_float = float(qcovs)
        except ValueError:
            query_cover_float = 0.0
        if identity_float < identity_min or e_value_float > e_value_max or query_cover_float < query_cover_min:
            continue
        gene_name = str(candidate.get("gene_name", "")).strip()
        accession = str(candidate.get("protein_accession", "")).strip()
        transcript_id = str(candidate.get("transcript_id", "")).strip()
        gene_names = str(candidate.get("gene_names", "")).strip()
        description = str(candidate.get("protein_description", "")).strip()
        total_sites = str(candidate.get("total_sites", "0") or "0")
        modification_type_count = str(candidate.get("modification_type_count", "0") or "0")
        hit_rows.append(
            {
                "query_id": "Query_1",
                "gene_name": gene_name,
                "protein_accession": accession,
                "uniprot_ids": accession,
                "organism": str(candidate.get("organism", "")).strip(),
                "protein_description": description,
                "description": description,
                "transcript_id": transcript_id,
                "gene_names": gene_names,
                "display_id": _format_overview_id("", gene_name, accession, transcript_id, gene_names),
                "plaza_gene_id": _format_plaza_gene_id(transcript_id, gene_names, gene_name),
                "modification_type_count": modification_type_count,
                "total_sites": total_sites,
                "has_ptm": total_sites not in {"", "0"},
                "ptm_types": str(candidate.get("ptm_types", "")).strip(),
                "select_key": str(candidate.get("select_key", "")).strip(),
                "bit_score": bit_score_float,
                "bit_score_display": f"{bit_score_float:.1f}",
                "e_value": e_value_float,
                "e_value_display": _format_e_value(e_value_float),
                "identity_percent": identity_float,
                "identity_display": f"{identity_float:.1f}%",
                "query_cover": query_cover_float,
                "query_cover_display": f"{query_cover_float:.1f}%",
                "gaps": gaps,
                "alignment_length": alignment_length,
            }
        )

    hit_rows.sort(
        key=lambda item: (
            float(item.get("e_value", float("inf"))),
            -float(item.get("bit_score", 0.0)),
            -float(item.get("identity_percent", 0.0)),
        )
    )
    selected_rows = hit_rows[:selected_max_hits]
    return selected_rows, {
        "candidate_count": candidate_count,
        "query_length": len(normalized_query),
        "matched_hits": len(selected_rows),
        "identity_min": identity_min,
        "e_value_max": e_value_max,
        "query_cover_min": query_cover_min,
    }


def _format_int(value: int) -> str:
    return f"{value:,}"


def _format_file_size(size_bytes: int) -> str:
    units = ["B", "KB", "MB", "GB"]
    size = float(size_bytes)
    for unit in units:
        if size < 1024 or unit == units[-1]:
            return f"{size:.1f} {unit}" if unit != "B" else f"{int(size)} {unit}"
        size /= 1024
    return f"{size_bytes} B"


def _build_home_data_tables(downloads_allowed: bool = False) -> list[dict[str, Any]]:
    tables: list[dict[str, Any]] = []
    for config in HOME_DATA_TABLES:
        filename = str(config["filename"])
        dataset = Path(filename).stem
        path = IMPORT_READY_DIR / filename
        exists = path.exists() and path.is_file()
        row_count = 0
        column_count = 0
        columns_preview = ""
        if exists:
            try:
                with path.open("r", encoding="utf-8-sig", newline="") as fh:
                    reader = csv.reader(fh)
                    headers = next(reader, [])
                    column_count = len(headers)
                    columns_preview = ", ".join(headers[:6])
                    if len(headers) > 6:
                        columns_preview += ", ..."
                    row_count = sum(1 for _ in reader)
            except Exception:
                row_count = 0
                column_count = 0
                columns_preview = "Unable to read CSV"

        tables.append(
            {
                "filename": filename,
                "dataset": dataset,
                "species": config["species"],
                "data_type": config["data_type"],
                "exists": exists,
                "rows": row_count,
                "rows_fmt": _format_int(row_count),
                "columns": column_count,
                "columns_preview": columns_preview or "-",
                "size": _format_file_size(path.stat().st_size) if exists else "-",
                "download_url": f"/downloads/import-ready/{urllib.parse.quote(filename)}" if exists and downloads_allowed else "",
                "view_url": f"/data/table?dataset={urllib.parse.quote(dataset)}" if exists and downloads_allowed else "",
            }
        )
    return tables


def _get_data_table_config(dataset: str) -> dict[str, Any] | None:
    clean_dataset = dataset.strip()
    for config in HOME_DATA_TABLES:
        filename = str(config["filename"])
        if Path(filename).stem == clean_dataset or filename == clean_dataset:
            return {**config, "dataset": Path(filename).stem}
    return None


def _read_csv_page(path: Path, page: int, page_size: int) -> tuple[list[str], list[list[str]], int]:
    headers: list[str] = []
    rows: list[list[str]] = []
    total_rows = 0
    start_index = (page - 1) * page_size
    end_index = start_index + page_size

    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        headers = next(reader, [])
        for row_index, row in enumerate(reader):
            if start_index <= row_index < end_index:
                normalized = list(row[: len(headers)])
                if len(normalized) < len(headers):
                    normalized.extend([""] * (len(headers) - len(normalized)))
                rows.append(normalized)
            total_rows += 1

    return headers, rows, total_rows


def _build_pie_gradient(items: list[dict[str, Any]]) -> str:
    positive = [item for item in items if int(item.get("site_count", 0)) > 0]
    total = sum(int(item.get("site_count", 0)) for item in positive)
    if total <= 0:
        return "conic-gradient(#dfe8c9 0 100%)"

    parts: list[str] = []
    cursor = 0.0
    for item in positive:
        share = (int(item["site_count"]) / total) * 100.0
        start = cursor
        end = cursor + share
        parts.append(f"{item['bg']} {start:.2f}% {end:.2f}%")
        cursor = end

    if cursor < 100:
        parts.append(f"#dfe8c9 {cursor:.2f}% 100%")
    return "conic-gradient(" + ", ".join(parts) + ")"


def _build_pie_items_with_percent(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    total = sum(int(item.get("site_count", 0)) for item in items if int(item.get("site_count", 0)) > 0)

    ranked = sorted(items, key=lambda x: int(x.get("site_count", 0)), reverse=True)
    output: list[dict[str, Any]] = []
    for item in ranked:
        site_count = int(item.get("site_count", 0))
        percent = (site_count / total) * 100.0 if total > 0 else 0.0
        copied = dict(item)
        copied["percent"] = percent
        copied["percent_fmt"] = f"{percent:.1f}%"
        output.append(copied)
    return output


def _hex_to_rgb(color: str) -> tuple[int, int, int]:
    value = color.strip().lstrip("#")
    if len(value) == 3:
        value = "".join(ch * 2 for ch in value)
    if len(value) != 6:
        return 79, 119, 45
    try:
        return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)
    except ValueError:
        return 79, 119, 45


def _text_color_for_bg(color: str) -> str:
    r, g, b = _hex_to_rgb(color)
    luminance = 0.299 * r + 0.587 * g + 0.114 * b
    return "#111111" if luminance >= 165 else "#ffffff"


def _polar_to_cartesian(cx: float, cy: float, r: float, angle_deg: float) -> tuple[float, float]:
    rad = math.radians(angle_deg - 90.0)
    return cx + r * math.cos(rad), cy + r * math.sin(rad)


def _build_pie_svg(items: list[dict[str, Any]], min_label_percent: float = 7.0) -> str:
    positive = [item for item in items if int(item.get("site_count", 0)) > 0]
    total = sum(int(item.get("site_count", 0)) for item in positive)
    size = 240.0
    cx = cy = size / 2
    radius = 90.0

    if total <= 0:
        return (
            f'<svg class="pie-svg" viewBox="0 0 {size:.0f} {size:.0f}" '
            'xmlns="http://www.w3.org/2000/svg" aria-hidden="true">'
            f'<circle cx="{cx:.2f}" cy="{cy:.2f}" r="{radius:.2f}" fill="#dfe8c9"/>'
            "</svg>"
        )

    parts: list[str] = [
        f'<svg class="pie-svg" viewBox="0 0 {size:.0f} {size:.0f}" xmlns="http://www.w3.org/2000/svg" aria-hidden="true">'
    ]

    cursor = 0.0
    for item in positive:
        site_count = int(item.get("site_count", 0))
        percent = (site_count / total) * 100.0
        start_angle = cursor * 3.6
        end_angle = (cursor + percent) * 3.6
        large_arc = 1 if (end_angle - start_angle) > 180 else 0

        x1, y1 = _polar_to_cartesian(cx, cy, radius, start_angle)
        x2, y2 = _polar_to_cartesian(cx, cy, radius, end_angle)
        fill = str(item.get("bg", "#4f772d"))

        path = (
            f"M {cx:.2f} {cy:.2f} "
            f"L {x1:.2f} {y1:.2f} "
            f"A {radius:.2f} {radius:.2f} 0 {large_arc} 1 {x2:.2f} {y2:.2f} Z"
        )
        parts.append(f'<path d="{path}" fill="{fill}" stroke="#ffffff" stroke-width="1"/>')

        if percent >= min_label_percent:
            mid = (start_angle + end_angle) / 2.0
            tx, ty = _polar_to_cartesian(cx, cy, radius * 0.56, mid)
            text_fill = _text_color_for_bg(fill)
            font_size = 13 if percent >= 20 else 11
            label = str(item.get("code", "")).strip()
            label_text = f"{label} {percent:.1f}%" if label else f"{percent:.1f}%"
            parts.append(
                f'<text x="{tx:.2f}" y="{ty:.2f}" text-anchor="middle" dominant-baseline="middle" '
                f'fill="{text_fill}" font-size="{font_size}" font-weight="700" '
                f'font-family="Arial, sans-serif" paint-order="stroke" stroke="rgba(0,0,0,0.2)" stroke-width="0.6">{label_text}</text>'
            )

        cursor += percent

    parts.append("</svg>")
    return "".join(parts)


def _build_species_case_sql(column: str = "organism") -> str:
    clauses = []
    for option in SPECIES_OPTIONS:
        terms = [str(term).strip().lower() for term in option.get("terms", []) if str(term).strip()]
        if not terms:
            continue
        term_clauses = [f"LOWER(COALESCE({column}, '')) LIKE '%{term}%'" for term in terms]
        clauses.append(f"WHEN {' OR '.join(term_clauses)} THEN {pg_quote(str(option['value']))}")
    return "CASE " + " ".join(clauses) + " ELSE NULL END"


def _empty_species_stats(option: dict[str, Any]) -> dict[str, Any]:
    return {
        "key": str(option["value"]),
        "label": str(option.get("label", option["value"])),
        "display": str(option.get("display", option.get("label", option["value"]))),
        "total_sites": 0,
        "total_sites_fmt": "0",
        "total_proteins": 0,
        "total_proteins_fmt": "0",
        "ptm_items": [],
        "ptm_items_nonzero": [],
        "pie_items": [],
        "pie_svg": "",
    }


def _build_ptm_items_from_counts(per_type_map: dict[str, dict[str, int]]) -> list[dict[str, Any]]:
    ptm_items: list[dict[str, Any]] = []
    for code in PTM_PRIORITY:
        style = get_ptm_style(code)
        counts = per_type_map.get(code, {"site_count": 0, "protein_count": 0})
        ptm_items.append(
            {
                "code": code,
                "label": style["label"],
                "bg": style["bg"],
                "fg": style["fg"],
                "site_count": counts["site_count"],
                "protein_count": counts["protein_count"],
                "site_count_fmt": _format_int(counts["site_count"]),
                "protein_count_fmt": _format_int(counts["protein_count"]),
            }
        )
    return ptm_items


def fetch_home_ptm_stats() -> dict[str, Any]:
    species_case_sql = _build_species_case_sql("organism")
    modification_condition = build_modification_sql_condition("ptm_types")
    total_rows = run_psql_csv_query(
        f"""
        SELECT
          species_key,
          COUNT(*) FILTER (WHERE NULLIF(BTRIM(ptm_types), '') IS NOT NULL) AS total_sites,
          COUNT(DISTINCT COALESCE(NULLIF(BTRIM(gene_name), ''), NULLIF(BTRIM(protein_accession), ''))) AS total_proteins
        FROM (
          SELECT
            {species_case_sql} AS species_key,
            ptm_types,
            gene_name,
            protein_accession
          FROM ptm.gs_ptm
        ) classified
        WHERE species_key IS NOT NULL
          AND {modification_condition}
        GROUP BY species_key;
        """
    )

    per_type_rows = run_psql_csv_query(
        f"""
        SELECT
          species_key,
          BTRIM(ptm_types) AS ptm_type,
          COUNT(*) AS site_count,
          COUNT(DISTINCT COALESCE(NULLIF(BTRIM(gene_name), ''), NULLIF(BTRIM(protein_accession), ''))) AS protein_count
        FROM (
          SELECT
            {species_case_sql} AS species_key,
            ptm_types,
            gene_name,
            protein_accession
          FROM ptm.gs_ptm
        ) classified
        WHERE species_key IS NOT NULL
          AND {modification_condition}
          AND NULLIF(BTRIM(ptm_types), '') IS NOT NULL
        GROUP BY species_key, BTRIM(ptm_types);
        """
    )

    canonical_map = {code.upper(): code for code in PTM_PRIORITY}
    species_totals: dict[str, dict[str, int]] = {}
    for row in total_rows:
        species_key = str(row.get("species_key", "")).strip()
        if not species_key:
            continue
        try:
            total_sites = int(row.get("total_sites", "0") or 0)
        except ValueError:
            total_sites = 0
        try:
            total_proteins = int(row.get("total_proteins", "0") or 0)
        except ValueError:
            total_proteins = 0
        species_totals[species_key] = {"total_sites": total_sites, "total_proteins": total_proteins}

    per_species_type_map: dict[str, dict[str, dict[str, int]]] = {}
    for row in per_type_rows:
        species_key = str(row.get("species_key", "")).strip()
        if not species_key:
            continue
        raw = str(row.get("ptm_type", "")).strip()
        if not raw:
            continue
        code = canonical_map.get(raw.upper(), raw)
        try:
            site_count = int(row.get("site_count", "0") or 0)
        except ValueError:
            site_count = 0
        try:
            protein_count = int(row.get("protein_count", "0") or 0)
        except ValueError:
            protein_count = 0
        per_species_type_map.setdefault(species_key, {})[code] = {
            "site_count": site_count,
            "protein_count": protein_count,
        }

    species_stats: list[dict[str, Any]] = []
    for option in SPECIES_OPTIONS:
        species_stat = _empty_species_stats(option)
        species_key = str(option["value"])
        totals = species_totals.get(species_key, {"total_sites": 0, "total_proteins": 0})
        ptm_items = _build_ptm_items_from_counts(per_species_type_map.get(species_key, {}))
        pie_items = [item for item in ptm_items if item["site_count"] > 0]
        species_stat.update(
            {
                "total_sites": totals["total_sites"],
                "total_sites_fmt": _format_int(totals["total_sites"]),
                "total_proteins": totals["total_proteins"],
                "total_proteins_fmt": _format_int(totals["total_proteins"]),
                "ptm_items": ptm_items,
                "ptm_items_nonzero": pie_items,
                "pie_items": _build_pie_items_with_percent(ptm_items),
                "pie_svg": _build_pie_svg(pie_items),
            }
        )
        species_stats.append(species_stat)

    total_sites = sum(int(item.get("total_sites", 0)) for item in species_stats)
    total_proteins = sum(int(item.get("total_proteins", 0)) for item in species_stats)

    return {
        "total_sites": total_sites,
        "total_sites_fmt": _format_int(total_sites),
        "total_proteins": total_proteins,
        "total_proteins_fmt": _format_int(total_proteins),
        "species_stats": species_stats,
    }


_HOME_STATS_CACHE_LOCK = threading.Lock()
_HOME_STATS_CACHE: dict[str, Any] | None = None
HOME_VISIT_METRIC_KEY = "home_visits"
_SITE_METRIC_TABLE_READY = False
_SITE_METRIC_TABLE_LOCK = threading.Lock()


def ensure_site_metric_table() -> None:
    global _SITE_METRIC_TABLE_READY
    if _SITE_METRIC_TABLE_READY:
        return

    with _SITE_METRIC_TABLE_LOCK:
        if _SITE_METRIC_TABLE_READY:
            return
        run_psql_csv_query(
            """
            CREATE TABLE IF NOT EXISTS ptm.site_metric (
              metric_key TEXT PRIMARY KEY,
              metric_value BIGINT NOT NULL DEFAULT 0 CHECK (metric_value >= 0),
              updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            );
            """
        )
        _SITE_METRIC_TABLE_READY = True


def increment_home_visit_count() -> int | None:
    ensure_site_metric_table()
    rows = run_psql_csv_query(
        f"""
        INSERT INTO ptm.site_metric AS metric (metric_key, metric_value)
        VALUES ({pg_quote(HOME_VISIT_METRIC_KEY)}, 1)
        ON CONFLICT (metric_key) DO UPDATE
        SET metric_value = metric.metric_value + 1,
            updated_at = NOW()
        RETURNING metric_value;
        """
    )
    if not rows:
        return None
    try:
        return int(rows[0].get("metric_value", "0") or 0)
    except (TypeError, ValueError):
        return None


def refresh_home_ptm_stats_cache() -> dict[str, Any]:
    global _HOME_STATS_CACHE
    stats = fetch_home_ptm_stats()
    with _HOME_STATS_CACHE_LOCK:
        _HOME_STATS_CACHE = stats
    return stats


def get_cached_home_ptm_stats() -> dict[str, Any]:
    global _HOME_STATS_CACHE
    with _HOME_STATS_CACHE_LOCK:
        if _HOME_STATS_CACHE is not None:
            return _HOME_STATS_CACHE

    stats = fetch_home_ptm_stats()

    with _HOME_STATS_CACHE_LOCK:
        if _HOME_STATS_CACHE is not None:
            return _HOME_STATS_CACHE
        _HOME_STATS_CACHE = stats
    return stats


def empty_home_ptm_stats() -> dict[str, Any]:
    species_stats = [_empty_species_stats(option) for option in SPECIES_OPTIONS]
    return {
        "total_sites": 0,
        "total_sites_fmt": "0",
        "total_proteins": 0,
        "total_proteins_fmt": "0",
        "species_stats": species_stats,
    }
SITE_DOWNLOAD_HEADERS = [
    "Gene name",
    "Protein accession",
    "Organism",
    "Description",
    "Position",
    "Amino acid",
    "PTM types",
    "Subcellular localization",
    "W4 mean linear",
    "W12 mean linear",
    "W12/W4 ratio",
    "W12/W4 P value",
    "W12/W4 Type",
]
def _identifier_batch_text(value: Any) -> str:
    if isinstance(value, (list, tuple)):
        return "\n".join(str(item).strip() for item in value if str(item).strip())
    return str(value or "")


def _protein_sites_batch_response(request: Request, identifiers: Any, species: str) -> Response:
    _require_downloads_enabled(request)
    selected_species = normalize_species_key(species)
    tokens = split_identifier_query(_identifier_batch_text(identifiers))
    if not tokens:
        raise HTTPException(status_code=400, detail="Missing identifiers")

    all_rows: list[list[Any]] = []
    for token in tokens:
        rows, _ = build_site_download_rows(token, selected_species)
        all_rows.extend(rows)
    if not all_rows:
        raise HTTPException(status_code=404, detail="No PTM sites found")
    return _csv_response("protein-sites-batch.csv", SITE_DOWNLOAD_HEADERS, all_rows)




__all__ = [name for name in globals() if not name.startswith("__")]
