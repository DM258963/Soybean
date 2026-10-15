#!/usr/bin/env python3
import argparse
import csv
import re
from collections import OrderedDict
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


SHEET_CONFIGS = [
    {
        "sheet": "Sobean",
        "prefix": "soybean",
        "organism": "Glycine max",
        "fasta_pattern": "Glycine_max*.fasta",
    },
    {
        "sheet": "Rhizobium",
        "prefix": "rhizobium",
        "organism": "Bradyrhizobium diazoefficiens",
        "fasta_pattern": "Bradyrhizobium*.fasta",
    },
]

GS_OUTPUT_HEADERS = [
    "Organism",
    "Class",
    "Type",
    "Protein accession",
    "Position",
    "Amino acid",
    "Protein description",
    "Gene name",
    "Transcript ID",
    "Gene Names",
    "W4M_1",
    "W4M_2",
    "W4M_3",
    "W12M_1",
    "W12M_2",
    "W12M_3",
    "W12M/W4M Ratio",
    "W12M/W4M P value",
    "W12M/W4M_Type",
    "Subcellular localization",
    "KEGG KO No",
    "KEGG Gene",
    "KEGG pathway",
    "Biological Process",
    "Cellular Component",
    "Molecular Function",
    "Pfam ID",
    "Domain description",
    "Site Domain",
    "KOG category",
    "KOG NO",
    "KOG description",
    "KEGG KO No 2",
    "COG category",
    "COG NO",
    "COG description",
]

HEADER_ALIASES = {
    "aa": "Amino acid",
    "aminoacid": "Amino acid",
    "class": "Class",
    "type": "Type",
    "proteinaccession": "Protein accession",
    "position": "Position",
    "proteindescription": "Protein description",
    "geneid": "Transcript ID",
    "transcriptid": "Transcript ID",
    "genenames": "Gene Names",
    "w4m1": "W4M_1",
    "w4m2": "W4M_2",
    "w4m3": "W4M_3",
    "w12m1": "W12M_1",
    "w12m2": "W12M_2",
    "w12m3": "W12M_3",
    "w12mw4mratio": "W12M/W4M Ratio",
    "w12mw4mpvalue": "W12M/W4M P value",
    "w12mw4mtype": "W12M/W4M_Type",
    "subcellularlocalization": "Subcellular localization",
    "keggkono": "KEGG KO No",
    "kegggene": "KEGG Gene",
    "keggpathway": "KEGG pathway",
    "biologicalprocess": "Biological Process",
    "cellularcomponent": "Cellular Component",
    "molecularfunction": "Molecular Function",
    "pfamid": "Pfam ID",
    "domaindescription": "Domain description",
    "sitedomain": "Site Domain",
    "kogcategory": "KOG category",
    "kogno": "KOG NO",
    "kogdescription": "KOG description",
    "keggkono1": "KEGG KO No 2",
    "keggkono2": "KEGG KO No 2",
    "cogcategory": "COG category",
    "cogno": "COG NO",
    "cogdescription": "COG description",
}


def normalize_header(value: object) -> str:
    return "".join(ch for ch in str(value or "").strip().lower() if ch.isalnum())


def canonical_header(value: object) -> str | None:
    raw = str(value or "").strip()
    lowered = raw.lower()
    if lowered == "gene name":
        return "Gene name"
    if raw == "GeneName" or lowered == "gene names":
        return "Gene Names"
    if raw == "GeneID" or lowered == "transcript id":
        return "Transcript ID"
    return HEADER_ALIASES.get(normalize_header(raw))


def clean_cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_accession(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9]", "", value.strip()).upper()


def extract_accession_from_header(header: str) -> str:
    matched = re.match(r"^[^|]*\|([^|]+)\|", header)
    if matched:
        return matched.group(1).strip()
    return header.split()[0].strip()


def parse_fasta(path: Path) -> dict[str, str]:
    accession_to_sequence: dict[str, str] = {}
    current_header = ""
    sequence_lines: list[str] = []

    def flush() -> None:
        nonlocal current_header, sequence_lines
        if not current_header:
            return
        accession = normalize_accession(extract_accession_from_header(current_header))
        sequence = "".join(ch for ch in "".join(sequence_lines).upper() if ch.isalpha())
        if accession and sequence:
            accession_to_sequence[accession] = sequence

    with path.open("r", encoding="utf-8", errors="ignore") as fh:
        for raw_line in fh:
            line = raw_line.strip()
            if not line:
                continue
            if line.startswith(">"):
                flush()
                current_header = line[1:].strip()
                sequence_lines = []
            else:
                sequence_lines.append(line)
        flush()

    return accession_to_sequence


def find_one_file(data_dir: Path, pattern: str) -> Path:
    matches = sorted(data_dir.glob(pattern))
    if not matches:
        raise FileNotFoundError(f"No file matched {data_dir / pattern}")
    return matches[0]


def find_data_workbook(data_dir: Path) -> Path:
    required_sheets = {str(config["sheet"]) for config in SHEET_CONFIGS}
    for path in sorted(data_dir.glob("*.xlsx")):
        wb = load_workbook(path, read_only=True, data_only=True)
        if required_sheets.issubset(set(wb.sheetnames)):
            return path
    raise FileNotFoundError(
        f"No xlsx workbook in {data_dir} contains sheets: {', '.join(sorted(required_sheets))}"
    )


def row_dict_from_sheet(headers: Iterable[object], row: Iterable[object]) -> dict[str, str]:
    output: dict[str, str] = {}
    for header, value in zip(headers, row):
        canonical = canonical_header(header)
        if canonical:
            cell_value = clean_cell(value)
            if cell_value or canonical not in output:
                output[canonical] = cell_value
    return output


def split_identifier_aliases(raw_value: str) -> list[str]:
    output: list[str] = []
    seen: set[str] = set()
    for token in re.split(r"[,\s;|/]+", raw_value.strip()):
        value = token.strip()
        if not value:
            continue
        lowered = value.lower()
        if lowered in seen:
            continue
        seen.add(lowered)
        output.append(value)
    return output


def collect_sequence_identifier_aliases(source: dict[str, str]) -> list[str]:
    output: list[str] = []
    seen: set[str] = set()
    for key in ("Protein accession", "Gene name", "Transcript ID", "Gene Names"):
        raw = source.get(key, "")
        values = [raw] if key != "Gene Names" else split_identifier_aliases(raw)
        for value in values:
            clean = value.strip()
            if not clean:
                continue
            lowered = clean.lower()
            if lowered in seen:
                continue
            seen.add(lowered)
            output.append(clean)
    return output


def write_sheet_outputs(
    xlsx_path: Path,
    sheet_name: str,
    organism: str,
    fasta_path: Path,
    gs_output: Path,
    protein_output: Path,
) -> tuple[int, int, int]:
    accession_to_sequence = parse_fasta(fasta_path)
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}. Available sheets: {wb.sheetnames}")

    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        raise ValueError(f"Sheet has no header: {sheet_name}")

    gs_output.parent.mkdir(parents=True, exist_ok=True)
    protein_output.parent.mkdir(parents=True, exist_ok=True)

    protein_rows: OrderedDict[str, str] = OrderedDict()
    missing_accessions: set[str] = set()
    gs_count = 0

    with gs_output.open("w", encoding="utf-8-sig", newline="") as gs_fh:
        writer = csv.DictWriter(gs_fh, fieldnames=GS_OUTPUT_HEADERS)
        writer.writeheader()

        for row in rows:
            source = row_dict_from_sheet(headers, row)
            if source.get("Class", "").strip().upper() == "DA":
                continue
            accession = source.get("Protein accession", "")
            gene_name = source.get("Gene name", "")
            if not accession or not gene_name:
                continue

            output_row = {header: source.get(header, "") for header in GS_OUTPUT_HEADERS}
            output_row["Organism"] = organism
            writer.writerow(output_row)
            gs_count += 1

            sequence = accession_to_sequence.get(normalize_accession(accession))
            if sequence:
                for identifier in collect_sequence_identifier_aliases(source):
                    protein_rows.setdefault(identifier, sequence)
            else:
                missing_accessions.add(accession)

    with protein_output.open("w", encoding="utf-8-sig", newline="") as protein_fh:
        writer = csv.writer(protein_fh)
        writer.writerow(["Gene name", "Protein"])
        writer.writerows((gene_name, sequence) for gene_name, sequence in protein_rows.items())

    return gs_count, len(protein_rows), len(missing_accessions)


def main() -> int:
    parser = argparse.ArgumentParser(description="Prepare GS and protein CSV files from data/*.xlsx and data/*.fasta.")
    parser.add_argument("--data-dir", default="data", help="Folder containing the xlsx and FASTA files.")
    parser.add_argument("--xlsx", default=None, help="Optional explicit xlsx path. Default: first data/*.xlsx file.")
    parser.add_argument("--output-dir", default="data/import_ready", help="Folder for generated CSV files.")
    args = parser.parse_args()

    data_dir = Path(args.data_dir)
    xlsx_path = Path(args.xlsx) if args.xlsx else find_data_workbook(data_dir)
    output_dir = Path(args.output_dir)

    for config in SHEET_CONFIGS:
        fasta_path = find_one_file(data_dir, str(config["fasta_pattern"]))
        prefix = str(config["prefix"])
        gs_output = output_dir / f"{prefix}_gs.csv"
        protein_output = output_dir / f"{prefix}_protein_sequence.csv"
        gs_count, protein_count, missing_count = write_sheet_outputs(
            xlsx_path=xlsx_path,
            sheet_name=str(config["sheet"]),
            organism=str(config["organism"]),
            fasta_path=fasta_path,
            gs_output=gs_output,
            protein_output=protein_output,
        )
        print(f"{config['sheet']}:")
        print(f"  GS CSV: {gs_output} ({gs_count} rows)")
        print(f"  Protein CSV: {protein_output} ({protein_count} rows)")
        print(f"  FASTA accessions not found: {missing_count}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
