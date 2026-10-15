import csv
import io
import json
import math
import os
import re
import queue
import shutil
import ssl
import statistics
import subprocess
import tempfile
import threading
import urllib.error
import urllib.parse
import urllib.request
from contextlib import contextmanager
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterator

import uvicorn
from fastapi import FastAPI, HTTPException, Query, Request
from fastapi.responses import FileResponse, HTMLResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

try:
    import psycopg
    from psycopg import OperationalError
except Exception:
    psycopg = None
    OperationalError = RuntimeError

try:
    import certifi
except Exception:
    certifi = None

BASE_DIR = Path(__file__).resolve().parent

DB_CONFIG = {
    "service": os.getenv("PTM_DB_SERVICE", "postgres"),
    "host": os.getenv("PTM_DB_HOST", "127.0.0.1"),
    "port": int(os.getenv("PTM_DB_PORT", "5432")),
    "dbname": os.getenv("PTM_DB_NAME", "ptm_prod"),
    "user": os.getenv("PTM_DB_USER", "myuser"),
    "password": os.getenv("PTM_DB_PASSWORD", ""),
    "connect_timeout": int(os.getenv("PTM_DB_CONNECT_TIMEOUT", "5")),
    "pool_min_size": int(os.getenv("PTM_DB_POOL_MIN_SIZE", "1")),
    "pool_max_size": int(os.getenv("PTM_DB_POOL_MAX_SIZE", "8")),
}

SPECIES_OPTIONS = [
    {
        "value": "glycine_max",
        "label": "Glycine max",
        "display": "Glycine max",
        "terms": ["glycine max", "soybean"],
    },
    {
        "value": "bradyrhizobium_diazoefficiens",
        "label": "Bradyrhizobium diazoefficiens",
        "display": "Bradyrhizobium diazoefficiens",
        "terms": ["bradyrhizobium diazoefficiens", "bradyrhizobium", "diazoefficiens"],
    },
]
DEFAULT_SPECIES = SPECIES_OPTIONS[0]["value"]
SPECIES_IDENTIFIER_EXAMPLES = {
    "glycine_max": [
        {"label": "Gene name", "value": "GLYMA_10G142500"},
        {"label": "Accession", "value": "I1LAY3"},
        {"label": "Short ID", "value": "10G142500"},
    ],
    "bradyrhizobium_diazoefficiens": [
        {"label": "Gene name", "value": "rpoC"},
        {"label": "Accession", "value": "Q89J75"},
        {"label": "Gene name", "value": "nifD"},
    ],
}
PTM_SEARCH_MAX_RESULTS = int(os.getenv("PTM_SEARCH_MAX_RESULTS", "500"))
PTM_SEARCH_PAGE_SIZE_DEFAULT = int(os.getenv("PTM_SEARCH_PAGE_SIZE_DEFAULT", "50"))
PTM_SEARCH_PAGE_SIZE_OPTIONS = [25, 50, 100]
DATA_TABLE_PAGE_SIZE_DEFAULT = int(os.getenv("DATA_TABLE_PAGE_SIZE_DEFAULT", "50"))
DATA_TABLE_PAGE_SIZE_OPTIONS = [25, 50, 100]
IMPORT_READY_DIR = BASE_DIR / "import_ready"
EXPRESSION_DATASETS = {
    "glycine_max": {
        "path": BASE_DIR / "data" / "Dataset S1-Soybean-DA-mean.xlsx",
        "sheet": "Soybean protein quant",
        "label": "Dataset S1-Soybean-DA-mean",
        "display": "Soybean protein quant",
    },
    "bradyrhizobium_diazoefficiens": {
        "path": BASE_DIR / "data" / "Dataset S2-Rhizobia-DA-mean.xlsx",
        "sheet": "Rhizobia protein quant",
        "label": "Dataset S2-Rhizobia-DA-mean",
        "display": "Rhizobia protein quant",
    },
}
EXPRESSION_TIMEPOINTS = ["W2", "W4", "W6", "W8", "W10", "W12"]
BLAST_DB_DIR = BASE_DIR / "blastdb_cache"
BLAST_MAX_DB_SEQUENCES = int(os.getenv("PTM_BLAST_MAX_DB_SEQUENCES", "200000"))
BLAST_METADATA_SCOPE = "all_proteins_v1"
BLAST_DEFAULT_IDENTITY_MIN = float(os.getenv("PTM_BLAST_DEFAULT_IDENTITY_MIN", "30"))
BLAST_DEFAULT_E_VALUE_MAX = float(os.getenv("PTM_BLAST_DEFAULT_E_VALUE_MAX", "1e-5"))
BLAST_DEFAULT_QUERY_COVER_MIN = float(os.getenv("PTM_BLAST_DEFAULT_QUERY_COVER_MIN", "70"))
BLAST_DEFAULT_MAX_HITS = int(os.getenv("PTM_BLAST_DEFAULT_MAX_HITS", "20"))
HOME_DATA_TABLES = [
    {
        "filename": "soybean_gs.csv",
        "species": "Glycine max",
        "data_type": "PTM sites / GS data",
    },
    {
        "filename": "soybean_protein_sequence.csv",
        "species": "Glycine max",
        "data_type": "Protein sequences",
    },
    {
        "filename": "rhizobium_gs.csv",
        "species": "Bradyrhizobium diazoefficiens",
        "data_type": "PTM sites / GS data",
    },
    {
        "filename": "rhizobium_protein_sequence.csv",
        "species": "Bradyrhizobium diazoefficiens",
        "data_type": "Protein sequences",
    },
]


def normalize_species_key(value: str) -> str:
    selected = (value or "").strip().lower()
    allowed = {str(option["value"]) for option in SPECIES_OPTIONS}
    return selected if selected in allowed else DEFAULT_SPECIES


def get_species_option(value: str) -> dict[str, Any]:
    species_key = normalize_species_key(value)
    for option in SPECIES_OPTIONS:
        if option["value"] == species_key:
            return option
    return SPECIES_OPTIONS[0]


def build_species_sql_condition(value: str, table_prefix: str = "") -> str:
    option = get_species_option(value)
    column = f"{table_prefix}.organism" if table_prefix else "organism"
    terms = [str(term).strip().lower() for term in option.get("terms", []) if str(term).strip()]
    if not terms:
        return "TRUE"
    clauses = [f"LOWER(COALESCE({column}, '')) LIKE '%{term}%'" for term in terms]
    return "(" + " OR ".join(clauses) + ")"


def build_modification_sql_condition(column: str = "ptm_types") -> str:
    excluded = ", ".join(pg_quote(code.upper()) for code in NON_MODIFICATION_CLASSES)
    return f"(NULLIF(BTRIM({column}), '') IS NOT NULL AND UPPER(BTRIM({column})) NOT IN ({excluded}))"


def _env_bool(name: str, default: bool = True) -> bool:
    raw = os.getenv(name)
    if raw is None:
        return default
    return raw.strip().lower() in {"1", "true", "yes", "y", "on"}


UNIPROT_TIMEOUT_SEC = float(os.getenv("PTM_UNIPROT_TIMEOUT_SEC", "10"))
UNIPROT_SSL_VERIFY = _env_bool("PTM_UNIPROT_SSL_VERIFY", True)
UNIPROT_CA_FILE = os.getenv("PTM_UNIPROT_CA_FILE", "").strip()


PTM_STYLE = {
    "Ub": {"tag": "ub", "label": "ubiquitylation(K)", "bg": "#f5da2f", "fg": "#c9211e"},
    "Ac": {"tag": "ac", "label": "Acetylation(K)", "bg": "#2f4ea9", "fg": "#ffe34a"},
    "ST": {"tag": "st", "label": "phosphorylation(S/T/Y)", "bg": "#e73b2f", "fg": "#ffffff"},
    "Sc": {"tag": "sc", "label": "succinylation(K)", "bg": "#258fc6", "fg": "#ffffff"},
    "Cr": {"tag": "cr", "label": "crotonylation(K)", "bg": "#7a5ec8", "fg": "#ffffff"},
    "3o": {"tag": "3o", "label": "β-hydroxybutyrylation(K)", "bg": "#c45a84", "fg": "#ffffff"},
    "LLa": {"tag": "lla", "label": "L-lactylation(K)", "bg": "#f6be6b", "fg": "#111111"},
    "DLa": {"tag": "dla", "label": "D-lactylation(K)", "bg": "#f09748", "fg": "#111111"},
    "Ig": {"tag": "ig", "label": "N-glycopeptide(N)", "bg": "#8e44ad", "fg": "#ffffff"},
    "IOg": {"tag": "iog", "label": "Intact O-GalNAc glycopeptide(S/T)", "bg": "#9a6d39", "fg": "#ffffff"},
    "Og": {"tag": "og", "label": "O-glycosylation(S/T)", "bg": "#b9770e", "fg": "#ffffff"},
    "Sn": {"tag": "sn", "label": "S-nitrosylation", "bg": "#d35454", "fg": "#ffffff"},
    "Sl": {"tag": "sl", "label": "Sulfenylation", "bg": "#1f9d8f", "fg": "#ffffff"},
    "Sg": {"tag": "sg", "label": "S-glutathionylation", "bg": "#20a08b", "fg": "#ffffff"},
    "Ss": {"tag": "ss", "label": "S-sulfhydration", "bg": "#5f7aa0", "fg": "#ffffff"},
}

PTM_PRIORITY = ["Ub", "Ac", "ST", "Sc", "Cr", "3o", "LLa", "DLa", "Ig", "IOg", "Og", "Sn", "Sl", "Sg", "Ss"]
NON_MODIFICATION_CLASSES = {"DA"}

NAV_ITEMS = [
    {"key": "home", "label": "Home", "endpoint": "/"},
    {"key": "proteome", "label": "Proteome", "endpoint": "/proteome"},
    {"key": "ptm", "label": "PTM", "endpoint": "/search/protein"},
    {"key": "analysis", "label": "Analysis", "endpoint": "/analysis"},
    {"key": "browse", "label": "Browse", "endpoint": "/browse"},
    {"key": "data", "label": "Data", "endpoint": "/data/raw"},
    {"key": "about", "label": "About", "endpoint": "/about"},
]

SEARCH_MENU = [
    {"key": "search-protein", "label": "Protein", "endpoint": "/search/protein"},
    {"key": "search-ptm", "label": "PTM", "endpoint": "/search/ptm"},
]

ANALYSIS_MENU = [
    {"key": "analysis", "label": "List Analysis", "endpoint": "/analysis"},
]

DATA_MENU = [
    {"key": "data-raw", "label": "Raw Data", "endpoint": "/data/raw"},
    {"key": "data-submit", "label": "Submit Data", "endpoint": "/data/submit"},
]

PAGES = {
    "proteome": {
        "title": "Protein Expression Profiles",
        "subtitle": "Explore protein abundance across nodule stages.",
        "description": (
            "Search proteome-level expression profiles by identifier or sequence, "
            "then open a protein-level view of stage-specific abundance."
        ),
        "actions": ["Identifier Search", "Sequence Search", "Expression Profiles"],
    },
    "search": {
        "title": "Search",
        "subtitle": "Find proteins and PTMs with flexible filters.",
        "description": (
            "Use this module to query proteins, peptide features, and PTM patterns. "
            "You can start from a protein ID, keyword, sequence motif, or curated PTM type."
        ),
        "actions": ["Protein Search", "PTM Search", "Advanced Filters"],
    },
    "search-protein": {
        "title": "Protein Search",
        "subtitle": "Search proteins and view sequence-level PTM annotations.",
        "description": (
            "Submit protein identifiers, descriptions, or sequences to display PTM-rich "
            "protein overviews and domain-level annotations."
        ),
        "actions": ["Input Protein IDs", "Upload FASTA", "Open Example Query"],
    },
    "search-ptm": {
        "title": "PTM Search",
        "subtitle": "Find species-specific target motifs and potential PTM crosstalk sites.",
        "description": (
            "Run motif-oriented PTM queries to detect co-occurrence, sequence context "
            "constraints, and experimentally supported PTM patterns across supported species."
        ),
        "actions": ["Motif Query", "Crosstalk Query", "Domain-aware Query"],
    },
    "analysis": {
        "title": "Analysis",
        "subtitle": "Analyze enrichment and PTM distribution across lists.",
        "description": (
            "Perform list analysis for soybean protein lists and compare PTM-level "
            "enrichment, depletion, and pathway-oriented signals."
        ),
        "actions": ["List Analysis", "PTM Statistics", "Export Results"],
    },
    "browse": {
        "title": "Browse",
        "subtitle": "Browse curated soybean PTM records.",
        "description": (
            "Explore soybean PTM data by experiment type, publication source, "
            "and confidence level in an easy browsing workflow."
        ),
        "actions": ["By PTM Type", "By Publication", "By Confidence"],
    },
    "data-submit": {
        "title": "Submit Data",
        "subtitle": "Contribute new soybean PTM studies to the community database.",
        "description": (
            "Upload peer-reviewed PTM datasets with metadata templates. "
            "The curation workflow helps keep records structured and reusable."
        ),
        "actions": ["Submission Guide", "Template Download", "Start Submission"],
    },
    "about": {
        "title": "About",
        "subtitle": "Project goals, design notes, and data policy.",
        "description": (
            "Soybean PTM Viewer 1.0 is a practical platform for sequence-centric PTM "
            "discovery, browsing, and systems-level analysis."
        ),
        "actions": ["Version Notes", "Citation", "Contact Team"],
    },
}


def pg_quote(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


class PgConnectionPool:
    def __init__(self, min_size: int, max_size: int) -> None:
        self.min_size = max(0, min_size)
        self.max_size = max(1, max_size)
        self._available: queue.LifoQueue[Any] = queue.LifoQueue(maxsize=self.max_size)
        self._lock = threading.Lock()
        self._created = 0

        for _ in range(min(self.min_size, self.max_size)):
            self._available.put(self._create_connection())
            self._created += 1

    def _create_connection(self) -> Any:
        if psycopg is None:
            raise RuntimeError(
                "psycopg is required for database access. Install it with: pip install 'psycopg[binary]>=3.1,<4.0'"
            )
        return psycopg.connect(
            host=DB_CONFIG["host"],
            port=DB_CONFIG["port"],
            dbname=DB_CONFIG["dbname"],
            user=DB_CONFIG["user"],
            password=DB_CONFIG["password"],
            connect_timeout=DB_CONFIG["connect_timeout"],
            autocommit=True,
        )

    def _get(self) -> Any:
        try:
            conn = self._available.get_nowait()
        except queue.Empty:
            with self._lock:
                if self._created < self.max_size:
                    self._created += 1
                    should_create = True
                else:
                    should_create = False
            if should_create:
                try:
                    return self._create_connection()
                except Exception:
                    with self._lock:
                        self._created = max(0, self._created - 1)
                    raise
            conn = self._available.get()

        if getattr(conn, "closed", True):
            with self._lock:
                self._created = max(0, self._created - 1)
            return self._get()
        return conn

    def _put(self, conn: Any) -> None:
        if getattr(conn, "closed", True):
            with self._lock:
                self._created = max(0, self._created - 1)
            return

        try:
            self._available.put_nowait(conn)
        except queue.Full:
            conn.close()
            with self._lock:
                self._created = max(0, self._created - 1)

    def discard(self, conn: Any) -> None:
        try:
            conn.close()
        finally:
            with self._lock:
                self._created = max(0, self._created - 1)

    @contextmanager
    def connection(self) -> Iterator[Any]:
        conn = self._get()
        try:
            yield conn
        finally:
            self._put(conn)

    def closeall(self) -> None:
        while True:
            try:
                conn = self._available.get_nowait()
            except queue.Empty:
                break
            try:
                conn.close()
            except Exception:
                pass
        with self._lock:
            self._created = 0


_DB_POOL: PgConnectionPool | None = None
_DB_POOL_LOCK = threading.Lock()


def get_db_pool() -> PgConnectionPool:
    global _DB_POOL
    if _DB_POOL is None:
        with _DB_POOL_LOCK:
            if _DB_POOL is None:
                _DB_POOL = PgConnectionPool(
                    min_size=int(DB_CONFIG["pool_min_size"]),
                    max_size=int(DB_CONFIG["pool_max_size"]),
                )
    return _DB_POOL


def run_psql_csv_query(sql: str) -> list[dict[str, str]]:
    pool = get_db_pool()
    last_error: Exception | None = None

    for _ in range(2):
        conn = pool._get()
        try:
            with conn.cursor() as cur:
                cur.execute(sql)
                if cur.description is None:
                    pool._put(conn)
                    return []

                headers = [column.name for column in cur.description]
                rows = cur.fetchall()
                pool._put(conn)
                return [
                    {header: "" if value is None else str(value) for header, value in zip(headers, row)}
                    for row in rows
                ]
        except OperationalError as exc:
            last_error = exc
            pool.discard(conn)
        except Exception:
            pool._put(conn)
            raise

    raise RuntimeError(str(last_error) if last_error else "Database query failed")


def normalize_sequence_input(raw_sequence: str) -> str:
    lines = []
    for line in raw_sequence.splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith(">"):
            continue
        lines.append(stripped)
    return "".join("".join(lines).split()).upper()


def normalize_identifier_token(raw_value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", raw_value.strip().lower())


def split_gene_aliases(raw_gene_name: str) -> list[str]:
    if not raw_gene_name:
        return []
    tokens = re.split(r"[,\s;|/]+", raw_gene_name.strip())
    output: list[str] = []
    seen: set[str] = set()
    for token in tokens:
        value = token.strip()
        if not value:
            continue
        lowered = value.lower()
        if lowered in seen:
            continue
        seen.add(lowered)
        output.append(value)
    return output


@lru_cache(maxsize=32)
def table_has_column(table_name: str, column_name: str) -> bool:
    if not re.fullmatch(r"[a-z_][a-z0-9_]*", table_name):
        return False
    if not re.fullmatch(r"[a-z_][a-z0-9_]*", column_name):
        return False

    rows = run_psql_csv_query(
        f"""
        SELECT 1 AS present
        FROM information_schema.columns
        WHERE table_schema = 'ptm'
          AND table_name = {pg_quote(table_name)}
          AND column_name = {pg_quote(column_name)}
        LIMIT 1;
        """
    )
    return bool(rows)


@lru_cache(maxsize=32)
def table_exists(table_name: str) -> bool:
    if not re.fullmatch(r"[a-z_][a-z0-9_]*", table_name):
        return False

    rows = run_psql_csv_query(
        f"""
        SELECT 1 AS present
        FROM information_schema.tables
        WHERE table_schema = 'ptm'
          AND table_name = {pg_quote(table_name)}
        LIMIT 1;
        """
    )
    return bool(rows)


def get_gs_identifier_columns() -> list[str]:
    columns: list[str] = []
    for column in ("transcript_id", "gene_name", "protein_accession", "gene_names"):
        if column in {"gene_name", "protein_accession"} or table_has_column("gs_ptm", column):
            columns.append(column)
    return columns


def build_gs_optional_identifier_select(table_prefix: str = "") -> str:
    prefix = f"{table_prefix}." if table_prefix else ""
    select_parts = []
    for column in ("transcript_id", "gene_names"):
        if table_has_column("gs_ptm", column):
            select_parts.append(f"{prefix}{column} AS {column}")
        else:
            select_parts.append(f"'' AS {column}")
    return ",\n            ".join(select_parts)


def build_identifier_match_condition(columns: list[str], quoted_query: str, normalized_quoted: str) -> str:
    clauses: list[str] = []
    for column in columns:
        clauses.append(f"LOWER(BTRIM({column})) = LOWER(BTRIM({quoted_query}))")
        clauses.append(
            f"""
            (
                {normalized_quoted} <> ''
                AND EXISTS (
                    SELECT 1
                    FROM regexp_split_to_table(LOWER(COALESCE({column}, '')), '[[:space:],;|/]+') AS tok
                    WHERE tok <> ''
                      AND REGEXP_REPLACE(tok, '[^a-z0-9]', '', 'g') = {normalized_quoted}
                )
            )
            """
        )
    return " OR ".join(f"({clause})" for clause in clauses)


def build_single_identifier_match_condition(column: str, quoted_query: str, normalized_quoted: str) -> str:
    return build_identifier_match_condition([column], quoted_query, normalized_quoted)


def build_identifier_fuzzy_condition(columns: list[str], normalized_quoted: str) -> str:
    return " OR ".join(
        (
            f"REGEXP_REPLACE(LOWER(COALESCE({column}, '')), '[^a-z0-9]', '', 'g') "
            f"LIKE '%' || {normalized_quoted} || '%'"
        )
        for column in columns
    )


def build_single_identifier_fuzzy_condition(column: str, normalized_quoted: str) -> str:
    return build_identifier_fuzzy_condition([column], normalized_quoted)


def collect_identifier_tokens(*values: Any) -> list[str]:
    output: list[str] = []
    seen: set[str] = set()

    for value in values:
        text = str(value or "").strip()
        if not text:
            continue
        candidates = [text]
        candidates.extend(split_gene_aliases(text))
        for matched in re.findall(r"\bGN=([^\s;|,/]+)", text):
            candidates.append(matched)

        for candidate in candidates:
            clean = candidate.strip()
            if not clean:
                continue
            lowered = clean.lower()
            if lowered in seen:
                continue
            seen.add(lowered)
            output.append(clean)

    return output


def fetch_protein_sequence_by_identifiers(identifiers: list[str]) -> str:
    tokens = collect_identifier_tokens(*identifiers)
    if not tokens:
        return ""

    for token in tokens:
        exact_value = pg_quote(token.lower())
        normalized_token = normalize_identifier_token(token)
        normalized_value = pg_quote(normalized_token) if normalized_token else "''"

        rows = run_psql_csv_query(
            f"""
            SELECT protein
            FROM ptm.protein_sequence
            WHERE LOWER(BTRIM(gene_name)) = {exact_value}
               OR (
                    {normalized_value} <> ''
                    AND REGEXP_REPLACE(LOWER(COALESCE(gene_name, '')), '[^a-z0-9]', '', 'g') = {normalized_value}
               )
            ORDER BY created_at DESC
            LIMIT 1;
            """
        )
        if rows:
            return "".join(str(rows[0].get("protein", "")).split()).upper()

    return ""


def fetch_protein_sequence_by_gene_name(gene_name: str) -> str:
    if not gene_name:
        return ""

    return fetch_protein_sequence_by_identifiers([gene_name])


def fetch_gene_name_by_exact_sequence(sequence: str) -> str:
    if not sequence:
        return ""

    rows = run_psql_csv_query(
        f"""
        SELECT gene_name
        FROM ptm.protein_sequence
        WHERE protein = {pg_quote(sequence)}
        ORDER BY created_at DESC
        LIMIT 1;
        """
    )
    if not rows:
        return ""
    return str(rows[0].get("gene_name", "")).strip()


def fetch_blast_candidate_sequences(species: str = DEFAULT_SPECIES, max_candidates: int = BLAST_MAX_DB_SEQUENCES) -> list[dict[str, str]]:
    species_condition = build_species_sql_condition(species, "gs")
    modification_condition = build_modification_sql_condition("gs.ptm_types")
    option = get_species_option(species)
    source_terms = [
        normalize_species_key(species),
        *[str(term) for term in option.get("terms", [])],
    ]
    if normalize_species_key(species) == "glycine_max":
        source_terms.extend(["glycine", "soybean"])
    elif normalize_species_key(species) == "bradyrhizobium_diazoefficiens":
        source_terms.extend(["bradyrhizobium", "diazoefficiens", "rhizobium"])
    source_terms = sorted({term.strip().lower().replace(" ", "_") for term in source_terms if term.strip()})
    source_clauses = []
    for term in source_terms:
        source_clauses.append(f"LOWER(COALESCE(ib.source_name, '')) LIKE '%{term}%'")
        source_clauses.append(f"LOWER(COALESCE(ib.notes, '')) LIKE '%{term}%'")
    source_condition = " OR ".join(source_clauses) if source_clauses else "FALSE"
    transcript_summary_select = (
        "MIN(NULLIF(BTRIM(transcript_id), '')) AS transcript_id,"
        if table_has_column("gs_ptm", "transcript_id")
        else "'' AS transcript_id,"
    )
    gene_names_summary_select = (
        "MIN(NULLIF(BTRIM(gene_names), '')) AS gene_names,"
        if table_has_column("gs_ptm", "gene_names")
        else "'' AS gene_names,"
    )
    return run_psql_csv_query(
        f"""
        WITH gs_summary AS (
            SELECT
                BTRIM(gene_name) AS gene_name,
                MIN(NULLIF(BTRIM(protein_accession), '')) AS protein_accession,
                MIN(NULLIF(BTRIM(organism), '')) AS organism,
                MIN(NULLIF(BTRIM(protein_description), '')) AS protein_description,
                {transcript_summary_select}
                {gene_names_summary_select}
                COUNT(*) FILTER (WHERE position > 0 AND NULLIF(BTRIM(ptm_types), '') IS NOT NULL) AS total_sites,
                COUNT(DISTINCT NULLIF(BTRIM(ptm_types), '')) FILTER (WHERE position > 0) AS modification_type_count,
                STRING_AGG(DISTINCT NULLIF(BTRIM(ptm_types), ''), ', ' ORDER BY NULLIF(BTRIM(ptm_types), '')) AS ptm_types
            FROM ptm.gs_ptm gs
            WHERE {species_condition}
              AND {modification_condition}
              AND NULLIF(BTRIM(gene_name), '') IS NOT NULL
            GROUP BY BTRIM(gene_name)
        ),
        selected_batches AS (
            SELECT DISTINCT ib.batch_id
            FROM ptm.import_batch ib
            WHERE ib.source_type = 'protein_sequence'
              AND ({source_condition})
        ),
        selected_batch_state AS (
            SELECT COUNT(*) AS selected_count
            FROM selected_batches
        )
        SELECT
            ps.gene_name,
            ps.protein,
            COALESCE(gs_summary.protein_accession, '') AS protein_accession,
            COALESCE(gs_summary.organism, {pg_quote(str(option.get("display", species)))}) AS organism,
            COALESCE(gs_summary.protein_description, '') AS protein_description,
            COALESCE(gs_summary.transcript_id, '') AS transcript_id,
            COALESCE(gs_summary.gene_names, '') AS gene_names,
            COALESCE(gs_summary.total_sites, 0) AS total_sites,
            COALESCE(gs_summary.modification_type_count, 0) AS modification_type_count,
            COALESCE(gs_summary.ptm_types, '') AS ptm_types
        FROM ptm.protein_sequence ps
        LEFT JOIN gs_summary
          ON LOWER(BTRIM(ps.gene_name)) = LOWER(BTRIM(gs_summary.gene_name))
        LEFT JOIN ptm.import_batch ib
          ON ps.source_batch_id = ib.batch_id
        CROSS JOIN selected_batch_state sbs
        WHERE NULLIF(BTRIM(ps.protein), '') IS NOT NULL
          AND (
            sbs.selected_count = 0
            OR ps.source_batch_id IN (SELECT batch_id FROM selected_batches)
          )
        ORDER BY COALESCE(gs_summary.total_sites, 0) DESC, ps.gene_name
        LIMIT {int(max_candidates)};
        """
    )


def keep_primary_target_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    if not rows:
        return []

    first = rows[0]
    primary_gene = (first.get("gene_name") or "").strip()
    primary_accession = (first.get("protein_accession") or "").strip()
    primary_organism = (first.get("organism") or "").strip()

    kept: list[dict[str, str]] = []
    for row in rows:
        row_gene = (row.get("gene_name") or "").strip()
        row_accession = (row.get("protein_accession") or "").strip()
        row_organism = (row.get("organism") or "").strip()
        if (
            row_gene == primary_gene
            and row_accession == primary_accession
            and row_organism == primary_organism
        ):
            kept.append(row)
    return kept if kept else rows


def split_identifier_query(raw_value: str) -> list[str]:
    tokens = re.split(r"[,;\s]+", raw_value.strip())
    output: list[str] = []
    seen: set[str] = set()
    for token in tokens:
        value = token.strip()
        if not value:
            continue
        lowered = value.lower()
        if lowered in seen:
            continue
        seen.add(lowered)
        output.append(value)
    return output


def _format_plaza_gene_id(transcript_id: str, gene_names: str = "", gene_name: str = "") -> str:
    candidates = collect_identifier_tokens(transcript_id, gene_names, gene_name)
    for candidate in candidates:
        matched = re.search(r"(glyma)[._](\d+g\d+)", candidate, flags=re.IGNORECASE)
        if matched:
            return f"Glyma.{matched.group(2).upper()}"
    return ""


def _format_overview_id(query: str, gene_name: str, accession: str, transcript_id: str, gene_names: str) -> str:
    clean_query = query.strip()
    if clean_query:
        return re.sub(r"[.]p$", "", clean_query, flags=re.IGNORECASE)
    for candidate in collect_identifier_tokens(transcript_id, gene_name, accession, gene_names):
        if candidate:
            return re.sub(r"[.]p$", "", candidate, flags=re.IGNORECASE)
    return ""


def build_protein_overview_row(query: str, rows: list[dict[str, str]]) -> dict[str, Any]:
    if not rows:
        return {
            "query": query,
            "display_id": query,
            "gene_name": "",
            "protein_accession": "",
            "uniprot_ids": "",
            "organism": "",
            "protein_description": "",
            "description": "",
            "transcript_id": "",
            "gene_names": "",
            "plaza_gene_id": "",
            "modification_type_count": 0,
            "total_sites": 0,
            "ptm_summary": "",
            "select_key": query,
            "status": "No result",
            "has_result": False,
        }

    first = rows[0]
    gene_name = str(first.get("gene_name", "")).strip()
    accession = str(first.get("protein_accession", "")).strip()
    organism = str(first.get("organism", "")).strip()
    description = str(first.get("protein_description", "")).strip()
    transcript_id = str(first.get("transcript_id", "")).strip()
    gene_names = str(first.get("gene_names", "")).strip()

    ptm_counts: dict[str, int] = {}
    valid_site_count = 0
    for row in rows:
        try:
            if int(row.get("position", "") or 0) <= 0:
                continue
        except (TypeError, ValueError):
            continue
        ptm_type = str(row.get("ptm_types", "")).strip()
        if not is_visual_ptm_code(ptm_type):
            continue
        valid_site_count += 1
        ptm_counts[ptm_type] = ptm_counts.get(ptm_type, 0) + 1

    ordered_ptm_types = order_ptm_codes(set(ptm_counts))
    ordered_ptm_types.extend(sorted(code for code in ptm_counts if code not in ordered_ptm_types))
    ptm_summary = "; ".join(f"{code}: {ptm_counts[code]}" for code in ordered_ptm_types)

    return {
        "query": query,
        "display_id": _format_overview_id(query, gene_name, accession, transcript_id, gene_names),
        "gene_name": gene_name,
        "protein_accession": accession,
        "uniprot_ids": accession,
        "organism": organism,
        "protein_description": description,
        "description": description,
        "transcript_id": transcript_id,
        "gene_names": gene_names,
        "plaza_gene_id": _format_plaza_gene_id(transcript_id, gene_names, gene_name),
        "modification_type_count": len(ptm_counts),
        "total_sites": valid_site_count,
        "ptm_summary": ptm_summary,
        "select_key": gene_name or accession or transcript_id or query,
        "status": "Matched" if valid_site_count else "Matched, no valid PTM site",
        "has_result": valid_site_count > 0,
    }


def build_identifier_match_rows(rows: list[dict[str, str]], query: str) -> list[dict[str, Any]]:
    query_tokens = {token.lower() for token in collect_identifier_tokens(query)}
    grouped: dict[tuple[str, str, str], dict[str, Any]] = {}

    for row in rows:
        gene_name = str(row.get("gene_name", "")).strip()
        accession = str(row.get("protein_accession", "")).strip()
        organism = str(row.get("organism", "")).strip()
        description = str(row.get("protein_description", "")).strip()
        transcript_id = str(row.get("transcript_id", "")).strip()
        gene_names = str(row.get("gene_names", "")).strip()
        key = (gene_name, accession, organism)

        bucket = grouped.get(key)
        if bucket is None:
            row_tokens = {
                token.lower()
                for token in collect_identifier_tokens(gene_name, accession, transcript_id, gene_names)
            }
            is_exact = bool(query_tokens & row_tokens)
            select_key = gene_name if gene_name else accession
            bucket = {
                "gene_name": gene_name,
                "protein_accession": accession,
                "transcript_id": transcript_id,
                "gene_names": gene_names,
                "organism": organism,
                "protein_description": description,
                "site_count": 0,
                "is_exact": is_exact,
                "select_key": select_key,
            }
            grouped[key] = bucket

        try:
            if int(row.get("position", "")) > 0:
                bucket["site_count"] += 1
        except (TypeError, ValueError):
            pass

    output = list(grouped.values())
    output.sort(
        key=lambda item: (
            0 if item.get("is_exact") else 1,
            -(int(item.get("site_count", 0))),
            str(item.get("gene_name", "")),
            str(item.get("protein_accession", "")),
        )
    )
    return output


def fetch_sites_from_identifier_lookup(gene_query: str, species: str = DEFAULT_SPECIES) -> list[dict[str, str]]:
    normalized_query = normalize_identifier_token(gene_query)
    if not normalized_query or not table_exists("identifier_lookup"):
        return []

    normalized_quoted = pg_quote(normalized_query)
    lookup_species_condition = build_species_sql_condition(species, "il")
    modification_condition = build_modification_sql_condition("gs.ptm_types")
    optional_identifier_select = build_gs_optional_identifier_select("gs")

    return run_psql_csv_query(
        f"""
        WITH lookup_matches AS (
            SELECT
                il.organism,
                il.gene_name,
                il.protein_accession,
                il.identifier_priority
            FROM ptm.identifier_lookup il
            WHERE il.identifier_norm = {normalized_quoted}
              AND {lookup_species_condition}
        ),
        best_priority AS (
            SELECT MIN(identifier_priority) AS identifier_priority
            FROM lookup_matches
        ),
        targets AS (
            SELECT DISTINCT
                lm.organism,
                lm.gene_name,
                lm.protein_accession,
                lm.identifier_priority
            FROM lookup_matches lm
            INNER JOIN best_priority bp
              ON lm.identifier_priority = bp.identifier_priority
        )
        SELECT
            gs.organism,
            gs.ptm_types,
            gs.protein_accession,
            gs.position,
            gs.amino_acid,
            gs.protein_description,
            gs.gene_name,
            {optional_identifier_select},
            gs.subcellular_localization,
            gs.w4m_1,
            gs.w4m_2,
            gs.w4m_3,
            gs.w12m_1,
            gs.w12m_2,
            gs.w12m_3,
            gs.w12m_w4m_ratio,
            gs.w12m_w4m_p_value,
            gs.w12m_w4m_type
        FROM ptm.gs_ptm gs
        INNER JOIN targets t
          ON gs.organism = t.organism
         AND gs.gene_name = t.gene_name
         AND gs.protein_accession = t.protein_accession
        WHERE {modification_condition}
        ORDER BY t.identifier_priority, gs.position ASC;
        """
    )


def fetch_sites_for_identifier(gene_query: str, species: str = DEFAULT_SPECIES) -> list[dict[str, str]]:
    lookup_rows = fetch_sites_from_identifier_lookup(gene_query, species)
    if lookup_rows:
        return lookup_rows

    quoted_query = pg_quote(gene_query)
    normalized_exact_query = normalize_identifier_token(gene_query)
    normalized_exact_quoted = pg_quote(normalized_exact_query) if normalized_exact_query else "''"
    species_condition = build_species_sql_condition(species)
    modification_condition = build_modification_sql_condition("ptm_types")
    identifier_columns = get_gs_identifier_columns()
    optional_identifier_select = build_gs_optional_identifier_select()

    def build_priority_case(conditions_by_column: dict[str, str]) -> str:
        clauses = [
            f"WHEN ({conditions_by_column[column]}) THEN {idx}"
            for idx, column in enumerate(identifier_columns)
            if column in conditions_by_column
        ]
        return "CASE " + " ".join(clauses) + " ELSE 99 END"

    def run_identifier_query(condition: str, priority_case: str) -> list[dict[str, str]]:
        return run_psql_csv_query(
            f"""
            SELECT
                organism,
                ptm_types,
                protein_accession,
                position,
                amino_acid,
                protein_description,
                gene_name,
                {optional_identifier_select},
                subcellular_localization,
                w4m_1,
                w4m_2,
                w4m_3,
                w12m_1,
                w12m_2,
                w12m_3,
                w12m_w4m_ratio,
                w12m_w4m_p_value,
                w12m_w4m_type
            FROM ptm.gs_ptm
            WHERE (
                {condition}
            )
              AND {species_condition}
              AND {modification_condition}
            ORDER BY {priority_case}, position ASC;
            """
        )

    exact_conditions_by_column = {
        column: build_single_identifier_match_condition(
            column,
            quoted_query,
            normalized_exact_quoted,
        )
        for column in identifier_columns
    }
    exact_condition = " OR ".join(f"({condition})" for condition in exact_conditions_by_column.values())
    exact_rows = run_identifier_query(exact_condition, build_priority_case(exact_conditions_by_column))
    if exact_rows:
        return exact_rows

    normalized_query = re.sub(r"[^a-z0-9]", "", gene_query.strip().lower())
    if not normalized_query:
        return []
    normalized_query = re.sub(r"^glyma", "", normalized_query)
    if len(normalized_query) < 4:
        return []

    normalized_quoted = pg_quote(normalized_query)
    fuzzy_conditions_by_column = {
        column: build_single_identifier_fuzzy_condition(column, normalized_quoted)
        for column in identifier_columns
    }
    fuzzy_condition = " OR ".join(f"({condition})" for condition in fuzzy_conditions_by_column.values())
    return run_identifier_query(fuzzy_condition, build_priority_case(fuzzy_conditions_by_column))


def fetch_sites_by_gene_name(gene_name: str, species: str = DEFAULT_SPECIES) -> list[dict[str, str]]:
    if not gene_name:
        return []
    normalized_exact_query = normalize_identifier_token(gene_name)
    normalized_exact_quoted = pg_quote(normalized_exact_query) if normalized_exact_query else "''"
    species_condition = build_species_sql_condition(species)
    modification_condition = build_modification_sql_condition("ptm_types")
    optional_identifier_select = build_gs_optional_identifier_select()

    return run_psql_csv_query(
        f"""
        SELECT
            organism,
            ptm_types,
            protein_accession,
            position,
            amino_acid,
            protein_description,
            gene_name,
            {optional_identifier_select},
            subcellular_localization,
            w4m_1,
            w4m_2,
            w4m_3,
            w12m_1,
            w12m_2,
            w12m_3,
            w12m_w4m_ratio,
            w12m_w4m_p_value,
            w12m_w4m_type
        FROM ptm.gs_ptm
        WHERE (
            LOWER(BTRIM(gene_name)) = LOWER(BTRIM({pg_quote(gene_name)}))
            OR (
                {normalized_exact_quoted} <> ''
                AND EXISTS (
                    SELECT 1
                    FROM regexp_split_to_table(LOWER(COALESCE(gene_name, '')), '[[:space:],;|/]+') AS tok
                    WHERE tok <> ''
                      AND REGEXP_REPLACE(tok, '[^a-z0-9]', '', 'g') = {normalized_exact_quoted}
                )
            )
        )
          AND {species_condition}
          AND {modification_condition}
        ORDER BY position ASC;
        """
    )


def _query_checkbox_checked(query_params: Any, key: str, default: bool) -> bool:
    if key in query_params:
        return True
    return default


def _clamp_int(value: int, minimum: int, maximum: int) -> int:
    return max(minimum, min(maximum, value))


def _clamp_float(value: float, minimum: float, maximum: float) -> float:
    if not math.isfinite(value):
        return minimum
    return max(minimum, min(maximum, value))


def _build_ptm_alias_map() -> dict[str, str]:
    aliases: dict[str, str] = {}
    for code, style in PTM_STYLE.items():
        aliases[code.lower()] = code
        tag = str(style.get("tag", "")).strip().lower()
        if tag:
            aliases[tag] = code
    return aliases


PTM_ALIAS_MAP = _build_ptm_alias_map()


def _parse_ptm_condition(text: str) -> tuple[dict[str, Any] | None, str]:
    raw = text.strip()
    if not raw:
        return None, "PTM 条件不能为空。"

    lowered = raw.lower()
    if lowered == "xx":
        return {"mode": "wildcard", "codes": [], "raw": raw}, ""

    if "&" in raw and "|" in raw:
        return None, "同一个 PTM 条件里不能同时混用 '&' 和 '|'。"

    sep = "&" if "&" in raw else "|" if "|" in raw else ""
    parts = [raw] if not sep else raw.split(sep)
    codes: list[str] = []
    for part in parts:
        token = part.strip().lower()
        if not token:
            return None, "PTM 条件格式错误，请检查括号里的内容。"
        mapped = PTM_ALIAS_MAP.get(token)
        if not mapped:
            return None, f"未识别的 PTM 类型: {part.strip()}。可用示例: Ub, Ac, ST, Sc, Cr, Og。"
        if mapped not in codes:
            codes.append(mapped)

    mode = "and" if sep == "&" else "or"
    return {"mode": mode, "codes": codes, "raw": raw}, ""


def parse_ptm_search_expression(expression: str) -> tuple[dict[str, Any] | None, str]:
    compact = "".join(expression.split())
    if not compact:
        return None, "请输入 Search Expression。"

    tokens: list[dict[str, Any]] = []
    constraints: list[dict[str, Any]] = []
    idx = 0
    while idx < len(compact):
        current = compact[idx]
        aa_display = ""
        aa_regex = ""

        if current == "[":
            end_idx = compact.find("]", idx + 1)
            if end_idx == -1:
                return None, "字符集合缺少右括号 ']'。"
            inner = compact[idx + 1:end_idx].upper()
            if not inner:
                return None, "字符集合 [] 不能为空。"
            if any(not ch.isalpha() for ch in inner):
                return None, "字符集合 [] 内仅支持氨基酸字母。"
            aa_display = f"[{inner}]"
            aa_regex = aa_display
            idx = end_idx + 1
        elif current in {"X", "x", "."}:
            aa_display = "X"
            aa_regex = "[A-Z]"
            idx += 1
        elif current.isalpha():
            aa_display = current.upper()
            aa_regex = aa_display
            idx += 1
        elif current in {"?", "+", "*", "{", "}", "|", "^", "$"}:
            return None, "当前版本暂不支持 ?, +, *, {}, |, ^, $ 等正则操作符。"
        else:
            return None, f"不支持的字符: {current}"

        token_offset = len(tokens)
        if idx < len(compact) and compact[idx] == "(":
            close_idx = compact.find(")", idx + 1)
            if close_idx == -1:
                return None, "PTM 条件缺少右括号 ')'。"
            parsed_ptm, error = _parse_ptm_condition(compact[idx + 1:close_idx])
            if error:
                return None, error
            constraints.append(
                {
                    "offset": token_offset,
                    "aa_display": aa_display,
                    "mode": parsed_ptm["mode"],
                    "codes": parsed_ptm["codes"],
                    "raw": parsed_ptm["raw"],
                }
            )
            idx = close_idx + 1

        tokens.append({"display": aa_display, "regex": aa_regex})

    if not tokens:
        return None, "Search Expression 不能为空。"
    if not constraints:
        return None, "每个表达式至少包含一个 PTM 条件，例如 M(Ub)。"

    motif_regex = "".join(token["regex"] for token in tokens)
    required_codes = sorted({code for rule in constraints for code in rule.get("codes", [])})
    return {
        "expression": compact,
        "tokens": tokens,
        "constraints": constraints,
        "motif_regex": motif_regex,
        "motif_length": len(tokens),
        "required_codes": required_codes,
    }, ""


def fetch_ptm_rows_with_sequence(required_codes: list[str], species: str = DEFAULT_SPECIES) -> list[dict[str, str]]:
    type_filter = ""
    if required_codes:
        quoted = ", ".join(pg_quote(code) for code in required_codes)
        type_filter = f"AND BTRIM(ptm.gs_ptm.ptm_types) IN ({quoted})"
    species_condition = build_species_sql_condition(species, "ptm.gs_ptm")
    modification_condition = build_modification_sql_condition("ptm.gs_ptm.ptm_types")

    return run_psql_csv_query(
        f"""
        SELECT
            ptm.gs_ptm.organism,
            ptm.gs_ptm.ptm_types,
            ptm.gs_ptm.protein_accession,
            ptm.gs_ptm.position,
            ptm.gs_ptm.amino_acid,
            ptm.gs_ptm.protein_description,
            ptm.gs_ptm.gene_name,
            ptm.gs_ptm.subcellular_localization,
            ptm.protein_sequence.protein AS protein_sequence
        FROM ptm.gs_ptm
        LEFT JOIN ptm.protein_sequence
          ON LOWER(BTRIM(ptm.gs_ptm.gene_name)) = LOWER(BTRIM(ptm.protein_sequence.gene_name))
        WHERE {species_condition}
          AND {modification_condition}
          AND NULLIF(BTRIM(ptm.gs_ptm.gene_name), '') IS NOT NULL
          {type_filter}
        ORDER BY ptm.gs_ptm.gene_name, ptm.gs_ptm.protein_accession, ptm.gs_ptm.position ASC;
        """
    )


def _gene_representative_key(gene_name: str) -> str:
    clean = gene_name.strip()
    matched = re.match(r"^(.*)\.\d+$", clean)
    if matched:
        return matched.group(1)
    return clean


def _format_ptm_hit_label(mode: str, matched_codes: list[str]) -> str:
    if not matched_codes:
        return "PTM"
    if mode == "and":
        return "&".join(matched_codes)
    return "|".join(matched_codes)


PTM_EXAMPLE_RESIDUE_BY_CODE = {
    "Ub": "K",
    "Ac": "K",
    "ST": "[STY]",
    "Sc": "K",
    "Cr": "K",
    "3o": "K",
    "LLa": "K",
    "DLa": "K",
    "lg": "N",
    "IOg": "[ST]",
    "Og": "[ST]",
    "Sn": "C",
    "Sl": "C",
    "Sg": "C",
    "Ss": "C",
}


def _example_residue_for_ptm(code: str, label: str) -> str:
    preset = PTM_EXAMPLE_RESIDUE_BY_CODE.get(code)
    if preset:
        return preset

    matched = re.search(r"\(([A-Za-z/]+)\)\s*$", label.strip())
    if not matched:
        return "X"

    token = matched.group(1).upper()
    letters = [ch for ch in token if "A" <= ch <= "Z"]
    if not letters:
        return "X"

    dedup: list[str] = []
    seen: set[str] = set()
    for ch in letters:
        if ch in seen:
            continue
        seen.add(ch)
        dedup.append(ch)

    if len(dedup) == 1:
        return dedup[0]
    return "[" + "".join(dedup) + "]"


def _build_ptm_reference_rows() -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for code in PTM_PRIORITY:
        style = get_ptm_style(code)
        label = style.get("label", code)
        tag = str(style.get("tag", "")).strip().lower() or code.lower()
        residue = _example_residue_for_ptm(code, str(label))
        rows.append(
            {
                "code": code,
                "tag": style.get("tag", ""),
                "label": label,
                "example": f"{residue}({tag})",
            }
        )
    return rows


def run_ptm_expression_search(
    parsed: dict[str, Any],
    *,
    sequence_length: int,
    representative_only: bool,
    peptides_unique: bool,
    max_results: int,
    species: str = DEFAULT_SPECIES,
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    rows = fetch_ptm_rows_with_sequence(parsed.get("required_codes", []), species)
    if not rows:
        return [], {"matched_proteins": 0, "matched_sites": 0, "truncated": 0}

    grouped: dict[tuple[str, str], dict[str, Any]] = {}
    for row in rows:
        gene_name = str(row.get("gene_name", "")).strip()
        if not gene_name:
            continue
        accession = str(row.get("protein_accession", "")).strip()
        key = (gene_name, accession)
        bucket = grouped.get(key)
        if bucket is None:
            raw_sequence = "".join(str(row.get("protein_sequence", "")).split()).upper()
            bucket = {
                "gene_name": gene_name,
                "protein_accession": accession,
                "organism": str(row.get("organism", "")).strip(),
                "protein_description": str(row.get("protein_description", "")).strip(),
                "subcellular_localization": str(row.get("subcellular_localization", "")).strip(),
                "sequence": raw_sequence,
                "ptm_map": {},
            }
            grouped[key] = bucket

        position_text = str(row.get("position", "")).strip()
        ptm_code = str(row.get("ptm_types", "")).strip()
        if not position_text or not ptm_code:
            continue
        try:
            position = int(position_text)
        except ValueError:
            continue
        bucket["ptm_map"].setdefault(position, set()).add(ptm_code)

    if not grouped:
        return [], {"matched_proteins": 0, "matched_sites": 0, "truncated": 0}

    motif_regex = str(parsed["motif_regex"])
    compiled = re.compile(f"(?=({motif_regex}))")
    motif_len = int(parsed["motif_length"])
    constraints = parsed["constraints"]

    matches: list[dict[str, Any]] = []
    for bucket in grouped.values():
        sequence = str(bucket.get("sequence", ""))
        if not sequence:
            continue

        for match in compiled.finditer(sequence):
            start_idx = match.start()
            motif_start = start_idx + 1
            motif_end = motif_start + motif_len - 1
            matched_sites: list[dict[str, Any]] = []
            passed = True

            for rule in constraints:
                site_pos = motif_start + int(rule["offset"])
                site_idx = site_pos - 1
                aa_here = sequence[site_idx]
                site_types = bucket["ptm_map"].get(site_pos, set())
                if not site_types:
                    passed = False
                    break

                if rule["mode"] == "wildcard":
                    matched_codes = order_ptm_codes(site_types)
                elif rule["mode"] == "and":
                    needed = set(rule.get("codes", []))
                    if not needed.issubset(site_types):
                        passed = False
                        break
                    matched_codes = [code for code in rule["codes"] if code in site_types]
                else:
                    matched_codes = [code for code in order_ptm_codes(site_types) if code in set(rule.get("codes", []))]
                    if not matched_codes:
                        passed = False
                        break

                matched_sites.append(
                    {
                        "position": site_pos,
                        "amino_acid": aa_here,
                        "matched_codes": matched_codes,
                        "mode": rule["mode"],
                        "left_context": sequence[max(0, site_idx - sequence_length):site_idx],
                        "right_context": sequence[site_idx + 1:min(len(sequence), site_idx + 1 + sequence_length)],
                    }
                )

            if not passed:
                continue

            left_start = max(0, start_idx - sequence_length)
            right_end = min(len(sequence), start_idx + motif_len + sequence_length)
            left_context = sequence[left_start:start_idx]
            motif_context = sequence[start_idx:start_idx + motif_len]
            right_context = sequence[start_idx + motif_len:right_end]

            site_labels = [
                f"{site['position']}:{site['amino_acid']}({ _format_ptm_hit_label(site['mode'], site['matched_codes']) })"
                for site in matched_sites
            ]
            matches.append(
                {
                    "gene_name": bucket["gene_name"],
                    "protein_accession": bucket["protein_accession"],
                    "organism": bucket["organism"],
                    "protein_description": bucket["protein_description"],
                    "subcellular_localization": bucket["subcellular_localization"],
                    "motif_start": motif_start,
                    "motif_end": motif_end,
                    "motif_sequence": motif_context,
                    "site_hits_text": "; ".join(site_labels),
                    "matched_site_count": len(matched_sites),
                    "matched_sites": matched_sites,
                    "left_context": left_context,
                    "motif_context": motif_context,
                    "right_context": right_context,
                }
            )

    if representative_only:
        kept: list[dict[str, Any]] = []
        seen_keys: set[str] = set()
        for row in sorted(matches, key=lambda item: (item["gene_name"], item["protein_accession"], item["motif_start"])):
            key = _gene_representative_key(str(row["gene_name"]))
            if key in seen_keys:
                continue
            seen_keys.add(key)
            kept.append(row)
        matches = kept

    if peptides_unique:
        motif_gene_map: dict[str, set[str]] = {}
        for row in matches:
            motif = str(row.get("motif_sequence", ""))
            motif_gene_map.setdefault(motif, set()).add(_gene_representative_key(str(row.get("gene_name", ""))))
        matches = [row for row in matches if len(motif_gene_map.get(str(row.get("motif_sequence", "")), set())) == 1]

    matches = sorted(matches, key=lambda item: (item["gene_name"], item["protein_accession"], item["motif_start"]))
    truncated = 1 if len(matches) > max_results else 0
    if truncated:
        matches = matches[:max_results]

    protein_keys = {(row["gene_name"], row["protein_accession"]) for row in matches}
    total_sites = sum(int(row.get("matched_site_count", 0)) for row in matches)
    return matches, {"matched_proteins": len(protein_keys), "matched_sites": total_sites, "truncated": truncated}


def _species_table_label(organism: str) -> str:
    normalized = organism.strip().lower()
    if "glycine max" in normalized or "soybean" in normalized:
        return "gmx"
    if "bradyrhizobium" in normalized or "diazoefficiens" in normalized:
        return "bdi"
    return organism.strip() or "-"


def build_ptm_site_rows(matches: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    seen: set[tuple[str, str, int, int, str]] = set()

    for match in matches:
        gene_name = str(match.get("gene_name", "")).strip()
        accession = str(match.get("protein_accession", "")).strip()
        description = str(match.get("protein_description", "")).strip()
        organism = str(match.get("organism", "")).strip()
        motif_start = int(match.get("motif_start", 0) or 0)
        motif_end = int(match.get("motif_end", 0) or 0)

        for site in match.get("matched_sites", []):
            pos = int(site.get("position", 0) or 0)
            aa = str(site.get("amino_acid", "")).strip()
            left_ctx = str(site.get("left_context", ""))
            right_ctx = str(site.get("right_context", ""))
            hit_codes = [str(code) for code in site.get("matched_codes", []) if str(code).strip()]
            hit_label = _format_ptm_hit_label(str(site.get("mode", "")), hit_codes)

            dedupe_key = (gene_name, accession, pos, motif_start, "|".join(hit_codes))
            if dedupe_key in seen:
                continue
            seen.add(dedupe_key)

            rows.append(
                {
                    "gene_name": gene_name,
                    "protein_accession": accession,
                    "protein_description": description,
                    "organism": organism,
                    "species_label": _species_table_label(organism),
                    "ptm_position": pos,
                    "hit_aa": aa,
                    "hit_label": hit_label,
                    "left_context": left_ctx,
                    "right_context": right_ctx,
                    "motif_start": motif_start,
                    "motif_end": motif_end,
                }
            )

    rows.sort(key=lambda item: (item["gene_name"], item["protein_accession"], item["ptm_position"], item["motif_start"]))
    return rows


def paginate_rows(rows: list[dict[str, Any]], page: int, page_size: int) -> tuple[list[dict[str, Any]], dict[str, int]]:
    total = len(rows)
    total_pages = max(1, math.ceil(total / page_size)) if total > 0 else 1
    current = _clamp_int(page, 1, total_pages)
    start = (current - 1) * page_size
    end = start + page_size
    sliced = rows[start:end]

    return sliced, {
        "total_rows": total,
        "total_pages": total_pages,
        "current_page": current,
        "page_size": page_size,
        "start_row": 0 if total == 0 else start + 1,
        "end_row": min(end, total),
        "prev_page": current - 1 if current > 1 else 1,
        "next_page": current + 1 if current < total_pages else total_pages,
    }


def group_ptm_rows_by_protein(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: list[dict[str, Any]] = []
    lookup: dict[tuple[str, str, str, str], dict[str, Any]] = {}

    for row in rows:
        key = (
            str(row.get("gene_name", "")),
            str(row.get("protein_accession", "")),
            str(row.get("protein_description", "")),
            str(row.get("species_label", "")),
        )
        group = lookup.get(key)
        if group is None:
            group = {
                "gene_name": key[0],
                "protein_accession": key[1],
                "protein_description": key[2],
                "species_label": key[3],
                "site_rows": [],
                "rowspan": 0,
            }
            lookup[key] = group
            groups.append(group)

        group["site_rows"].append(
            {
                "ptm_position": row.get("ptm_position", ""),
                "left_context": row.get("left_context", ""),
                "hit_aa": row.get("hit_aa", ""),
                "right_context": row.get("right_context", ""),
                "hit_label": row.get("hit_label", ""),
            }
        )

    for group in groups:
        group["rowspan"] = len(group["site_rows"])

    return groups


def _uniq_preserve(items: list[str]) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for item in items:
        text = item.strip()
        if not text or text in seen:
            continue
        seen.add(text)
        out.append(text)
    return out


def _nested_value(node: Any, path: list[str], default: str = "") -> str:
    current = node
    for key in path:
        if not isinstance(current, dict):
            return default
        current = current.get(key)
    if current is None:
        return default
    return str(current).strip()


def _extract_comment_texts(comment: dict[str, Any]) -> list[str]:
    texts: list[str] = []
    for text_item in comment.get("texts", []) or []:
        value = _nested_value(text_item, ["value"])
        if value:
            texts.append(value)
    return texts


def _normalize_accession(accession: str) -> str:
    if not accession:
        return ""
    token = accession.strip().split(";")[0].split(",")[0].strip()
    return token


def _build_uniprot_ssl_context() -> tuple[ssl.SSLContext, str]:
    if not UNIPROT_SSL_VERIFY:
        return ssl._create_unverified_context(), "verify=off"

    if UNIPROT_CA_FILE:
        return ssl.create_default_context(cafile=UNIPROT_CA_FILE), f"cafile={UNIPROT_CA_FILE}"

    if certifi is not None:
        try:
            ca_path = certifi.where()
            if ca_path:
                return ssl.create_default_context(cafile=ca_path), f"certifi={ca_path}"
        except Exception:
            pass

    return ssl.create_default_context(), "system-default"


@lru_cache(maxsize=256)
def fetch_uniprot_entry(accession: str) -> tuple[dict[str, Any] | None, str]:
    normalized = _normalize_accession(accession)
    if not normalized:
        return None, "Empty accession"

    url = f"https://rest.uniprot.org/uniprotkb/{urllib.parse.quote(normalized)}.json"
    request = urllib.request.Request(
        url,
        headers={
            "Accept": "application/json",
            "User-Agent": "SoybeanPTMViewer/1.0",
        },
    )
    ssl_context, ssl_source = _build_uniprot_ssl_context()
    try:
        with urllib.request.urlopen(request, timeout=UNIPROT_TIMEOUT_SEC, context=ssl_context) as response:
            payload = response.read().decode("utf-8")
        parsed = json.loads(payload)
        if isinstance(parsed, dict):
            return parsed, ""
        return None, "Invalid JSON payload"
    except urllib.error.HTTPError as exc:
        return None, f"HTTP {exc.code} [{ssl_source}]"
    except urllib.error.URLError as exc:
        return None, f"Network error [{ssl_source}]: {exc.reason}"
    except TimeoutError:
        return None, f"Request timeout [{ssl_source}]"
    except json.JSONDecodeError:
        return None, f"Invalid JSON response [{ssl_source}]"
    except Exception as exc:
        return None, f"{type(exc).__name__} [{ssl_source}]: {exc}"


def build_uniprot_context(accession: str, result_rows: list[dict[str, Any]], result_meta: dict[str, Any]) -> dict[str, Any]:
    normalized = _normalize_accession(accession)
    entry, error_reason = fetch_uniprot_entry(normalized)
    if not entry:
        fallback_entry_url = f"https://www.uniprot.org/uniprotkb/{normalized}/entry" if normalized else ""
        fallback_af_url = f"https://alphafold.ebi.ac.uk/entry/{normalized}" if normalized else ""
        reason = f" ({error_reason})" if error_reason else ""
        return {
            "available": False,
            "query_accession": normalized,
            "primary_accession": normalized,
            "entry_url": fallback_entry_url,
            "alphafold_id": normalized,
            "alphafold_url": fallback_af_url,
            "protein_name": str(result_meta.get("protein_description", "")).strip(),
            "db_subcellular_localization": str(result_meta.get("subcellular_localization", "")).strip(),
            "message": f"UniProt data is currently unavailable for this accession{reason}.",
        }

    primary_accession = _nested_value(entry, ["primaryAccession"]) or normalized
    entry_url = f"https://www.uniprot.org/uniprotkb/{primary_accession}/entry"

    protein_name = _nested_value(entry, ["proteinDescription", "recommendedName", "fullName", "value"])
    if not protein_name:
        submission_names = (entry.get("proteinDescription", {}) or {}).get("submissionNames", []) or []
        if submission_names:
            protein_name = _nested_value(submission_names[0], ["fullName", "value"])
    if not protein_name:
        protein_name = str(result_meta.get("protein_description", "")).strip()

    gene_names: list[str] = []
    for gene in entry.get("genes", []) or []:
        main_name = _nested_value(gene, ["geneName", "value"])
        if main_name:
            gene_names.append(main_name)
        for syn in gene.get("synonyms", []) or []:
            syn_name = _nested_value(syn, ["value"])
            if syn_name:
                gene_names.append(syn_name)
    gene_names = _uniq_preserve(gene_names)

    organism_name = _nested_value(entry, ["organism", "scientificName"]) or str(result_meta.get("organism", "")).strip()

    function_texts: list[str] = []
    catalytic_activities: list[str] = []
    subcellular_locations: list[str] = []

    for comment in entry.get("comments", []) or []:
        comment_type = str(comment.get("commentType", "")).strip().upper()
        if comment_type == "FUNCTION":
            function_texts.extend(_extract_comment_texts(comment))
        elif comment_type == "CATALYTIC ACTIVITY":
            reaction_name = _nested_value(comment, ["reaction", "name"])
            ec_number = _nested_value(comment, ["reaction", "ecNumber"])
            assembled = reaction_name
            if ec_number:
                assembled = f"{assembled} (EC {ec_number})" if assembled else f"EC {ec_number}"
            if assembled:
                catalytic_activities.append(assembled)
        elif comment_type == "SUBCELLULAR LOCATION":
            for loc in comment.get("subcellularLocations", []) or []:
                location = _nested_value(loc, ["location", "value"])
                topology = _nested_value(loc, ["topology", "value"])
                orientation = _nested_value(loc, ["orientation", "value"])
                bits = [location, topology, orientation]
                merged = " | ".join([b for b in bits if b])
                if merged:
                    subcellular_locations.append(merged)
            subcellular_locations.extend(_extract_comment_texts(comment))

    function_texts = _uniq_preserve(function_texts)
    catalytic_activities = _uniq_preserve(catalytic_activities)
    subcellular_locations = _uniq_preserve(subcellular_locations)

    go_terms: list[dict[str, str]] = []
    pdb_entries: list[dict[str, str]] = []
    alphafold_id = ""

    for ref in entry.get("uniProtKBCrossReferences", []) or []:
        db_name = str(ref.get("database", "")).strip()
        db_id = str(ref.get("id", "")).strip()
        props = {str(p.get("key", "")).strip(): str(p.get("value", "")).strip() for p in ref.get("properties", []) or []}

        if db_name == "GO" and db_id:
            go_terms.append(
                {
                    "id": db_id,
                    "term": props.get("GoTerm", ""),
                    "evidence": props.get("GoEvidenceType", ""),
                }
            )
        elif db_name == "PDB" and db_id:
            pdb_entries.append(
                {
                    "id": db_id,
                    "chains": props.get("Chains", ""),
                    "method": props.get("Method", ""),
                    "resolution": props.get("Resolution", ""),
                }
            )
        elif db_name == "AlphaFoldDB" and db_id:
            alphafold_id = db_id

    if not alphafold_id:
        alphafold_id = primary_accession

    feature_types = {"DOMAIN", "REGION", "ACT_SITE", "BINDING", "MOTIF", "SITE"}
    feature_items: list[dict[str, str]] = []
    for feat in entry.get("features", []) or []:
        feature_type = str(feat.get("type", "")).strip()
        if feature_type not in feature_types:
            continue
        begin = _nested_value(feat, ["location", "start", "value"])
        end = _nested_value(feat, ["location", "end", "value"])
        if begin and end and begin != end:
            pos_text = f"{begin}-{end}"
        else:
            pos_text = begin or end
        description = str(feat.get("description", "")).strip()
        feature_items.append({"type": feature_type, "position": pos_text, "description": description})
        if len(feature_items) >= 20:
            break

    ptm_sites = []
    for row in result_rows:
        ptm_sites.append(
            {
                "position": row.get("Position"),
                "amino_acid": row.get("Amino acid", ""),
                "ptm_type": row.get("PTM types", ""),
            }
        )

    return {
        "available": True,
        "query_accession": normalized,
        "primary_accession": primary_accession,
        "entry_url": entry_url,
        "protein_name": protein_name,
        "gene_names": gene_names[:12],
        "organism_name": organism_name,
        "function_texts": function_texts[:8],
        "catalytic_activities": catalytic_activities[:8],
        "subcellular_locations": subcellular_locations[:12],
        "go_terms": go_terms[:20],
        "alphafold_id": alphafold_id,
        "alphafold_url": f"https://alphafold.ebi.ac.uk/entry/{alphafold_id}" if alphafold_id else "",
        "pdb_entries": pdb_entries[:12],
        "feature_items": feature_items,
        "ptm_sites": ptm_sites[:100],
        "db_subcellular_localization": str(result_meta.get("subcellular_localization", "")).strip(),
        "message": "",
    }


app = FastAPI(title="Soybean PTM Viewer 1.0")
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")

templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))
templates.env.globals["static_url"] = lambda request, path: request.url_for("static", path=path)
templates.env.globals["route_url"] = lambda request, route_name: request.url_for(route_name)


@app.on_event("shutdown")
def close_db_pool() -> None:
    if _DB_POOL is not None:
        _DB_POOL.closeall()


@app.on_event("startup")
def warm_home_stats_on_startup() -> None:
    try:
        refresh_home_ptm_stats_cache()
    except Exception:
        pass


def render_template(
    request: Request,
    template_name: str,
    *,
    title: str,
    active: str,
    extra_context: dict[str, Any] | None = None,
) -> HTMLResponse:
    context: dict[str, Any] = {
        "request": request,
        "title": title,
        "active": active,
        "nav_items": NAV_ITEMS,
        "search_menu": SEARCH_MENU,
        "analysis_menu": ANALYSIS_MENU,
        "data_menu": DATA_MENU,
    }
    if extra_context:
        context.update(extra_context)
    try:
        # Starlette/FastAPI newer signature
        return templates.TemplateResponse(request=request, name=template_name, context=context)
    except TypeError:
        # Backward-compatible fallback for older signature
        return templates.TemplateResponse(template_name, context=context)


def render_page(request: Request, page_key: str, active_key: str) -> HTMLResponse:
    page = PAGES[page_key]
    return render_template(
        request,
        "simple_page.html",
        title=page["title"],
        active=active_key,
        extra_context={"page": page},
    )


def order_ptm_codes(ptm_codes: set[str]) -> list[str]:
    ordered = [ptm for ptm in PTM_PRIORITY if ptm in ptm_codes]
    remaining = sorted([ptm for ptm in ptm_codes if ptm not in ordered])
    return ordered + remaining


def is_visual_ptm_code(ptm_code: str) -> bool:
    code = str(ptm_code or "").strip()
    return bool(code) and code.upper() not in NON_MODIFICATION_CLASSES


def get_ptm_style(ptm_code: str) -> dict[str, str]:
    return PTM_STYLE.get(
        ptm_code,
        {"tag": ptm_code.lower()[:3] if ptm_code else "", "label": ptm_code or "Unknown PTM", "bg": "#e73b2f", "fg": "#ffffff"},
    )


def build_multi_ptm_background(ordered_codes: list[str]) -> str:
    if not ordered_codes:
        return ""
    if len(ordered_codes) == 1:
        return get_ptm_style(ordered_codes[0])["bg"]
    return "#000000"


def build_ptm_legend_from_sequence(sequence_rows: list[dict[str, Any]]) -> list[dict[str, str]]:
    seen = set()
    has_multiple_types = False
    for line in sequence_rows:
        for cell in line.get("cells", []):
            if cell.get("highlight"):
                ptm_codes = cell.get("ptm_codes", [])
                if len(ptm_codes) > 1:
                    has_multiple_types = True
                for code in ptm_codes:
                    seen.add(code)

    legend_codes = order_ptm_codes(seen)

    legend_items = []
    for code in legend_codes:
        style = get_ptm_style(code)
        legend_items.append({"code": code, "label": style["label"], "bg": style["bg"], "fg": style["fg"]})
    if has_multiple_types:
        legend_items.append({"code": "__multiple__", "label": "Multiple types", "bg": "#000000", "fg": "#ffffff"})
    return legend_items


def build_sequence_rows(sequence: str, site_rows: list[dict[str, Any]], line_width: int = 75) -> list[dict[str, Any]]:
    if not sequence:
        return []

    def build_plain_cell(amino_acid: str, pos: int) -> dict[str, Any]:
        return {
            "aa": amino_acid,
            "position": pos,
            "highlight": False,
            "ptm_text": "",
            "bg_style": "",
            "fg": "",
            "ptm_codes": [],
            "primary_ptm": "",
        }

    site_map: dict[int, dict[str, set[str]]] = {}
    for row in site_rows:
        position = row.get("Position")
        if not isinstance(position, int):
            continue
        site_map.setdefault(position, {"ptm_types": set()})
        site_map[position]["ptm_types"].add(row.get("PTM types", ""))

    rendered_rows: list[dict[str, Any]] = []
    for start in range(0, len(sequence), line_width):
        segment = sequence[start:start + line_width]
        cells: list[dict[str, Any]] = []
        for offset, amino_acid in enumerate(segment):
            pos = start + offset + 1
            ptm_info = site_map.get(pos)
            if ptm_info:
                ptm_set = {item for item in ptm_info["ptm_types"] if is_visual_ptm_code(item)}
                ordered_codes = order_ptm_codes(ptm_set)
                if not ordered_codes:
                    cells.append(build_plain_cell(amino_acid, pos))
                    continue
                primary_ptm = ordered_codes[0] if ordered_codes else ""
                style = get_ptm_style(primary_ptm)
                cells.append(
                    {
                        "aa": amino_acid,
                        "position": pos,
                        "highlight": True,
                        "ptm_text": ", ".join(ordered_codes),
                        "bg_style": build_multi_ptm_background(ordered_codes),
                        "fg": style["fg"] if len(ordered_codes) == 1 else "#ffffff",
                        "ptm_codes": ordered_codes,
                        "primary_ptm": primary_ptm,
                    }
                )
            else:
                cells.append(build_plain_cell(amino_acid, pos))

        rendered_rows.append(
            {
                "start": start + 1,
                "end": start + len(segment),
                "cells": cells,
                "is_full": len(segment) == line_width,
            }
        )
    return rendered_rows


def build_result_rows_and_meta(matched_rows: list[dict[str, str]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    def to_float(value: Any) -> float | None:
        if value is None:
            return None
        text = str(value).strip()
        if not text:
            return None
        try:
            return float(text)
        except (TypeError, ValueError):
            return None

    def mean(values: list[float]) -> float | None:
        if not values:
            return None
        return sum(values) / len(values)

    def safe_pow2(value: float | None) -> float | None:
        if value is None:
            return None
        return 2 ** value

    result_rows: list[dict[str, Any]] = []
    for idx, row in enumerate(matched_rows, start=1):
        try:
            position = int(row.get("position", ""))
        except (TypeError, ValueError):
            continue

        w4_reps_log2_raw = [to_float(row.get("w4m_1")), to_float(row.get("w4m_2")), to_float(row.get("w4m_3"))]
        w12_reps_log2_raw = [to_float(row.get("w12m_1")), to_float(row.get("w12m_2")), to_float(row.get("w12m_3"))]
        w4_reps_log2 = [v for v in w4_reps_log2_raw if v is not None]
        w12_reps_log2 = [v for v in w12_reps_log2_raw if v is not None]
        w4_reps_linear = [safe_pow2(v) for v in w4_reps_log2]
        w12_reps_linear = [safe_pow2(v) for v in w12_reps_log2]

        w4_mean_linear = mean([v for v in w4_reps_linear if v is not None])
        w12_mean_linear = mean([v for v in w12_reps_linear if v is not None])

        ratio = to_float(row.get("w12m_w4m_ratio"))
        if ratio is None and w4_mean_linear and w12_mean_linear:
            ratio = w12_mean_linear / w4_mean_linear

        result_rows.append(
            {
                "Site id": f"site-{idx}",
                "Position": position,
                "Position display": str(position) if position > 0 else "-",
                "Amino acid": row.get("amino_acid", ""),
                "Amino acid display": (row.get("amino_acid", "") or "").strip() if position > 0 else "-",
                "PTM types": row.get("ptm_types", ""),
                "Subcellular localization": (row.get("subcellular_localization") or "").strip(),
                "W4 reps log2": w4_reps_log2,
                "W12 reps log2": w12_reps_log2,
                "W4 mean linear": w4_mean_linear,
                "W12 mean linear": w12_mean_linear,
                "W12/W4 ratio": ratio,
                "W12/W4 P value": to_float(row.get("w12m_w4m_p_value")),
                "W12/W4 Type": (row.get("w12m_w4m_type") or "").strip(),
            }
        )

    if not result_rows:
        return [], {}

    first = matched_rows[0]
    result_meta = {
        "gene_name": first.get("gene_name", ""),
        "organism": first.get("organism", ""),
        "protein_accession": first.get("protein_accession", ""),
        "transcript_id": first.get("transcript_id", ""),
        "gene_names": first.get("gene_names", ""),
        "protein_description": first.get("protein_description", ""),
        "subcellular_localization": first.get("subcellular_localization", ""),
        "total_sites": len(result_rows),
    }
    return result_rows, result_meta


def build_site_chart_data(result_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    chart_data: list[dict[str, Any]] = []
    for row in result_rows:
        chart_data.append(
            {
                "site_id": row.get("Site id", ""),
                "position": row.get("Position"),
                "amino_acid": row.get("Amino acid", ""),
                "ptm_type": row.get("PTM types", ""),
                "w4_reps_log2": row.get("W4 reps log2", []),
                "w12_reps_log2": row.get("W12 reps log2", []),
                "w4_mean_linear": row.get("W4 mean linear"),
                "w12_mean_linear": row.get("W12 mean linear"),
                "ratio": row.get("W12/W4 ratio"),
                "subcellular_localization": row.get("Subcellular localization", ""),
                "p_value": row.get("W12/W4 P value"),
                "ratio_type": row.get("W12/W4 Type", ""),
            }
        )
    return chart_data


def _safe_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _mean(values: list[float]) -> float | None:
    return sum(values) / len(values) if values else None


@lru_cache(maxsize=4)
def load_expression_dataset(species: str) -> dict[str, dict[str, Any]]:
    species_key = normalize_species_key(species)
    config = EXPRESSION_DATASETS.get(species_key)
    if not config:
        return {}

    path = Path(config["path"])
    if not path.exists():
        return {}

    try:
        from openpyxl import load_workbook
    except Exception:
        return {}

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[str(config["sheet"])]
        rows = worksheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows)]
        index = {header: idx for idx, header in enumerate(headers)}
        if "Protein accession" not in index:
            return {}

        dataset: dict[str, dict[str, Any]] = {}
        for row in rows:
            accession_idx = index["Protein accession"]
            accession = str(row[accession_idx] or "").strip() if accession_idx < len(row) else ""
            if not accession:
                continue

            timepoints: list[dict[str, Any]] = []
            for timepoint in EXPRESSION_TIMEPOINTS:
                mean_idx = index.get(timepoint)
                mean_value = _safe_float(row[mean_idx]) if mean_idx is not None and mean_idx < len(row) else None
                reps: list[float] = []
                for replicate_idx in range(1, 4):
                    col_idx = index.get(f"{timepoint}_{replicate_idx}")
                    value = _safe_float(row[col_idx]) if col_idx is not None and col_idx < len(row) else None
                    if value is not None:
                        reps.append(value)

                if mean_value is None:
                    mean_value = _mean(reps)
                sd_value = statistics.stdev(reps) if len(reps) > 1 else None
                timepoints.append(
                    {
                        "label": timepoint,
                        "mean": mean_value,
                        "sd": sd_value,
                        "replicates": reps,
                        "replicate_count": len(reps),
                    }
                )

            dataset[accession] = {
                "protein_accession": accession,
                "protein_description": str(row[index["Protein description"]] or "").strip()
                if "Protein description" in index and index["Protein description"] < len(row)
                else "",
                "gene_name": str(row[index["Gene name"]] or "").strip()
                if "Gene name" in index and index["Gene name"] < len(row)
                else "",
                "dataset_label": str(config["label"]),
                "dataset_display": str(config["display"]),
                "timepoints": timepoints,
            }
    finally:
        workbook.close()

    return dataset


def build_expression_profile(accession: str, species: str) -> dict[str, Any]:
    clean_accession = str(accession or "").strip()
    species_key = normalize_species_key(species)
    config = EXPRESSION_DATASETS.get(species_key)
    profile: dict[str, Any] = {
        "available": False,
        "protein_accession": clean_accession,
        "dataset_label": str(config["label"]) if config else "",
        "dataset_display": str(config["display"]) if config else "",
        "timepoints": [],
        "axis_max": 0,
    }
    if not clean_accession:
        return profile

    record = load_expression_dataset(species_key).get(clean_accession)
    if not record:
        return profile

    raw_timepoints = list(record.get("timepoints", []))
    max_value = 0.0
    for item in raw_timepoints:
        mean_value = item.get("mean")
        sd_value = item.get("sd")
        if isinstance(mean_value, (int, float)) and math.isfinite(float(mean_value)):
            max_value = max(max_value, float(mean_value) + (float(sd_value) if isinstance(sd_value, (int, float)) else 0.0))

    axis_max = build_chart_axis_max(max_value)
    timepoints: list[dict[str, Any]] = []
    plot_scale = axis_max if axis_max > 0 else 1.0
    for item in raw_timepoints:
        mean_value = item.get("mean")
        sd_value = item.get("sd")
        mean_number = float(mean_value) if isinstance(mean_value, (int, float)) and math.isfinite(float(mean_value)) else None
        sd_number = float(sd_value) if isinstance(sd_value, (int, float)) and math.isfinite(float(sd_value)) else None
        bar_percent = (mean_number / plot_scale * 100) if mean_number is not None else 0
        err_percent = (sd_number / plot_scale * 100) if sd_number is not None else 0
        error_bottom_percent = max(0, bar_percent - err_percent) if mean_number is not None else 0
        timepoints.append(
            {
                **item,
                "mean_display": format_chart_value(mean_number),
                "sd_display": format_chart_value(sd_number),
                "bar_percent": max(2, min(100, bar_percent)) if mean_number is not None else 0,
                "error_percent": max(0, min(100, err_percent)),
                "error_bottom_percent": min(100, error_bottom_percent),
                "error_top_percent": max(0, min(100, bar_percent + err_percent)) if mean_number is not None else 0,
            }
        )

    return {
        **profile,
        **record,
        "available": any(item.get("mean") is not None for item in timepoints),
        "axis_max": axis_max,
        "axis_labels": [format_chart_value(axis_max * ratio) for ratio in (1, 0.75, 0.5, 0.25, 0)],
        "timepoints": timepoints,
    }


def build_chart_axis_max(max_value: float | None) -> float:
    if max_value is None or not math.isfinite(max_value) or max_value <= 0:
        return 1.0
    padded = max_value * 1.1
    exponent = math.floor(math.log10(padded))
    fraction = padded / (10 ** exponent)
    if fraction <= 1:
        nice_fraction = 1
    elif fraction <= 2:
        nice_fraction = 2
    elif fraction <= 5:
        nice_fraction = 5
    else:
        nice_fraction = 10
    return nice_fraction * (10 ** exponent)


def format_chart_value(value: float | None) -> str:
    if value is None or not math.isfinite(value):
        return "-"
    if abs(value) >= 1000 or (abs(value) > 0 and abs(value) < 0.001):
        return f"{value:.2e}"
    if abs(value) >= 10:
        return f"{value:.1f}".rstrip("0").rstrip(".")
    return f"{value:.4f}".rstrip("0").rstrip(".")


def _expression_na_value(value: float | None) -> str:
    return "NA" if value is None else format_chart_value(value)


def _expression_timepoint_values(timepoints: list[dict[str, Any]]) -> dict[str, str]:
    values = {timepoint: "NA" for timepoint in EXPRESSION_TIMEPOINTS}
    for point in timepoints:
        label = str(point.get("label", ""))
        if label in values:
            values[label] = _expression_na_value(point.get("mean"))
    return values


def _expression_replicate_rows(timepoints: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for point in timepoints:
        replicates = list(point.get("replicates", []))
        rows.append(
            {
                "label": point.get("label", ""),
                "mean": _expression_na_value(point.get("mean")),
                "sd": _expression_na_value(point.get("sd")),
                "replicate_count": len(replicates),
                "replicates": [_expression_na_value(value) for value in replicates],
            }
        )
    return rows


def _expression_search_text(record: dict[str, Any]) -> str:
    return " ".join(
        str(record.get(key, ""))
        for key in ("protein_accession", "gene_name", "protein_description")
        if str(record.get(key, "")).strip()
    )


def build_proteome_overview_row(
    record: dict[str, Any],
    species: str,
    *,
    query: str = "",
    has_expression: bool = True,
) -> dict[str, Any]:
    accession = str(record.get("protein_accession", "") or record.get("gene_name", "") or query).strip()
    gene_name = str(record.get("gene_name", "")).strip()
    description = str(record.get("protein_description", "") or record.get("description", "")).strip()
    timepoints = list(record.get("timepoints", [])) if has_expression else []
    return {
        "query": query,
        "display_id": accession or gene_name or query,
        "protein_accession": accession,
        "gene_name": gene_name,
        "description": description,
        "organism": str(get_species_option(species).get("display", species)),
        "dataset_label": str(record.get("dataset_label", "")),
        "dataset_display": str(record.get("dataset_display", "")),
        "timepoint_values": _expression_timepoint_values(timepoints),
        "has_expression": has_expression,
        "status": "Matched" if has_expression else "NA",
        "select_key": accession or gene_name or query,
    }


def search_expression_records(query: str, species: str, limit: int = 200) -> list[dict[str, Any]]:
    clean_query = query.strip()
    if not clean_query:
        return []

    records = list(load_expression_dataset(species).values())
    tokens = split_identifier_query(clean_query) or [clean_query]
    exact_tokens = {token.lower() for token in tokens if token.strip()}
    normalized_tokens = {normalize_identifier_token(token) for token in tokens if normalize_identifier_token(token)}

    exact_matches: list[dict[str, Any]] = []
    fuzzy_matches: list[dict[str, Any]] = []
    seen: set[str] = set()

    for record in records:
        accession = str(record.get("protein_accession", "")).strip()
        gene_name = str(record.get("gene_name", "")).strip()
        candidates = {accession.lower(), gene_name.lower()}
        normalized_candidates = {
            normalize_identifier_token(accession),
            normalize_identifier_token(gene_name),
        }
        if exact_tokens.intersection(candidates) or normalized_tokens.intersection(normalized_candidates):
            if accession not in seen:
                exact_matches.append(record)
                seen.add(accession)

    if exact_matches:
        return [
            build_proteome_overview_row(record, species, query=clean_query, has_expression=True)
            for record in exact_matches[:limit]
        ]

    normalized_query = normalize_identifier_token(clean_query)
    if len(normalized_query) < 3:
        return []

    for record in records:
        accession = str(record.get("protein_accession", "")).strip()
        search_text = _expression_search_text(record).lower()
        normalized_text = normalize_identifier_token(search_text)
        if normalized_query in normalized_text and accession not in seen:
            fuzzy_matches.append(record)
            seen.add(accession)
        if len(fuzzy_matches) >= limit:
            break

    fuzzy_matches.sort(
        key=lambda item: (
            str(item.get("gene_name", "")),
            str(item.get("protein_accession", "")),
        )
    )
    return [
        build_proteome_overview_row(record, species, query=clean_query, has_expression=True)
        for record in fuzzy_matches[:limit]
    ]


def build_proteome_blast_rows(blast_rows: list[dict[str, Any]], species: str) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for row in blast_rows:
        accession = str(row.get("protein_accession") or row.get("gene_name") or row.get("display_id") or "").strip()
        profile = build_expression_profile(accession, species)
        overview_record = profile if profile.get("available") else {
            "protein_accession": accession,
            "gene_name": str(row.get("gene_name", "")).strip(),
            "protein_description": str(row.get("description") or row.get("protein_description") or "").strip(),
        }
        overview = build_proteome_overview_row(
            overview_record,
            species,
            query=str(row.get("display_id", "")),
            has_expression=bool(profile.get("available")),
        )
        overview.update(
            {
                "display_id": row.get("display_id") or overview.get("display_id"),
                "description": row.get("description") or overview.get("description"),
                "e_value_display": row.get("e_value_display", "-"),
                "bit_score_display": row.get("bit_score_display", "-"),
                "identity_display": row.get("identity_display", "-"),
                "query_cover_display": row.get("query_cover_display", "-"),
                "gaps": row.get("gaps", "-"),
            }
        )
        output.append(overview)
    return output


def build_proteome_detail(identifier: str, species: str) -> dict[str, Any]:
    clean_identifier = identifier.strip()
    profile = build_expression_profile(clean_identifier, species)
    if not profile.get("available"):
        matches = search_expression_records(clean_identifier, species, limit=1)
        if matches:
            profile = build_expression_profile(str(matches[0].get("protein_accession", "")), species)

    accession = str(profile.get("protein_accession") or clean_identifier).strip()
    gene_name = str(profile.get("gene_name", "")).strip()
    description = str(profile.get("protein_description", "")).strip()
    sequence = fetch_protein_sequence_by_identifiers([accession, gene_name, clean_identifier])
    return {
        "profile": profile,
        "meta": {
            "protein_accession": accession,
            "gene_name": gene_name,
            "description": description,
            "organism": str(get_species_option(species).get("display", species)),
            "sequence_length": len(sequence),
        },
        "replicate_rows": _expression_replicate_rows(list(profile.get("timepoints", []))),
    }


def _format_csv_cell(value: Any) -> str:
    text = "" if value is None else str(value)
    if any(ch in text for ch in ['"', ",", "\r", "\n"]):
        return '"' + text.replace('"', '""') + '"'
    return text


def _csv_response(filename: str, headers: list[str], rows: list[list[Any]]) -> Response:
    output = io.StringIO()
    output.write(",".join(_format_csv_cell(header) for header in headers) + "\r\n")
    for row in rows:
        output.write(",".join(_format_csv_cell(cell) for cell in row) + "\r\n")
    return Response(
        content="\ufeff" + output.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def build_protein_summary_rows(identifier_tokens: list[str], species: str = DEFAULT_SPECIES) -> list[dict[str, Any]]:
    summary_rows: list[dict[str, Any]] = []
    emitted: set[tuple[str, str, str, str]] = set()

    for token in identifier_tokens:
        matched_rows = fetch_sites_for_identifier(token, species)
        if not matched_rows:
            summary_rows.append(build_protein_overview_row(token, []))
            continue

        grouped: dict[tuple[str, str, str, str, str, str], list[dict[str, str]]] = {}
        for row in matched_rows:
            key = (
                str(row.get("gene_name", "")).strip(),
                str(row.get("protein_accession", "")).strip(),
                str(row.get("organism", "")).strip(),
                str(row.get("protein_description", "")).strip(),
                str(row.get("transcript_id", "")).strip(),
                str(row.get("gene_names", "")).strip(),
            )
            grouped.setdefault(key, []).append(row)

        for key, rows in grouped.items():
            emit_key = (token.lower(), key[0].lower(), key[1].lower(), key[2].lower())
            if emit_key in emitted:
                continue
            emitted.add(emit_key)
            summary_rows.append(build_protein_overview_row(token, rows))

    return summary_rows


def build_site_download_rows(identifier: str, species: str = DEFAULT_SPECIES) -> tuple[list[list[Any]], str]:
    matched_rows = fetch_sites_for_identifier(identifier, species)
    matched_rows = keep_primary_target_rows(matched_rows)
    result_rows, result_meta = build_result_rows_and_meta(matched_rows)
    if not result_rows:
        return [], identifier

    export_rows: list[list[Any]] = []
    for row in result_rows:
        export_rows.append(
            [
                result_meta.get("gene_name", ""),
                result_meta.get("protein_accession", ""),
                result_meta.get("organism", ""),
                result_meta.get("protein_description", ""),
                row.get("Position display", ""),
                row.get("Amino acid display", ""),
                row.get("PTM types", ""),
                row.get("Subcellular localization", ""),
                row.get("W4 mean linear", ""),
                row.get("W12 mean linear", ""),
                row.get("W12/W4 ratio", ""),
                row.get("W12/W4 P value", ""),
                row.get("W12/W4 Type", ""),
            ]
        )
    filename_key = str(result_meta.get("gene_name") or result_meta.get("protein_accession") or identifier)
    return export_rows, filename_key


def _safe_download_name(value: str, fallback: str = "protein") -> str:
    cleaned = re.sub(r'[\\/:*?"<>|\s]+', "_", value.strip())
    cleaned = cleaned.strip("._")
    return cleaned or fallback


def _format_e_value(value: float) -> str:
    if not math.isfinite(value):
        return "-"
    if value == 0:
        return "0"
    if value < 0.001 or value >= 10000:
        return f"{value:.2e}"
    return f"{value:.4f}".rstrip("0").rstrip(".")


def _wrap_fasta_sequence(sequence: str, line_width: int = 80) -> str:
    clean = normalize_sequence_input(sequence)
    return "\n".join(clean[index:index + line_width] for index in range(0, len(clean), line_width))


def _blast_db_prefix(species: str) -> Path:
    return BLAST_DB_DIR / normalize_species_key(species) / "all_proteins"


def _blast_metadata_path(species: str) -> Path:
    return BLAST_DB_DIR / normalize_species_key(species) / "all_proteins.metadata.json"


def _blast_fasta_path(species: str) -> Path:
    return BLAST_DB_DIR / normalize_species_key(species) / "all_proteins.fasta"


def _blast_db_exists(prefix: Path) -> bool:
    return any(prefix.with_suffix(suffix).exists() for suffix in [".pin", ".psq", ".phr", ".pdb"])


def _blast_metadata_excludes_non_modification_classes(metadata: dict[str, Any]) -> bool:
    excluded = {code.upper() for code in NON_MODIFICATION_CLASSES}
    for item in metadata.values():
        if not isinstance(item, dict):
            return False
        ptm_codes = {code.strip().upper() for code in str(item.get("ptm_types", "")).split(",") if code.strip()}
        if ptm_codes.intersection(excluded):
            return False
    return True


def _check_blast_tools() -> tuple[str, str]:
    blastp_path = shutil.which("blastp")
    makeblastdb_path = shutil.which("makeblastdb")
    if not blastp_path or not makeblastdb_path:
        raise RuntimeError(
            "NCBI BLAST+ is not installed or not on PATH. Please install BLAST+ so blastp and makeblastdb are available."
        )
    return blastp_path, makeblastdb_path


def ensure_blast_database(species: str = DEFAULT_SPECIES) -> tuple[Path, dict[str, dict[str, Any]], int]:
    _, makeblastdb_path = _check_blast_tools()
    selected_species = normalize_species_key(species)
    prefix = _blast_db_prefix(selected_species)
    metadata_path = _blast_metadata_path(selected_species)
    fasta_path = _blast_fasta_path(selected_species)

    if _blast_db_exists(prefix) and metadata_path.exists():
        try:
            with metadata_path.open("r", encoding="utf-8") as fh:
                metadata = json.load(fh)
            if (
                isinstance(metadata, dict)
                and metadata
                and all("transcript_id" in item and "gene_names" in item for item in metadata.values() if isinstance(item, dict))
                and all(item.get("database_scope") == BLAST_METADATA_SCOPE for item in metadata.values() if isinstance(item, dict))
                and _blast_metadata_excludes_non_modification_classes(metadata)
            ):
                return prefix, metadata, len(metadata)
        except Exception:
            pass

    prefix.parent.mkdir(parents=True, exist_ok=True)
    candidates = fetch_blast_candidate_sequences(selected_species)
    if not candidates:
        raise RuntimeError("No protein sequences are available for this species.")

    metadata: dict[str, dict[str, Any]] = {}
    fasta_parts: list[str] = []
    for idx, candidate in enumerate(candidates, start=1):
        sequence = normalize_sequence_input(str(candidate.get("protein", "")))
        if not sequence:
            continue
        blast_id = f"ptm_{idx}"
        gene_name = str(candidate.get("gene_name", "")).strip()
        accession = str(candidate.get("protein_accession", "")).strip()
        metadata[blast_id] = {
            "gene_name": gene_name,
            "protein_accession": accession,
            "organism": str(candidate.get("organism", "")).strip(),
            "protein_description": str(candidate.get("protein_description", "")).strip(),
            "transcript_id": str(candidate.get("transcript_id", "")).strip(),
            "gene_names": str(candidate.get("gene_names", "")).strip(),
            "modification_type_count": str(candidate.get("modification_type_count", "0") or "0"),
            "total_sites": str(candidate.get("total_sites", "0") or "0"),
            "ptm_types": str(candidate.get("ptm_types", "")).strip(),
            "select_key": gene_name or accession,
            "database_scope": BLAST_METADATA_SCOPE,
        }
        fasta_parts.append(f">{blast_id}\n{_wrap_fasta_sequence(sequence)}")

    if not metadata:
        raise RuntimeError("No valid protein sequences are available for BLAST database creation.")

    fasta_path.write_text("\n".join(fasta_parts) + "\n", encoding="utf-8")
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")

    result = subprocess.run(
        [
            makeblastdb_path,
            "-in",
            str(fasta_path),
            "-dbtype",
            "prot",
            "-parse_seqids",
            "-out",
            str(prefix),
        ],
        capture_output=True,
        text=True,
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip() or result.stdout.strip() or "makeblastdb failed.")

    return prefix, metadata, len(metadata)


def build_blast_summary_rows(
    query_sequence: str,
    species: str = DEFAULT_SPECIES,
    max_hits: int = 20,
    identity_min: float = BLAST_DEFAULT_IDENTITY_MIN,
    e_value_max: float = BLAST_DEFAULT_E_VALUE_MAX,
    query_cover_min: float = BLAST_DEFAULT_QUERY_COVER_MIN,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    normalized_query = normalize_sequence_input(query_sequence)
    blastp_path, _ = _check_blast_tools()
    db_prefix, metadata, candidate_count = ensure_blast_database(species)
    hit_rows: list[dict[str, Any]] = []
    task = "blastp-short" if len(normalized_query) < 30 else "blastp"
    outfmt_fields = "6 qseqid sseqid pident length mismatch gaps qstart qend sstart send evalue bitscore qcovs"
    selected_max_hits = _clamp_int(max_hits, 5, 100)
    blast_target_limit = min(max(selected_max_hits * 5, 100), 500)

    with tempfile.TemporaryDirectory() as tmpdir:
        query_path = Path(tmpdir) / "query.fasta"
        output_path = Path(tmpdir) / "blast.tsv"
        query_path.write_text(f">Query_1\n{_wrap_fasta_sequence(normalized_query)}\n", encoding="utf-8")

        result = subprocess.run(
            [
                blastp_path,
                "-task",
                task,
                "-query",
                str(query_path),
                "-db",
                str(db_prefix),
                "-outfmt",
                outfmt_fields,
                "-evalue",
                str(e_value_max),
                "-max_target_seqs",
                str(blast_target_limit),
                "-out",
                str(output_path),
            ],
            capture_output=True,
            text=True,
        )
        if result.returncode != 0:
            raise RuntimeError(result.stderr.strip() or result.stdout.strip() or "blastp failed.")
        blast_lines = output_path.read_text(encoding="utf-8").splitlines() if output_path.exists() else []

    for line in blast_lines:
        parts = line.split("\t")
        if len(parts) < 13:
            continue
        (
            _query_id,
            subject_id,
            pident,
            alignment_length,
            _mismatch,
            gaps,
            _qstart,
            _qend,
            _sstart,
            _send,
            evalue,
            bitscore,
            qcovs,
        ) = parts[:13]
        candidate = metadata.get(subject_id, {})
        if not candidate:
            continue
        try:
            e_value_float = float(evalue)
        except ValueError:
            e_value_float = float("inf")
        try:
            bit_score_float = float(bitscore)
        except ValueError:
            bit_score_float = 0.0
        try:
            identity_float = float(pident)
        except ValueError:
            identity_float = 0.0
        try:
            query_cover_float = float(qcovs)
        except ValueError:
            query_cover_float = 0.0
        if identity_float < identity_min or e_value_float > e_value_max or query_cover_float < query_cover_min:
            continue
        gene_name = str(candidate.get("gene_name", "")).strip()
        accession = str(candidate.get("protein_accession", "")).strip()
        transcript_id = str(candidate.get("transcript_id", "")).strip()
        gene_names = str(candidate.get("gene_names", "")).strip()
        description = str(candidate.get("protein_description", "")).strip()
        total_sites = str(candidate.get("total_sites", "0") or "0")
        modification_type_count = str(candidate.get("modification_type_count", "0") or "0")
        hit_rows.append(
            {
                "query_id": "Query_1",
                "gene_name": gene_name,
                "protein_accession": accession,
                "uniprot_ids": accession,
                "organism": str(candidate.get("organism", "")).strip(),
                "protein_description": description,
                "description": description,
                "transcript_id": transcript_id,
                "gene_names": gene_names,
                "display_id": _format_overview_id("", gene_name, accession, transcript_id, gene_names),
                "plaza_gene_id": _format_plaza_gene_id(transcript_id, gene_names, gene_name),
                "modification_type_count": modification_type_count,
                "total_sites": total_sites,
                "has_ptm": total_sites not in {"", "0"},
                "ptm_types": str(candidate.get("ptm_types", "")).strip(),
                "select_key": str(candidate.get("select_key", "")).strip(),
                "bit_score": bit_score_float,
                "bit_score_display": f"{bit_score_float:.1f}",
                "e_value": e_value_float,
                "e_value_display": _format_e_value(e_value_float),
                "identity_percent": identity_float,
                "identity_display": f"{identity_float:.1f}%",
                "query_cover": query_cover_float,
                "query_cover_display": f"{query_cover_float:.1f}%",
                "gaps": gaps,
                "alignment_length": alignment_length,
            }
        )

    hit_rows.sort(
        key=lambda item: (
            float(item.get("e_value", float("inf"))),
            -float(item.get("bit_score", 0.0)),
            -float(item.get("identity_percent", 0.0)),
        )
    )
    selected_rows = hit_rows[:selected_max_hits]
    return selected_rows, {
        "candidate_count": candidate_count,
        "query_length": len(normalized_query),
        "matched_hits": len(selected_rows),
        "identity_min": identity_min,
        "e_value_max": e_value_max,
        "query_cover_min": query_cover_min,
    }


def _format_int(value: int) -> str:
    return f"{value:,}"


def _format_file_size(size_bytes: int) -> str:
    units = ["B", "KB", "MB", "GB"]
    size = float(size_bytes)
    for unit in units:
        if size < 1024 or unit == units[-1]:
            return f"{size:.1f} {unit}" if unit != "B" else f"{int(size)} {unit}"
        size /= 1024
    return f"{size_bytes} B"


def _build_home_data_tables() -> list[dict[str, Any]]:
    tables: list[dict[str, Any]] = []
    for config in HOME_DATA_TABLES:
        filename = str(config["filename"])
        dataset = Path(filename).stem
        path = IMPORT_READY_DIR / filename
        exists = path.exists() and path.is_file()
        row_count = 0
        column_count = 0
        columns_preview = ""
        if exists:
            try:
                with path.open("r", encoding="utf-8-sig", newline="") as fh:
                    reader = csv.reader(fh)
                    headers = next(reader, [])
                    column_count = len(headers)
                    columns_preview = ", ".join(headers[:6])
                    if len(headers) > 6:
                        columns_preview += ", ..."
                    row_count = sum(1 for _ in reader)
            except Exception:
                row_count = 0
                column_count = 0
                columns_preview = "Unable to read CSV"

        tables.append(
            {
                "filename": filename,
                "dataset": dataset,
                "species": config["species"],
                "data_type": config["data_type"],
                "exists": exists,
                "rows": row_count,
                "rows_fmt": _format_int(row_count),
                "columns": column_count,
                "columns_preview": columns_preview or "-",
                "size": _format_file_size(path.stat().st_size) if exists else "-",
                "download_url": f"/downloads/import-ready/{urllib.parse.quote(filename)}" if exists else "",
                "view_url": f"/data/table?dataset={urllib.parse.quote(dataset)}" if exists else "",
            }
        )
    return tables


def _get_data_table_config(dataset: str) -> dict[str, Any] | None:
    clean_dataset = dataset.strip()
    for config in HOME_DATA_TABLES:
        filename = str(config["filename"])
        if Path(filename).stem == clean_dataset or filename == clean_dataset:
            return {**config, "dataset": Path(filename).stem}
    return None


def _read_csv_page(path: Path, page: int, page_size: int) -> tuple[list[str], list[list[str]], int]:
    headers: list[str] = []
    rows: list[list[str]] = []
    total_rows = 0
    start_index = (page - 1) * page_size
    end_index = start_index + page_size

    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        headers = next(reader, [])
        for row_index, row in enumerate(reader):
            if start_index <= row_index < end_index:
                normalized = list(row[: len(headers)])
                if len(normalized) < len(headers):
                    normalized.extend([""] * (len(headers) - len(normalized)))
                rows.append(normalized)
            total_rows += 1

    return headers, rows, total_rows


def _build_pie_gradient(items: list[dict[str, Any]]) -> str:
    positive = [item for item in items if int(item.get("site_count", 0)) > 0]
    total = sum(int(item.get("site_count", 0)) for item in positive)
    if total <= 0:
        return "conic-gradient(#dfe8c9 0 100%)"

    parts: list[str] = []
    cursor = 0.0
    for item in positive:
        share = (int(item["site_count"]) / total) * 100.0
        start = cursor
        end = cursor + share
        parts.append(f"{item['bg']} {start:.2f}% {end:.2f}%")
        cursor = end

    if cursor < 100:
        parts.append(f"#dfe8c9 {cursor:.2f}% 100%")
    return "conic-gradient(" + ", ".join(parts) + ")"


def _build_pie_items_with_percent(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    total = sum(int(item.get("site_count", 0)) for item in items if int(item.get("site_count", 0)) > 0)

    ranked = sorted(items, key=lambda x: int(x.get("site_count", 0)), reverse=True)
    output: list[dict[str, Any]] = []
    for item in ranked:
        site_count = int(item.get("site_count", 0))
        percent = (site_count / total) * 100.0 if total > 0 else 0.0
        copied = dict(item)
        copied["percent"] = percent
        copied["percent_fmt"] = f"{percent:.1f}%"
        output.append(copied)
    return output


def _hex_to_rgb(color: str) -> tuple[int, int, int]:
    value = color.strip().lstrip("#")
    if len(value) == 3:
        value = "".join(ch * 2 for ch in value)
    if len(value) != 6:
        return 79, 119, 45
    try:
        return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)
    except ValueError:
        return 79, 119, 45


def _text_color_for_bg(color: str) -> str:
    r, g, b = _hex_to_rgb(color)
    luminance = 0.299 * r + 0.587 * g + 0.114 * b
    return "#111111" if luminance >= 165 else "#ffffff"


def _polar_to_cartesian(cx: float, cy: float, r: float, angle_deg: float) -> tuple[float, float]:
    rad = math.radians(angle_deg - 90.0)
    return cx + r * math.cos(rad), cy + r * math.sin(rad)


def _build_pie_svg(items: list[dict[str, Any]], min_label_percent: float = 7.0) -> str:
    positive = [item for item in items if int(item.get("site_count", 0)) > 0]
    total = sum(int(item.get("site_count", 0)) for item in positive)
    size = 240.0
    cx = cy = size / 2
    radius = 90.0

    if total <= 0:
        return (
            f'<svg class="pie-svg" viewBox="0 0 {size:.0f} {size:.0f}" '
            'xmlns="http://www.w3.org/2000/svg" aria-hidden="true">'
            f'<circle cx="{cx:.2f}" cy="{cy:.2f}" r="{radius:.2f}" fill="#dfe8c9"/>'
            "</svg>"
        )

    parts: list[str] = [
        f'<svg class="pie-svg" viewBox="0 0 {size:.0f} {size:.0f}" xmlns="http://www.w3.org/2000/svg" aria-hidden="true">'
    ]

    cursor = 0.0
    for item in positive:
        site_count = int(item.get("site_count", 0))
        percent = (site_count / total) * 100.0
        start_angle = cursor * 3.6
        end_angle = (cursor + percent) * 3.6
        large_arc = 1 if (end_angle - start_angle) > 180 else 0

        x1, y1 = _polar_to_cartesian(cx, cy, radius, start_angle)
        x2, y2 = _polar_to_cartesian(cx, cy, radius, end_angle)
        fill = str(item.get("bg", "#4f772d"))

        path = (
            f"M {cx:.2f} {cy:.2f} "
            f"L {x1:.2f} {y1:.2f} "
            f"A {radius:.2f} {radius:.2f} 0 {large_arc} 1 {x2:.2f} {y2:.2f} Z"
        )
        parts.append(f'<path d="{path}" fill="{fill}" stroke="#ffffff" stroke-width="1"/>')

        if percent >= min_label_percent:
            mid = (start_angle + end_angle) / 2.0
            tx, ty = _polar_to_cartesian(cx, cy, radius * 0.56, mid)
            text_fill = _text_color_for_bg(fill)
            font_size = 13 if percent >= 20 else 11
            label = str(item.get("code", "")).strip()
            label_text = f"{label} {percent:.1f}%" if label else f"{percent:.1f}%"
            parts.append(
                f'<text x="{tx:.2f}" y="{ty:.2f}" text-anchor="middle" dominant-baseline="middle" '
                f'fill="{text_fill}" font-size="{font_size}" font-weight="700" '
                f'font-family="Arial, sans-serif" paint-order="stroke" stroke="rgba(0,0,0,0.2)" stroke-width="0.6">{label_text}</text>'
            )

        cursor += percent

    parts.append("</svg>")
    return "".join(parts)


def _build_species_case_sql(column: str = "organism") -> str:
    clauses = []
    for option in SPECIES_OPTIONS:
        terms = [str(term).strip().lower() for term in option.get("terms", []) if str(term).strip()]
        if not terms:
            continue
        term_clauses = [f"LOWER(COALESCE({column}, '')) LIKE '%{term}%'" for term in terms]
        clauses.append(f"WHEN {' OR '.join(term_clauses)} THEN {pg_quote(str(option['value']))}")
    return "CASE " + " ".join(clauses) + " ELSE NULL END"


def _empty_species_stats(option: dict[str, Any]) -> dict[str, Any]:
    return {
        "key": str(option["value"]),
        "label": str(option.get("label", option["value"])),
        "display": str(option.get("display", option.get("label", option["value"]))),
        "total_sites": 0,
        "total_sites_fmt": "0",
        "total_proteins": 0,
        "total_proteins_fmt": "0",
        "ptm_items": [],
        "ptm_items_nonzero": [],
        "pie_items": [],
        "pie_svg": "",
    }


def _build_ptm_items_from_counts(per_type_map: dict[str, dict[str, int]]) -> list[dict[str, Any]]:
    ptm_items: list[dict[str, Any]] = []
    for code in PTM_PRIORITY:
        style = get_ptm_style(code)
        counts = per_type_map.get(code, {"site_count": 0, "protein_count": 0})
        ptm_items.append(
            {
                "code": code,
                "label": style["label"],
                "bg": style["bg"],
                "fg": style["fg"],
                "site_count": counts["site_count"],
                "protein_count": counts["protein_count"],
                "site_count_fmt": _format_int(counts["site_count"]),
                "protein_count_fmt": _format_int(counts["protein_count"]),
            }
        )
    return ptm_items


def fetch_home_ptm_stats() -> dict[str, Any]:
    species_case_sql = _build_species_case_sql("organism")
    modification_condition = build_modification_sql_condition("ptm_types")
    total_rows = run_psql_csv_query(
        f"""
        SELECT
          species_key,
          COUNT(*) FILTER (WHERE NULLIF(BTRIM(ptm_types), '') IS NOT NULL) AS total_sites,
          COUNT(DISTINCT COALESCE(NULLIF(BTRIM(gene_name), ''), NULLIF(BTRIM(protein_accession), ''))) AS total_proteins
        FROM (
          SELECT
            {species_case_sql} AS species_key,
            ptm_types,
            gene_name,
            protein_accession
          FROM ptm.gs_ptm
        ) classified
        WHERE species_key IS NOT NULL
          AND {modification_condition}
        GROUP BY species_key;
        """
    )

    per_type_rows = run_psql_csv_query(
        f"""
        SELECT
          species_key,
          BTRIM(ptm_types) AS ptm_type,
          COUNT(*) AS site_count,
          COUNT(DISTINCT COALESCE(NULLIF(BTRIM(gene_name), ''), NULLIF(BTRIM(protein_accession), ''))) AS protein_count
        FROM (
          SELECT
            {species_case_sql} AS species_key,
            ptm_types,
            gene_name,
            protein_accession
          FROM ptm.gs_ptm
        ) classified
        WHERE species_key IS NOT NULL
          AND {modification_condition}
          AND NULLIF(BTRIM(ptm_types), '') IS NOT NULL
        GROUP BY species_key, BTRIM(ptm_types);
        """
    )

    canonical_map = {code.upper(): code for code in PTM_PRIORITY}
    species_totals: dict[str, dict[str, int]] = {}
    for row in total_rows:
        species_key = str(row.get("species_key", "")).strip()
        if not species_key:
            continue
        try:
            total_sites = int(row.get("total_sites", "0") or 0)
        except ValueError:
            total_sites = 0
        try:
            total_proteins = int(row.get("total_proteins", "0") or 0)
        except ValueError:
            total_proteins = 0
        species_totals[species_key] = {"total_sites": total_sites, "total_proteins": total_proteins}

    per_species_type_map: dict[str, dict[str, dict[str, int]]] = {}
    for row in per_type_rows:
        species_key = str(row.get("species_key", "")).strip()
        if not species_key:
            continue
        raw = str(row.get("ptm_type", "")).strip()
        if not raw:
            continue
        code = canonical_map.get(raw.upper(), raw)
        try:
            site_count = int(row.get("site_count", "0") or 0)
        except ValueError:
            site_count = 0
        try:
            protein_count = int(row.get("protein_count", "0") or 0)
        except ValueError:
            protein_count = 0
        per_species_type_map.setdefault(species_key, {})[code] = {
            "site_count": site_count,
            "protein_count": protein_count,
        }

    species_stats: list[dict[str, Any]] = []
    for option in SPECIES_OPTIONS:
        species_stat = _empty_species_stats(option)
        species_key = str(option["value"])
        totals = species_totals.get(species_key, {"total_sites": 0, "total_proteins": 0})
        ptm_items = _build_ptm_items_from_counts(per_species_type_map.get(species_key, {}))
        pie_items = [item for item in ptm_items if item["site_count"] > 0]
        species_stat.update(
            {
                "total_sites": totals["total_sites"],
                "total_sites_fmt": _format_int(totals["total_sites"]),
                "total_proteins": totals["total_proteins"],
                "total_proteins_fmt": _format_int(totals["total_proteins"]),
                "ptm_items": ptm_items,
                "ptm_items_nonzero": pie_items,
                "pie_items": _build_pie_items_with_percent(ptm_items),
                "pie_svg": _build_pie_svg(pie_items),
            }
        )
        species_stats.append(species_stat)

    total_sites = sum(int(item.get("total_sites", 0)) for item in species_stats)
    total_proteins = sum(int(item.get("total_proteins", 0)) for item in species_stats)

    return {
        "total_sites": total_sites,
        "total_sites_fmt": _format_int(total_sites),
        "total_proteins": total_proteins,
        "total_proteins_fmt": _format_int(total_proteins),
        "species_stats": species_stats,
    }


_HOME_STATS_CACHE_LOCK = threading.Lock()
_HOME_STATS_CACHE: dict[str, Any] | None = None


def refresh_home_ptm_stats_cache() -> dict[str, Any]:
    global _HOME_STATS_CACHE
    stats = fetch_home_ptm_stats()
    with _HOME_STATS_CACHE_LOCK:
        _HOME_STATS_CACHE = stats
    return stats


def get_cached_home_ptm_stats() -> dict[str, Any]:
    global _HOME_STATS_CACHE
    with _HOME_STATS_CACHE_LOCK:
        if _HOME_STATS_CACHE is not None:
            return _HOME_STATS_CACHE

    stats = fetch_home_ptm_stats()

    with _HOME_STATS_CACHE_LOCK:
        if _HOME_STATS_CACHE is not None:
            return _HOME_STATS_CACHE
        _HOME_STATS_CACHE = stats
    return stats


def empty_home_ptm_stats() -> dict[str, Any]:
    species_stats = [_empty_species_stats(option) for option in SPECIES_OPTIONS]
    return {
        "total_sites": 0,
        "total_sites_fmt": "0",
        "total_proteins": 0,
        "total_proteins_fmt": "0",
        "species_stats": species_stats,
    }


@app.get("/", response_class=HTMLResponse)
async def home(request: Request) -> HTMLResponse:
    stats_error = ""
    stats = empty_home_ptm_stats()

    try:
        stats = get_cached_home_ptm_stats()
    except Exception:
        stats_error = "PTM statistics are temporarily unavailable. Please check database connection."

    return render_template(
        request,
        "home.html",
        title="Home",
        active="home",
        extra_context={
            "stats": stats,
            "stats_error": stats_error,
        },
    )


@app.get("/admin/refresh-home-stats")
async def refresh_home_stats() -> dict[str, Any]:
    stats = refresh_home_ptm_stats_cache()
    return {
        "status": "ok",
        "total_sites": stats.get("total_sites", 0),
        "total_proteins": stats.get("total_proteins", 0),
    }


@app.get("/downloads/import-ready/{filename}")
async def download_import_ready(filename: str) -> FileResponse:
    allowed = {str(item["filename"]) for item in HOME_DATA_TABLES}
    if filename not in allowed:
        raise HTTPException(status_code=404, detail="File not found")
    path = IMPORT_READY_DIR / filename
    if not path.exists() or not path.is_file():
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(
        path,
        media_type="text/csv",
        filename=filename,
    )


SITE_DOWNLOAD_HEADERS = [
    "Gene name",
    "Protein accession",
    "Organism",
    "Description",
    "Position",
    "Amino acid",
    "PTM types",
    "Subcellular localization",
    "W4 mean linear",
    "W12 mean linear",
    "W12/W4 ratio",
    "W12/W4 P value",
    "W12/W4 Type",
]


@app.get("/downloads/protein-sites.csv")
async def download_protein_sites(
    identifier: str = Query(default="", alias="identifier"),
    species: str = Query(default=DEFAULT_SPECIES, alias="species"),
) -> Response:
    selected_species = normalize_species_key(species)
    clean_identifier = identifier.strip()
    if not clean_identifier:
        raise HTTPException(status_code=400, detail="Missing identifier")

    rows, filename_key = build_site_download_rows(clean_identifier, selected_species)
    if not rows:
        raise HTTPException(status_code=404, detail="No PTM sites found")
    filename = f"protein-sites-{_safe_download_name(filename_key)}.csv"
    return _csv_response(filename, SITE_DOWNLOAD_HEADERS, rows)


@app.get("/downloads/protein-sites-batch.csv")
async def download_protein_sites_batch(
    identifiers: str = Query(default="", alias="identifiers"),
    species: str = Query(default=DEFAULT_SPECIES, alias="species"),
) -> Response:
    selected_species = normalize_species_key(species)
    tokens = split_identifier_query(identifiers)
    if not tokens:
        raise HTTPException(status_code=400, detail="Missing identifiers")

    all_rows: list[list[Any]] = []
    for token in tokens:
        rows, _ = build_site_download_rows(token, selected_species)
        all_rows.extend(rows)
    if not all_rows:
        raise HTTPException(status_code=404, detail="No PTM sites found")
    return _csv_response("protein-sites-batch.csv", SITE_DOWNLOAD_HEADERS, all_rows)


@app.get("/search", response_class=HTMLResponse)
async def search(request: Request) -> RedirectResponse:
    return RedirectResponse(url="/proteome", status_code=302)


@app.get("/proteome", response_class=HTMLResponse)
async def proteome_search(
    request: Request,
    tab: str = Query(default="identifier"),
    species: str = Query(default=DEFAULT_SPECIES, alias="species"),
    gene_name: str = Query(default="", alias="gene_name"),
    sequence: str = Query(default="", alias="sequence"),
    identity_min: float = Query(default=BLAST_DEFAULT_IDENTITY_MIN, alias="identity_min"),
    e_value_max: float = Query(default=BLAST_DEFAULT_E_VALUE_MAX, alias="e_value_max"),
    query_cover_min: float = Query(default=BLAST_DEFAULT_QUERY_COVER_MIN, alias="query_cover_min"),
    max_hits: int = Query(default=BLAST_DEFAULT_MAX_HITS, alias="max_hits"),
) -> HTMLResponse:
    search_tab = tab.strip().lower() or "identifier"
    if search_tab not in {"identifier", "sequence"}:
        search_tab = "identifier"
    selected_species = normalize_species_key(species)
    selected_species_label = str(get_species_option(selected_species).get("display", selected_species))
    selected_identity_min = _clamp_float(identity_min, 0, 100)
    selected_e_value_max = _clamp_float(e_value_max, 0, 1000)
    selected_query_cover_min = _clamp_float(query_cover_min, 0, 100)
    selected_max_hits = _clamp_int(max_hits, 5, 100)

    gene_query = gene_name.strip()
    sequence_query = sequence.strip()
    normalized_sequence_query = normalize_sequence_input(sequence_query)
    proteome_rows: list[dict[str, Any]] = []
    sequence_blast_summary = {
        "candidate_count": 0,
        "query_length": len(normalized_sequence_query),
        "matched_hits": 0,
        "identity_min": selected_identity_min,
        "e_value_max": selected_e_value_max,
        "query_cover_min": selected_query_cover_min,
    }
    message = ""

    if search_tab == "identifier" and gene_query:
        try:
            proteome_rows = search_expression_records(gene_query, selected_species)
            if proteome_rows:
                message = f"Found {len(proteome_rows)} protein expression profile rows in {selected_species_label}."
            else:
                message = f'No expression profile found for "{gene_query}" in {selected_species_label}.'
        except Exception:
            message = "Expression profile query failed. Please check the expression dataset files."

    if search_tab == "sequence" and sequence_query:
        try:
            if len(normalized_sequence_query) < 10:
                message = "Sequence is too short. Please input at least 10 amino acids for BLAST search."
            else:
                blast_rows, sequence_blast_summary = build_blast_summary_rows(
                    normalized_sequence_query,
                    species=selected_species,
                    max_hits=selected_max_hits,
                    identity_min=selected_identity_min,
                    e_value_max=selected_e_value_max,
                    query_cover_min=selected_query_cover_min,
                )
                proteome_rows = build_proteome_blast_rows(blast_rows, selected_species)
                if proteome_rows:
                    message = (
                        f"Found {len(proteome_rows)} BLAST hits in {selected_species_label}. "
                        "Proteins without quantitative expression data are shown as NA."
                    )
                else:
                    message = f"No BLAST hits passed the selected thresholds in {selected_species_label}."
        except Exception as exc:
            message = f"BLAST query failed: {exc}"

    return render_template(
        request,
        "proteome_search.html",
        title="Protein Expression Profiles",
        active="proteome",
        extra_context={
            "search_tab": search_tab,
            "species_options": SPECIES_OPTIONS,
            "selected_species": selected_species,
            "selected_species_label": selected_species_label,
            "gene_query": gene_query,
            "sequence_query": sequence_query,
            "proteome_rows": proteome_rows,
            "expression_timepoints": EXPRESSION_TIMEPOINTS,
            "sequence_blast_summary": sequence_blast_summary,
            "blast_identity_min": selected_identity_min,
            "blast_e_value_max": selected_e_value_max,
            "blast_query_cover_min": selected_query_cover_min,
            "blast_max_hits": selected_max_hits,
            "message": message,
        },
    )


@app.get("/proteome/detail", response_class=HTMLResponse)
async def proteome_detail(
    request: Request,
    species: str = Query(default=DEFAULT_SPECIES, alias="species"),
    identifier: str = Query(default="", alias="identifier"),
    return_tab: str = Query(default="identifier", alias="return_tab"),
    return_query: str = Query(default="", alias="return_query"),
) -> HTMLResponse:
    selected_species = normalize_species_key(species)
    clean_identifier = identifier.strip()
    detail = build_proteome_detail(clean_identifier, selected_species) if clean_identifier else {
        "profile": {},
        "meta": {},
        "replicate_rows": [],
    }
    message = "" if detail.get("profile", {}).get("available") else "No quantitative expression data matched this protein accession."

    return render_template(
        request,
        "proteome_detail.html",
        title="Protein Expression Profile",
        active="proteome",
        extra_context={
            "selected_species": selected_species,
            "selected_species_label": str(get_species_option(selected_species).get("display", selected_species)),
            "identifier": clean_identifier,
            "return_tab": return_tab,
            "return_query": return_query,
            "expression_profile": detail.get("profile", {}),
            "expression_meta": detail.get("meta", {}),
            "replicate_rows": detail.get("replicate_rows", []),
            "message": message,
        },
    )


@app.get("/search/protein", response_class=HTMLResponse)
async def search_protein(
    request: Request,
    tab: str = Query(default="identifier"),
    view: str = Query(default="overview", alias="view"),
    species: str = Query(default=DEFAULT_SPECIES, alias="species"),
    gene_name: str = Query(default="", alias="gene_name"),
    return_gene_name: str = Query(default="", alias="return_gene_name"),
    sequence: str = Query(default="", alias="sequence"),
    identity_min: float = Query(default=BLAST_DEFAULT_IDENTITY_MIN, alias="identity_min"),
    e_value_max: float = Query(default=BLAST_DEFAULT_E_VALUE_MAX, alias="e_value_max"),
    query_cover_min: float = Query(default=BLAST_DEFAULT_QUERY_COVER_MIN, alias="query_cover_min"),
    max_hits: int = Query(default=BLAST_DEFAULT_MAX_HITS, alias="max_hits"),
) -> HTMLResponse:
    search_tab = tab.strip().lower() or "identifier"
    if search_tab not in {"identifier", "sequence"}:
        search_tab = "identifier"
    result_view = view.strip().lower() or "overview"
    if result_view not in {"overview", "detail"}:
        result_view = "overview"
    selected_species = normalize_species_key(species)
    selected_species_label = str(get_species_option(selected_species).get("display", selected_species))
    selected_identity_min = _clamp_float(identity_min, 0, 100)
    selected_e_value_max = _clamp_float(e_value_max, 0, 1000)
    selected_query_cover_min = _clamp_float(query_cover_min, 0, 100)
    selected_max_hits = _clamp_int(max_hits, 5, 100)

    gene_query = gene_name.strip()
    return_gene_query = return_gene_name.strip()
    sequence_query = sequence.strip()
    normalized_sequence_query = normalize_sequence_input(sequence_query)

    result_rows: list[dict[str, Any]] = []
    sequence_rows: list[dict[str, Any]] = []
    legend_items: list[dict[str, str]] = []
    result_meta: dict[str, Any] = {}
    site_chart_data: list[dict[str, Any]] = []
    site_chart_default_id = ""
    expression_profile: dict[str, Any] = {}
    sequence_length = 0
    protein_sequence_export = ""
    message = ""
    identifier_match_rows: list[dict[str, Any]] = []
    show_identifier_match_table = False
    identifier_tokens = split_identifier_query(gene_query) if search_tab == "identifier" else []
    batch_summary_rows: list[dict[str, Any]] = []
    show_batch_summary_table = False
    sequence_blast_rows: list[dict[str, Any]] = []
    sequence_blast_summary = {
        "candidate_count": 0,
        "query_length": len(normalized_sequence_query),
        "matched_hits": 0,
        "identity_min": selected_identity_min,
        "e_value_max": selected_e_value_max,
        "query_cover_min": selected_query_cover_min,
    }

    if search_tab == "identifier" and gene_query:
        try:
            if result_view != "detail":
                tokens_for_overview = identifier_tokens if identifier_tokens else [gene_query]
                batch_summary_rows = build_protein_summary_rows(tokens_for_overview, selected_species)
                show_batch_summary_table = True
                matched_count = sum(1 for row in batch_summary_rows if row.get("has_result"))
                if matched_count:
                    message = (
                        f"Found {matched_count} matched protein overview rows from {len(tokens_for_overview)} submitted IDs. "
                        "Open a row to view detailed PTM sites."
                    )
                else:
                    message = f'No result found for "{gene_query}" in {selected_species_label}. Please try another identifier.'
                matched_rows = []
            else:
                matched_rows = fetch_sites_for_identifier(identifier_tokens[0] if identifier_tokens else gene_query, selected_species)

            if show_batch_summary_table:
                pass
            elif not matched_rows:
                message = f'No result found for "{gene_query}" in {selected_species_label}. Please try another Gene name.'
            else:
                identifier_match_rows = build_identifier_match_rows(matched_rows, gene_query)
                has_exact = any(bool(item.get("is_exact")) for item in identifier_match_rows)

                if not has_exact and len(identifier_match_rows) > 1:
                    show_identifier_match_table = True
                    message = (
                        f'Found {len(identifier_match_rows)} matches for "{gene_query}". '
                        "Please choose one row to view details."
                    )
                else:
                    selected_key = identifier_match_rows[0].get("select_key", "") if identifier_match_rows else gene_query
                    selected_rows = fetch_sites_for_identifier(str(selected_key), selected_species)
                    selected_rows = keep_primary_target_rows(selected_rows)

                    result_rows, result_meta = build_result_rows_and_meta(selected_rows)
                    if result_rows:
                        site_chart_data = build_site_chart_data(result_rows)
                        site_chart_default_id = str(result_rows[0].get("Site id", ""))
                        expression_profile = build_expression_profile(
                            str(result_meta.get("protein_accession", "")),
                            selected_species,
                        )
                        linked_sequence = fetch_protein_sequence_by_identifiers(
                            collect_identifier_tokens(
                                result_meta.get("gene_name", ""),
                                result_meta.get("protein_accession", ""),
                                result_meta.get("gene_names", ""),
                                result_meta.get("transcript_id", ""),
                                result_meta.get("protein_description", ""),
                            )
                        )
                        protein_sequence_export = linked_sequence
                        sequence_length = len(linked_sequence)
                        sequence_rows = build_sequence_rows(linked_sequence, result_rows)
                        legend_items = build_ptm_legend_from_sequence(sequence_rows)
                        if not sequence_rows:
                            message = (
                                "Matched GS PTM sites were found, but no linked protein sequence was found "
                                "in ptm.protein_sequence for this Gene name."
                            )
                    else:
                        message = f'No valid PTM position found for "{gene_query}".'
        except Exception:
            message = "Database query failed. Please check Docker/PostgreSQL status and imported data."

    if search_tab == "sequence" and sequence_query:
        try:
            if len(normalized_sequence_query) < 10:
                message = "Sequence is too short. Please input at least 10 amino acids for BLAST search."
            else:
                sequence_blast_rows, sequence_blast_summary = build_blast_summary_rows(
                    normalized_sequence_query,
                    species=selected_species,
                    max_hits=selected_max_hits,
                    identity_min=selected_identity_min,
                    e_value_max=selected_e_value_max,
                    query_cover_min=selected_query_cover_min,
                )
                if sequence_blast_rows:
                    message = (
                        f"Found {len(sequence_blast_rows)} BLAST hits in {selected_species_label}. "
                        "Proteins without PTM annotation are retained with #PTMs shown as 0."
                    )
                else:
                    message = f"No BLAST hits passed the selected thresholds in {selected_species_label}."
        except Exception as exc:
            message = f"BLAST query failed: {exc}"

    return render_template(
        request,
        "protein_search.html",
        title="Protein Search",
        active="search-protein",
        extra_context={
            "search_tab": search_tab,
            "result_view": result_view,
            "species_options": SPECIES_OPTIONS,
            "selected_species": selected_species,
            "selected_species_label": selected_species_label,
            "species_identifier_examples": SPECIES_IDENTIFIER_EXAMPLES,
            "gene_query": gene_query,
            "return_gene_query": return_gene_query,
            "sequence_query": sequence_query,
            "result_rows": result_rows,
            "sequence_rows": sequence_rows,
            "legend_items": legend_items,
            "result_meta": result_meta,
            "site_chart_data_json": json.dumps(site_chart_data, ensure_ascii=False),
            "site_chart_default_id": site_chart_default_id,
            "expression_profile": expression_profile,
            "sequence_length": sequence_length,
            "protein_sequence_export": protein_sequence_export,
            "message": message,
            "identifier_match_rows": identifier_match_rows,
            "show_identifier_match_table": show_identifier_match_table,
            "identifier_tokens": identifier_tokens,
            "batch_summary_rows": batch_summary_rows,
            "show_batch_summary_table": show_batch_summary_table,
            "sequence_blast_rows": sequence_blast_rows,
            "sequence_blast_summary": sequence_blast_summary,
            "blast_identity_min": selected_identity_min,
            "blast_e_value_max": selected_e_value_max,
            "blast_query_cover_min": selected_query_cover_min,
            "blast_max_hits": selected_max_hits,
        },
    )


@app.get("/search/ptm", response_class=HTMLResponse)
async def search_ptm(
    request: Request,
    species: str = Query(default=DEFAULT_SPECIES, alias="species"),
    expression: str = Query(default="", alias="expression"),
    sequence_length: int = Query(default=10, alias="sequence_length"),
    page: int = Query(default=1, alias="page"),
    page_size: int = Query(default=PTM_SEARCH_PAGE_SIZE_DEFAULT, alias="page_size"),
) -> HTMLResponse:
    selected_species = normalize_species_key(species)
    selected_species_label = str(get_species_option(selected_species).get("display", selected_species))
    expression_query = expression.strip()
    query_params = request.query_params
    if "representative_only_present" in query_params:
        representative_only = "representative_only" in query_params
    else:
        representative_only = True
    peptides_unique = _query_checkbox_checked(query_params, "peptides_unique", default=False)
    flank_length = _clamp_int(sequence_length, 3, 30)
    selected_page_size = page_size if page_size in PTM_SEARCH_PAGE_SIZE_OPTIONS else PTM_SEARCH_PAGE_SIZE_DEFAULT
    selected_page = max(1, page)

    result_rows: list[dict[str, Any]] = []
    grouped_rows: list[dict[str, Any]] = []
    summary = {"matched_proteins": 0, "matched_sites": 0, "truncated": 0}
    page_meta = {
        "total_rows": 0,
        "total_pages": 1,
        "current_page": 1,
        "page_size": selected_page_size,
        "start_row": 0,
        "end_row": 0,
        "prev_page": 1,
        "next_page": 1,
    }
    message = ""

    if expression_query:
        parsed, parse_error = parse_ptm_search_expression(expression_query)
        if parse_error:
            message = parse_error
        else:
            try:
                raw_matches, summary = run_ptm_expression_search(
                    parsed,
                    sequence_length=flank_length,
                    representative_only=representative_only,
                    peptides_unique=peptides_unique,
                    max_results=PTM_SEARCH_MAX_RESULTS,
                    species=selected_species,
                )

                all_site_rows = build_ptm_site_rows(raw_matches)
                result_rows, page_meta = paginate_rows(all_site_rows, selected_page, selected_page_size)
                grouped_rows = group_ptm_rows_by_protein(result_rows)

                if not all_site_rows:
                    message = f'No records matched "{expression_query}" in {selected_species_label}.'
                elif summary["truncated"]:
                    message = (
                        f"Result set is large. Showing up to {PTM_SEARCH_MAX_RESULTS} matches before pagination. "
                        "Refine expression for narrower results."
                    )
            except Exception:
                message = "PTM query failed. Please check Docker/PostgreSQL status and imported data."
    elif query_params:
        message = "Please enter a Search Expression."

    return render_template(
        request,
        "ptm_search.html",
        title="PTM Search",
        active="search-ptm",
        extra_context={
            "expression_query": expression_query,
            "species_options": SPECIES_OPTIONS,
            "selected_species": selected_species,
            "selected_species_label": selected_species_label,
            "sequence_length": flank_length,
            "representative_only": representative_only,
            "peptides_unique": peptides_unique,
            "fixed_species": selected_species_label,
            "result_rows": result_rows,
            "grouped_rows": grouped_rows,
            "summary": summary,
            "page_meta": page_meta,
            "message": message,
            "ptm_reference_rows": _build_ptm_reference_rows(),
            "max_results": PTM_SEARCH_MAX_RESULTS,
            "page_size_options": PTM_SEARCH_PAGE_SIZE_OPTIONS,
        },
    )


@app.get("/search/blast", response_class=HTMLResponse)
async def search_blast(
    request: Request,
    species: str = Query(default=DEFAULT_SPECIES, alias="species"),
    sequence: str = Query(default="", alias="sequence"),
    max_hits: int = Query(default=BLAST_DEFAULT_MAX_HITS, alias="max_hits"),
) -> RedirectResponse:
    selected_species = normalize_species_key(species)
    sequence_query = sequence.strip()
    selected_max_hits = _clamp_int(max_hits, 5, 100)
    params = {
        "tab": "sequence",
        "species": selected_species,
        "max_hits": str(selected_max_hits),
        "identity_min": str(BLAST_DEFAULT_IDENTITY_MIN),
        "e_value_max": str(BLAST_DEFAULT_E_VALUE_MAX),
        "query_cover_min": str(BLAST_DEFAULT_QUERY_COVER_MIN),
    }
    if sequence_query:
        params["sequence"] = sequence_query
    return RedirectResponse(
        url=f"{request.url_for('search_protein')}?{urllib.parse.urlencode(params)}",
        status_code=302,
    )


@app.get("/search/blast/detail", response_class=HTMLResponse)
async def search_blast_detail(
    request: Request,
    species: str = Query(default=DEFAULT_SPECIES, alias="species"),
    identifier: str = Query(default="", alias="identifier"),
    sequence: str = Query(default="", alias="sequence"),
    max_hits: int = Query(default=BLAST_DEFAULT_MAX_HITS, alias="max_hits"),
    identity_min: float = Query(default=BLAST_DEFAULT_IDENTITY_MIN, alias="identity_min"),
    e_value_max: float = Query(default=BLAST_DEFAULT_E_VALUE_MAX, alias="e_value_max"),
    query_cover_min: float = Query(default=BLAST_DEFAULT_QUERY_COVER_MIN, alias="query_cover_min"),
) -> HTMLResponse:
    selected_species = normalize_species_key(species)
    clean_identifier = identifier.strip()
    sequence_query = sequence.strip()
    selected_max_hits = _clamp_int(max_hits, 5, 100)
    selected_identity_min = _clamp_float(identity_min, 0, 100)
    selected_e_value_max = _clamp_float(e_value_max, 0, 1000)
    selected_query_cover_min = _clamp_float(query_cover_min, 0, 100)
    result_rows: list[dict[str, Any]] = []
    sequence_rows: list[dict[str, Any]] = []
    legend_items: list[dict[str, str]] = []
    result_meta: dict[str, Any] = {}
    site_chart_data: list[dict[str, Any]] = []
    site_chart_default_id = ""
    expression_profile: dict[str, Any] = {}
    sequence_length = 0
    protein_sequence_export = ""
    message = ""

    if not clean_identifier:
        message = "Missing BLAST hit identifier."
    else:
        try:
            matched_rows = fetch_sites_for_identifier(clean_identifier, selected_species)
            matched_rows = keep_primary_target_rows(matched_rows)
            result_rows, result_meta = build_result_rows_and_meta(matched_rows)
            if not result_rows:
                message = f'No PTM details found for "{clean_identifier}".'
            else:
                site_chart_data = build_site_chart_data(result_rows)
                site_chart_default_id = str(result_rows[0].get("Site id", ""))
                expression_profile = build_expression_profile(
                    str(result_meta.get("protein_accession", "")),
                    selected_species,
                )
                linked_sequence = fetch_protein_sequence_by_identifiers(
                    collect_identifier_tokens(
                        result_meta.get("gene_name", ""),
                        result_meta.get("protein_accession", ""),
                        result_meta.get("gene_names", ""),
                        result_meta.get("transcript_id", ""),
                        result_meta.get("protein_description", ""),
                    )
                )
                protein_sequence_export = linked_sequence
                sequence_length = len(linked_sequence)
                sequence_rows = build_sequence_rows(linked_sequence, result_rows)
                legend_items = build_ptm_legend_from_sequence(sequence_rows)
                if not sequence_rows:
                    message = (
                        "Matched PTM sites were found, but no linked protein sequence was found "
                        "in ptm.protein_sequence for this Gene name."
                    )
        except Exception:
            message = "BLAST hit detail query failed. Please check Docker/PostgreSQL status and imported data."

    return render_template(
        request,
        "blast_detail.html",
        title="BLAST Hit Details",
        active="search-protein",
        extra_context={
            "selected_species": selected_species,
            "identifier": clean_identifier,
            "sequence_query": sequence_query,
            "max_hits": selected_max_hits,
            "identity_min": selected_identity_min,
            "e_value_max": selected_e_value_max,
            "query_cover_min": selected_query_cover_min,
            "result_rows": result_rows,
            "sequence_rows": sequence_rows,
            "legend_items": legend_items,
            "result_meta": result_meta,
            "site_chart_data_json": json.dumps(site_chart_data, ensure_ascii=False),
            "site_chart_default_id": site_chart_default_id,
            "expression_profile": expression_profile,
            "sequence_length": sequence_length,
            "protein_sequence_export": protein_sequence_export,
            "message": message,
        },
    )


@app.get("/analysis", response_class=HTMLResponse)
async def analysis(request: Request) -> HTMLResponse:
    return render_page(request, "analysis", "analysis")


@app.get("/browse", response_class=HTMLResponse)
async def browse(request: Request) -> HTMLResponse:
    return render_page(request, "browse", "browse")


@app.get("/data/raw", response_class=HTMLResponse)
async def raw_data(request: Request) -> HTMLResponse:
    return render_template(
        request,
        "raw_data.html",
        title="Raw Data",
        active="data-raw",
        extra_context={"home_data_tables": _build_home_data_tables()},
    )


@app.get("/data/table", response_class=HTMLResponse)
async def data_table(
    request: Request,
    dataset: str = Query(default="soybean_gs", alias="dataset"),
    page: int = Query(default=1, alias="page"),
    page_size: int = Query(default=DATA_TABLE_PAGE_SIZE_DEFAULT, alias="page_size"),
) -> HTMLResponse:
    selected_page_size = page_size if page_size in DATA_TABLE_PAGE_SIZE_OPTIONS else DATA_TABLE_PAGE_SIZE_DEFAULT
    selected_page = max(1, page)
    config = _get_data_table_config(dataset)
    message = ""
    headers: list[str] = []
    rows: list[list[str]] = []
    total_rows = 0
    table_meta: dict[str, Any] = {
        "dataset": dataset,
        "filename": "",
        "species": "",
        "data_type": "",
        "download_url": "",
    }

    if config is None:
        message = "Unknown dataset."
    else:
        filename = str(config["filename"])
        path = IMPORT_READY_DIR / filename
        table_meta = {
            "dataset": str(config["dataset"]),
            "filename": filename,
            "species": str(config["species"]),
            "data_type": str(config["data_type"]),
            "download_url": f"/downloads/import-ready/{urllib.parse.quote(filename)}",
        }
        if not path.exists() or not path.is_file():
            message = "Dataset file is missing."
        else:
            try:
                headers, rows, total_rows = _read_csv_page(path, selected_page, selected_page_size)
            except Exception:
                message = "Unable to read dataset CSV."

    total_pages = max(1, math.ceil(total_rows / selected_page_size)) if total_rows > 0 else 1
    selected_page = _clamp_int(selected_page, 1, total_pages)
    if total_rows > 0 and rows == [] and not message:
        config_path = IMPORT_READY_DIR / str(config["filename"]) if config else Path()
        headers, rows, total_rows = _read_csv_page(config_path, selected_page, selected_page_size)

    page_meta = {
        "total_rows": total_rows,
        "total_pages": total_pages,
        "current_page": selected_page,
        "page_size": selected_page_size,
        "start_row": 0 if total_rows == 0 else ((selected_page - 1) * selected_page_size) + 1,
        "end_row": min(selected_page * selected_page_size, total_rows),
        "prev_page": selected_page - 1 if selected_page > 1 else 1,
        "next_page": selected_page + 1 if selected_page < total_pages else total_pages,
    }

    return render_template(
        request,
        "data_table.html",
        title="Data Table",
        active="data-raw",
        extra_context={
            "home_data_tables": _build_home_data_tables(),
            "table_meta": table_meta,
            "headers": headers,
            "rows": rows,
            "page_meta": page_meta,
            "page_size_options": DATA_TABLE_PAGE_SIZE_OPTIONS,
            "message": message,
        },
    )


@app.get("/data/submit", response_class=HTMLResponse)
async def submit_data(request: Request) -> HTMLResponse:
    return render_page(request, "data-submit", "data-submit")


@app.get("/submit-data", response_class=HTMLResponse)
async def submit_data_legacy(request: Request) -> HTMLResponse:
    return render_page(request, "data-submit", "data-submit")


@app.get("/about", response_class=HTMLResponse)
async def about(request: Request) -> HTMLResponse:
    return render_page(request, "about", "about")


if __name__ == "__main__":
    uvicorn.run("app:app", host="127.0.0.1", port=5000, reload=True)
